# api/services/export_word_service.py
from __future__ import annotations
from io import BytesIO
from typing import Optional, Tuple
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from api.core.presenters.text_sanitize import normalize_text_for_docx
from api.core.presenters.docx_helpers import (
    add_markdown_to_doc,
    add_data_agent_section,
)
from api.core.presenters.docx.config import load_docx_config_from_env
from api.core.presenters.extractors import extract_presentation
from shared.citations.citation_view_core import build_citation_core_views
from shared.http.download_filename import build_docx_filename_from_presentation

# ----------------------------
# Service
# ----------------------------

class ExportWordService:
    def build_docx_bytes(
        self,
        out: dict,
        filename: Optional[str] = "agent_output.docx",
    ) -> Tuple[bytes, str]:
        pres = extract_presentation(out)

        # 若未特別指定檔名（或仍用預設），就用「可含中文」的檔名
        if (not filename) or (filename == "agent_output.docx"):
            filename = build_docx_filename_from_presentation(
                pres,
                fallback="agent_output.docx",
                ext=".docx",
            )
        doc = Document()
        cfg = load_docx_config_from_env()
        
        # 預設字型
        style = doc.styles["Normal"]
        style.font.name = "標楷體"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "標楷體")
        style.font.size = Pt(12)

        doc.add_heading("AI 代理人輸出報告", level=1)

        # 使用者提問
        if pres.user_question:
            doc.add_heading("使用者提問", level=2)
            doc.add_paragraph(normalize_text_for_docx(pres.user_question))

        # 回答
        if pres.answer:
            doc.add_heading("AI 代理人回答", level=2)
            add_markdown_to_doc(doc, pres.answer)

        # 引用來源（改成吃同一套 shared view core）
        if pres.citations:
            doc.add_page_break()
            doc.add_heading("引用來源 / Citations", level=2)

            core_views = build_citation_core_views(pres.citations, text_limit=300)
            for v in core_views:
                doc_id = v.doc_id or ""

                # 一些顯示補充（可選）
                pid_str = f"（段落 {v.paragraph_id}）" if v.paragraph_id else ""
                date_str = f" · {v.date_iso}" if v.date_iso else ""

                # docx 文字清理（避免奇怪字元）
                short_txt = normalize_text_for_docx(v.short_text or "")

                # List Number 本身會編號，所以不用再塞 v.label
                p = doc.add_paragraph(style="List Number")
                # 第一段：標題/主旨
                p.add_run(f"{v.display}{pid_str}{date_str}")
                # 第二段：doc_id + 引文
                if doc_id:
                    p.add_run(f"\n({doc_id}) {short_txt}")
                else:
                    p.add_run(f"\n{short_txt}")

        # Data Agent
        add_data_agent_section(doc, pres.da_section, cfg=cfg)

        # 推理內容
        if pres.reasoning:
            doc.add_page_break()
            doc.add_heading("附錄：推理內容", level=2)
            add_markdown_to_doc(doc, pres.reasoning)

        buf = BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf.getvalue(), filename
