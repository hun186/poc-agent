# api/routers/export.py
from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import Response
from api.services.sessions import get_last_structured
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import re
from docx.oxml.ns import qn
from docx.shared import Pt

router = APIRouter()

def _add_bold_italic_runs(paragraph, text: str):
    """
    處理 **粗體** / *斜體* 的簡單情況。
    """
    # 先處理粗體，再處理斜體
    # 分段：保留 **...** 作為 token
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            inner = part[2:-2]
            run = paragraph.add_run(inner)
            run.bold = True
        else:
            # 再細拆 *斜體*
            sub_parts = re.split(r"(\*.*?\*)", part)
            for s in sub_parts:
                if not s:
                    continue
                if s.startswith("*") and s.endswith("*") and len(s) > 2:
                    inner_i = s[1:-1]
                    run = paragraph.add_run(inner_i)
                    run.italic = True
                else:
                    paragraph.add_run(s)

def _add_markdown_table(doc, header_line: str, sep_line: str, body_lines: list[str]):
    """
    把 markdown 表格轉成 Word 表格。
    header_line / body_lines 形如：| 日期 | 宮古水道 | ...
    """
    def _split_row(line: str) -> list[str]:
        line = line.strip()
        if line.startswith("|"):
            line = line[1:]
        if line.endswith("|"):
            line = line[:-1]
        return [cell.strip() for cell in line.split("|")]

    headers = _split_row(header_line)
    rows = [_split_row(l) for l in body_lines if l.strip()]

    if not headers:
        return

    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"

    # 表頭
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h

    # 內容
    for r_idx, row in enumerate(rows, start=1):
        cells = table.rows[r_idx].cells
        for c_idx, cell_text in enumerate(row):
            cells[c_idx].text = cell_text

def _normalize_hyphens_for_word(text: str) -> str:
    """
    把各種特殊連字號換成一般半形減號，避免 Word 出現灰底標記。
    """
    if not isinstance(text, str):
        return text

    return (
        text
        .replace("\u2010", "-")  # hyphen
        .replace("\u2011", "-")  # non-breaking hyphen
        .replace("\u2012", "-")  # figure dash
        .replace("\u2013", "-")  # en dash
        .replace("\u2014", "-")  # em dash
        .replace("\u2212", "-")  # minus sign
    )

def _add_markdown_to_doc(doc: Document, text: str):
    """
    把 answer 的 Markdown-ish 內容轉成較漂亮的 Word 排版：
    - 標題行（**一、...**）→ Heading 2
    - - 開頭 → 項目符號
    - 表格（| ... |）→ Word 表格
    - 其餘當一般段落，支援 **粗體** / *斜體*
    """
    text = _normalize_hyphens_for_word(text)
    lines = text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            # 空行 → 加一個空段落作為間距
            doc.add_paragraph("")
            i += 1
            continue

        # 1) Markdown 表格：連續含 '|' 的行 + 第二行是 --- 分隔
        if "|" in line and i + 1 < len(lines) and re.match(r"^\s*\|?\s*-", lines[i+1]):
            header_line = line
            sep_line = lines[i+1]
            body_lines = []
            j = i + 2
            while j < len(lines) and "|" in lines[j]:
                body_lines.append(lines[j])
                j += 1
            _add_markdown_table(doc, header_line, sep_line, body_lines)
            i = j
            continue

        # 2) 粗體標題行：例如 **一、資料概況**
        m_heading = re.match(r"^\s*\*\*(.+?)\*\*\s*$", line)
        if m_heading:
            title_text = m_heading.group(1).strip()
            p = doc.add_paragraph()
            p.style = "Heading 2"
            p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
            p.add_run(title_text)
            i += 1
            continue

        # 3) 一般 Markdown 標題：## 標題 / # 標題
        m_md_h = re.match(r"^\s*(#{1,3})\s+(.*)$", line)
        if m_md_h:
            level = len(m_md_h.group(1))
            title_text = m_md_h.group(2).strip()
            style_name = "Heading 1" if level == 1 else "Heading 2" if level == 2 else "Heading 3"
            p = doc.add_paragraph()
            p.style = style_name
            p.add_run(title_text)
            i += 1
            continue

        # 4) 引用（> 開頭）
        if line.lstrip().startswith(">"):
            content = line.lstrip()[1:].strip()
            p = doc.add_paragraph()
            p.style = "Intense Quote" if "Intense Quote" in [s.name for s in doc.styles if s.type == WD_STYLE_TYPE.PARAGRAPH] else None
            _add_bold_italic_runs(p, content)
            i += 1
            continue

        # 5) 項目符號行：- ... 或 * ...
        m_bullet = re.match(r"^\s*[-*]\s+(.*)$", line)
        if m_bullet:
            bullet_text = m_bullet.group(1)
            p = doc.add_paragraph(style="List Bullet")
            _add_bold_italic_runs(p, bullet_text)
            i += 1
            continue

        # 6) 一般段落
        p = doc.add_paragraph()
        _add_bold_italic_runs(p, line)
        i += 1


def _add_data_agent_section(doc: Document, out: dict):
    """
    依照 render_markdown 的邏輯，把 data_agent.runs[-1]
    寫成 Word 裡的「資料整備紀錄（Data Agent）」段落。
    """
    data_agent = out.get("data_agent") or {}
    da_runs = data_agent.get("runs") or []
    if not da_runs:
        return

    last_run = da_runs[-1]
    mode = last_run.get("mode") or out.get("mode")
    elapsed = last_run.get("elapsed_sec")
    question_raw = last_run.get("question") or out.get("question") or out.get("q")

    stats = last_run.get("stats") or {}
    query_meta = last_run.get("query_meta") or {}
    time_range = last_run.get("time_range") or {}

    core_terms = (
        query_meta.get("core_terms")
        or last_run.get("core_terms")
        or []
    )
    topic_terms = (
        query_meta.get("topic_terms")
        or last_run.get("topic_terms")
        or []
    )
    style_terms = (
        query_meta.get("style_terms")
        or last_run.get("style_terms")
        or []
    )

    syn_terms = query_meta.get("syn_terms") or []
    effective_terms = query_meta.get("effective_terms") or []

    # ==== 標題 ====
    doc.add_page_break()
    doc.add_heading("資料整備紀錄（Data Agent）", level=2)

    # ==== 基本資訊 ====
    if mode:
        p = doc.add_paragraph()
        p.add_run("模式：").bold = True
        p.add_run(str(mode))

    if elapsed is not None:
        p = doc.add_paragraph()
        p.add_run("資料準備耗時：").bold = True
        p.add_run(f"約 {elapsed} 秒")

    if question_raw:
        p = doc.add_paragraph()
        p.add_run("原始問題：").bold = True
        p.add_run(str(question_raw))

    # ==== 時間範圍 ====
    dt_from = time_range.get("date_from") or query_meta.get("date_from")
    dt_to = time_range.get("date_to") or query_meta.get("date_to")
    
    # 先正規化 dash
    dt_from_norm = _normalize_hyphens_for_word(dt_from or "N/A")
    dt_to_norm = _normalize_hyphens_for_word(dt_to or "N/A")
    
    if dt_from or dt_to:
        p = doc.add_paragraph()
        p.add_run("查詢時間範圍：").bold = True
        p.add_run(f"{dt_from_norm} ～ {dt_to_norm}")

    # ==== 統計資訊 ====
    if stats:
        doc_cnt = stats.get("doc_count")
        snip_cnt = stats.get("snippet_count")
        active_days = stats.get("active_days")
        full_doc_count = stats.get("full_doc_count")

        items = []
        if doc_cnt is not None:
            items.append(f"文件數：{doc_cnt}")
        if snip_cnt is not None:
            items.append(f"片段數：{snip_cnt}")
        if active_days is not None:
            items.append(f"有活動天數：{active_days}")
        if full_doc_count is not None:
            items.append(f"風格樣本文檔數：{full_doc_count}")

        if items:
            p = doc.add_paragraph()
            p.add_run("統計概況：").bold = True
            p.add_run("；".join(items))

    # ==== 關鍵字樣態 ====
    # report 模式：核心關鍵詞
    if mode == "report":
        if core_terms:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run("核心關鍵詞：").bold = True
            p.add_run(", ".join(core_terms))

    # style 模式：分風格 / 主題 / 合併核心
    elif mode == "style":
        if style_terms:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run("風格關鍵詞（style_terms）：").bold = True
            p.add_run(", ".join(style_terms))
        if topic_terms:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run("主題關鍵詞（topic_terms）：").bold = True
            p.add_run(", ".join(topic_terms))
        if core_terms:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run("檢索核心關鍵詞（style_terms + topic_terms）：").bold = True
            p.add_run(", ".join(core_terms))

    # ==== ES 查詢字串 ====
    es_query = (
        query_meta.get("es_query")
        or query_meta.get("core_query")
    )
    if mode == "style" and not es_query:
        es_query = last_run.get("style_query")

    if es_query:
        p = doc.add_paragraph()
        p.add_run("代理人 ES 文本查詢工具查詢字串：").bold = True
        p.add_run(str(es_query))

    # ==== 同義詞 / effective_terms ====
    if syn_terms:
        syn_str = ", ".join(syn_terms)
        if len(syn_str) > 120:
            syn_str = syn_str[:120] + "…"
        p = doc.add_paragraph()
        p.add_run("同義／擴充關鍵詞：").bold = True
        p.add_run(syn_str)

    if effective_terms:
        eff_str = ", ".join(effective_terms)
        if len(eff_str) > 120:
            eff_str = eff_str[:120] + "…"
        p = doc.add_paragraph()
        p.add_run("實際使用關鍵詞（effective terms）：").bold = True
        p.add_run(eff_str)
        
@router.get("/v1/export/word")
def export_word(session_id: str = Query(..., description="會話 session_id")):
    """
    依 session_id 取出上一輪 structured out，轉成美觀的 Word 檔。
    """
    last_structured = get_last_structured()
    out = last_structured.get(session_id)

    if not isinstance(out, dict):
        raise HTTPException(status_code=404, detail="找不到可匯出的報告內容（可能尚未產生結構化輸出）。")

    answer = out.get("answer") or ""
    citations = out.get("citations") or []
    reasoning = out.get("reasoning") or ""

    doc = Document()
    # === 設定預設字型為「標楷體」 ===
    style = doc.styles["Normal"]
    style.font.name = "標楷體"          # Windows: 標楷體（DFKai-SB）
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '標楷體')
    style.font.size = Pt(12)            # 可自行調整預設字級
    doc.add_heading("AI 代理人輸出報告", level=1)

    # 回答主文（用簡易 Markdown → Word）
    if answer:
        _add_markdown_to_doc(doc, answer)
        
    # 2) 資料整備紀錄（Data Agent）
    _add_data_agent_section(doc, out)
    
    # 引用來源（用乾淨 bullet 列）
    if citations:
        doc.add_page_break()
        doc.add_heading("引用來源 / Citations", level=2)
        for c in citations:
            txt = (c.get("text") or "").replace("\n", " ").strip()
            txt = _normalize_hyphens_for_word(txt)
            uuid_ = c.get("uuid") or c.get("source_id") or c.get("doc_id") or c.get("id") or ""
            if len(txt) > 300:
                txt = txt[:300] + "…"
            p = doc.add_paragraph(style="List Number")
            p.add_run(f"({uuid_}) {txt}")

    # 若你也想把 reasoning 放進去，可以當附錄
    if reasoning:
        doc.add_page_break()
        doc.add_heading("附錄：推理內容", level=2)
        _add_markdown_to_doc(doc, reasoning)

    from io import BytesIO
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)

    return Response(
        buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": 'attachment; filename="agent_output.docx"'
        },
    )
