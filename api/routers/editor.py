# api/routers/editor.py
from __future__ import annotations

from pathlib import Path
import hashlib
import json
from datetime import datetime, timezone
from typing import Any

from fastapi import APIRouter, Depends, File, Header, HTTPException, Query, UploadFile, status
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field

from agent.es import es_client
from api.services.auth_store import AUTH_STORE, AuthUser, create_token, parse_token
from api.services.editor_audit import append_audit_event, diff_values
from api.services.editor_form_config import get_editor_schema, public_editor_form_config

router = APIRouter(tags=["editor"])


class AuthPayload(BaseModel):
    username: str = Field(..., min_length=1, max_length=64)
    password: str = Field(..., min_length=1)
    full_name: str | None = Field(default=None, max_length=120)
    department: str | None = Field(default=None, max_length=120)
    note: str | None = Field(default=None, max_length=500)


class ApprovePayload(BaseModel):
    role: str = Field(...)


class SaveDocumentPayload(BaseModel):
    index: str = Field(..., min_length=1)
    id: str | None = Field(default=None, min_length=1)
    source: dict[str, Any]

class DeleteDocumentsPayload(BaseModel):
    index: str = Field(..., min_length=1)
    ids: list[str] = Field(..., min_length=1)


def _load_audit_records(limit: int = 5000) -> list[dict[str, Any]]:
    from api.services.editor_audit import audit_log_path

    path = audit_log_path()
    if not path.exists():
        return []

    records: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as f:
        for line in f:
            text = line.strip()
            if not text:
                continue
            try:
                obj = json.loads(text)
            except Exception:
                continue
            if isinstance(obj, dict):
                records.append(obj)
    if len(records) > limit:
        return records[-limit:]
    return records



def _make_uuid_from_source(source: dict[str, Any]) -> str:
    """Generate deterministic sha1 id from source payload.

    `importDT` is runtime metadata and should not affect duplicate detection.
    """
    normalized = dict(source)
    normalized.pop("importDT", None)
    basis = json.dumps(normalized, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha1(basis.encode("utf-8")).hexdigest()

def _pick_source_uuid(source: dict[str, Any]) -> str | None:
    value = source.get("uuid")
    if isinstance(value, list) and value:
        candidate = str(value[0] or "").strip()
        return candidate or None
    if isinstance(value, str):
        candidate = value.strip()
        return candidate or None
    return None


def _utc_now_iso_z() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def _extract_text_from_upload(filename: str, content: bytes) -> str:
    ext = Path(filename or "").suffix.lower()
    if ext in {".txt", ".md", ".csv"}:
        return content.decode("utf-8", errors="ignore")
    if ext == ".pdf":
        try:
            from pypdf import PdfReader  # type: ignore
            import io
            reader = PdfReader(io.BytesIO(content))
            return "\n".join((p.extract_text() or "") for p in reader.pages)
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"PDF 文字擷取失敗：{exc}") from exc
    if ext in {".docx"}:
        try:
            import docx  # type: ignore
            import io
            doc = docx.Document(io.BytesIO(content))
            return "\n".join(p.text for p in doc.paragraphs)
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"Word 文字擷取失敗：{exc}") from exc
    if ext in {".xlsx"}:
        try:
            import io
            import openpyxl  # type: ignore
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            lines: list[str] = []
            for ws in wb.worksheets:
                lines.append(f"# {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    vals = [str(c).strip() for c in row if c is not None and str(c).strip()]
                    if vals:
                        lines.append(" | ".join(vals))
            return "\n".join(lines)
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"Excel 文字擷取失敗：{exc}") from exc
    raise HTTPException(status_code=400, detail="不支援的檔案格式，請上傳 txt、md、csv、pdf、docx、xlsx。")


def _mapping_object_paths(properties: dict[str, Any], prefix: str = "") -> set[str]:
    paths: set[str] = set()
    for name, definition in (properties or {}).items():
        if not isinstance(definition, dict):
            continue
        path = f"{prefix}.{name}" if prefix else name
        field_type = str(definition.get("type") or "")
        child_props = definition.get("properties")
        if field_type in {"object", "nested"} or isinstance(child_props, dict):
            paths.add(path)
            if isinstance(child_props, dict):
                paths |= _mapping_object_paths(child_props, path)
    return paths


def _get_index_object_paths(index: str) -> set[str]:
    try:
        mapping = es_client.get_es().indices.get_mapping(index=index)
    except Exception:
        return set()

    object_paths: set[str] = set()
    for _, body in (mapping or {}).items():
        m = body.get("mappings") if isinstance(body, dict) else {}
        props = m.get("properties") if isinstance(m, dict) else {}
        if isinstance(props, dict):
            object_paths |= _mapping_object_paths(props)
    return object_paths


def _get_index_field_type_map(index: str, targets: list[str]) -> dict[str, str]:
    try:
        mapping = es_client.get_es().indices.get_mapping(index=index)
    except Exception:
        return {}

    out: dict[str, str] = {}

    def _walk(props: dict[str, Any], prefix: str = "") -> None:
        for name, definition in (props or {}).items():
            if not isinstance(definition, dict):
                continue
            path = f"{prefix}.{name}" if prefix else name
            ftype = str(definition.get("type") or "")
            if path in targets and ftype:
                out[path] = ftype
            child_props = definition.get("properties")
            if isinstance(child_props, dict):
                _walk(child_props, path)

    for _, body in (mapping or {}).items():
        m = body.get("mappings") if isinstance(body, dict) else {}
        props = m.get("properties") if isinstance(m, dict) else {}
        if isinstance(props, dict):
            _walk(props)

    return out


def _coerce_source_by_mapping(source: dict[str, Any], object_paths: set[str]) -> dict[str, Any]:
    patched = dict(source)
    for path in sorted(object_paths):
        parts = path.split(".")
        node: Any = patched
        ok = True
        for part in parts[:-1]:
            if not isinstance(node, dict) or part not in node:
                ok = False
                break
            node = node.get(part)
        if not ok or not isinstance(node, dict):
            continue
        leaf = parts[-1]
        if leaf not in node:
            continue
        value = node.get(leaf)
        if value is None or isinstance(value, dict) or isinstance(value, list):
            continue
        node[leaf] = {"value": value}
    return patched




def _normalize_sra_source_shapes(source: dict[str, Any], field_type_map: dict[str, str]) -> dict[str, Any]:
    """Normalize known SRP fields based on Elasticsearch mapping types.

    - nested  -> list[dict]
    - object  -> dict
    """
    normalized = dict(source)

    def _to_list_of_dict(value: Any) -> list[dict[str, Any]]:
        if isinstance(value, list):
            return [v for v in value if isinstance(v, dict)]
        if isinstance(value, dict):
            return [value]
        return []

    def _to_dict(value: Any) -> dict[str, Any]:
        if isinstance(value, dict):
            return value
        if isinstance(value, list):
            for item in value:
                if isinstance(item, dict):
                    return item
        return {}

    for key in ("multidim_event_json", "route_audit_json"):
        if key not in normalized:
            continue
        ftype = (field_type_map.get(key) or "").lower()
        if ftype == "nested":
            normalized[key] = _to_list_of_dict(normalized.get(key))
        elif ftype == "object":
            normalized[key] = _to_dict(normalized.get(key))

    return normalized

def _static_path(filename: str) -> Path:
    return Path(__file__).resolve().parents[1] / "static" / filename


@router.get("/editor", include_in_schema=False)
def editor_page() -> FileResponse:
    return FileResponse(_static_path("editor.html"))


def _bearer_token(authorization: str | None) -> str:
    if not authorization or not authorization.lower().startswith("bearer "):
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="請先登入。")
    return authorization.split(" ", 1)[1].strip()


def current_user(authorization: str | None = Header(default=None)) -> AuthUser:
    payload = parse_token(_bearer_token(authorization))
    if not payload:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="登入已失效，請重新登入。")
    username = str(payload.get("sub") or "")
    stored = AUTH_STORE.get_user(username)
    if not stored or stored.get("status") != "active":
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="帳號尚未啟用。")
    return AuthUser(username=username, role=str(stored.get("role") or "user"), status="active")


def require_superuser(user: AuthUser = Depends(current_user)) -> AuthUser:
    if user.role != "superuser":
        append_audit_event({
            "action": "authz.denied",
            "username": user.username,
            "role": user.role,
            "reason": "require_superuser",
        })
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="需要超級使用者權限。")
    return user


@router.post("/v1/auth/register")
def register(payload: AuthPayload) -> dict[str, Any]:
    try:
        user = AUTH_STORE.register(
            payload.username,
            payload.password,
            full_name=payload.full_name,
            department=payload.department,
            note=payload.note,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"user": user, "message": "帳號已建立，狀態為 pending；請等待超級使用者審核。"}


@router.post("/v1/auth/login")
def login(payload: AuthPayload) -> dict[str, Any]:
    user = AUTH_STORE.authenticate(payload.username, payload.password)
    if user is None:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="帳號或密碼錯誤。")
    if user.status != "active":
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail=f"帳號狀態為 {user.status}，尚不可使用。")
    return {"access_token": create_token(user), "token_type": "bearer", "user": user.__dict__}


@router.get("/v1/auth/me")
def me(user: AuthUser = Depends(current_user)) -> dict[str, Any]:
    return {"user": user.__dict__}


@router.get("/v1/auth/users")
def list_users(_: AuthUser = Depends(require_superuser)) -> dict[str, Any]:
    return {"users": AUTH_STORE.list_users()}


@router.post("/v1/auth/users/{username}/approve")
def approve_user(username: str, payload: ApprovePayload, user: AuthUser = Depends(require_superuser)) -> dict[str, Any]:
    try:
        approved = AUTH_STORE.approve_user(username, payload.role, approved_by=user.username)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"user": approved}


@router.post("/v1/auth/users/{username}/reject")
def reject_user(username: str, user: AuthUser = Depends(require_superuser)) -> dict[str, Any]:
    try:
        rejected = AUTH_STORE.reject_user(username, rejected_by=user.username)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    return {"user": rejected}


@router.delete("/v1/auth/users/{username}")
def delete_user(username: str, user: AuthUser = Depends(require_superuser)) -> dict[str, Any]:
    if username == user.username:
        raise HTTPException(status_code=400, detail="不可刪除目前登入的超級使用者。")
    try:
        deleted = AUTH_STORE.delete_user(username)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    return {"user": deleted}


@router.post("/v1/editor/extract-text")
async def extract_text_from_file(
    file: UploadFile = File(...),
    _: AuthUser = Depends(current_user),
) -> dict[str, Any]:
    content = await file.read()
    text = _extract_text_from_upload(file.filename or "", content)
    return {"filename": file.filename, "text": text}


def _current_indexes() -> list[str]:
    """Return currently-live ES indexes for editor dropdown (excluding hidden/system indexes)."""
    es_client.init_index_cache(live_check=True, force=True)
    live_from_cache = list(es_client.ALL_INDEXES)
    default_live = es_client.ES_INDEX_DEFAULT

    try:
        alias_map = es_client.get_es().indices.get_alias(index="*")
        raw_indexes = list((alias_map or {}).keys())
    except Exception:
        raw_indexes = live_from_cache

    merged: list[str] = []
    for idx in raw_indexes:
        s = str(idx or "").strip()
        if not s or s.startswith("."):
            continue
        if s not in merged:
            merged.append(s)

    if default_live and default_live not in merged:
        merged.insert(0, default_live)

    return merged


def _ensure_allowed_index(index: str) -> None:
    indexes = _current_indexes()
    if index not in indexes:
        raise HTTPException(status_code=400, detail=f"不允許或不存在的 index：{index}")


@router.get("/v1/editor/indexes")
def list_indexes(_: AuthUser = Depends(current_user)) -> dict[str, Any]:
    indexes = _current_indexes()
    default_index = es_client.ES_INDEX_DEFAULT or (indexes[0] if indexes else "")
    return {"indexes": indexes, "default_index": default_index}


@router.get("/v1/editor/form-config")
def editor_form_config(
    index: str | None = Query(None, description="用 index 選擇對應表單 schema"),
    schema: str | None = Query(None, description="明確指定 schema 名稱"),
    _: AuthUser = Depends(current_user),
) -> dict[str, Any]:
    try:
        return public_editor_form_config(index=index, schema_name=schema)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"讀取表單設定失敗：{exc}") from exc


@router.get("/v1/editor/search")
def search_documents(
    index: str = Query(..., description="要查詢的 Elasticsearch index"),
    q: str = Query("", description="關鍵字；空白時列出最近文件"),
    size: int = Query(20, ge=1, le=200),
    page: int = Query(1, ge=1),
    start_ts: str = Query("", description="起始時間（ISO8601）"),
    end_ts: str = Query("", description="結束時間（ISO8601）"),
    _: AuthUser = Depends(current_user),
) -> dict[str, Any]:
    _ensure_allowed_index(index)
    schema = get_editor_schema(index=index)
    filters: list[dict[str, Any]] = []
    date_field = str(schema.get("search_date_field") or "importDT")
    range_filter: dict[str, str] = {}
    if start_ts.strip():
        range_filter["gte"] = start_ts.strip()
    if end_ts.strip():
        range_filter["lte"] = end_ts.strip()
    if range_filter:
        filters.append({"range": {date_field: range_filter}})

    text_query: dict[str, Any]
    if q.strip():
        text_query = {
            "simple_query_string": {
                "query": q.strip(),
                "fields": schema.get("search_fields") or ["title^3", "subject^3", "uuid^2", "content", "full_text", "rawInfo.content"],
                "default_operator": "and",
            }
        }
    else:
        text_query = {"match_all": {}}

    query: dict[str, Any] = {"bool": {"must": [text_query]}}
    if filters:
        query["bool"]["filter"] = filters

    offset = (page - 1) * size
    res = es_client.get_es().search(
        index=index,
        body={
            "from": offset,
            "size": size,
            "track_total_hits": True,
            "query": query,
            "sort": [
                {date_field: {"order": "desc", "unmapped_type": "date", "missing": "_last"}},
                {"_score": {"order": "desc"}},
            ],
            "_source": schema.get("search_source_fields") or ["uuid", "title", "subject", "communication.subject", "rawInfo.title", "rawInfo.date", "importDT"],
        },
    )
    body = es_client._unwrap_es(res)
    hits = []
    for hit in body.get("hits", {}).get("hits", []) or []:
        src = hit.get("_source") or {}
        hits.append({
            "id": hit.get("_id"),
            "index": hit.get("_index"),
            "score": hit.get("_score"),
            "source": src,
        })
    total_obj = (body.get("hits", {}) or {}).get("total", 0)
    total = total_obj.get("value", 0) if isinstance(total_obj, dict) else int(total_obj or 0)
    return {"hits": hits, "total": total, "size": size, "page": page, "has_more": (offset + len(hits)) < total}


@router.get("/v1/editor/doc")
def get_document(
    index: str = Query(...),
    id: str = Query(...),
    _: AuthUser = Depends(current_user),
) -> dict[str, Any]:
    _ensure_allowed_index(index)
    try:
        got = es_client.get_es().get(index=index, id=id, _source=True)
    except Exception as exc:
        raise HTTPException(status_code=404, detail="文件不存在或無法讀取。") from exc
    body = es_client._unwrap_es(got)
    return {"index": body.get("_index", index), "id": body.get("_id", id), "source": body.get("_source") or {}}


@router.put("/v1/editor/doc")
def save_document(payload: SaveDocumentPayload, user: AuthUser = Depends(current_user)) -> dict[str, Any]:
    _ensure_allowed_index(payload.index)
    is_create = not (payload.id and payload.id.strip())
    if any(k.startswith("_") for k in payload.source.keys()):
        raise HTTPException(status_code=400, detail="source 不可包含 _ 開頭的 Elasticsearch metadata 欄位。")
    es = es_client.get_es()
    before_source: dict[str, Any] = {}
    import_dt = payload.source.get("importDT")
    if is_create and not (isinstance(import_dt, str) and import_dt.strip()):
        payload.source["importDT"] = _utc_now_iso_z()
    doc_id = payload.id.strip() if isinstance(payload.id, str) and payload.id.strip() else None
    if not doc_id:
        doc_id = _pick_source_uuid(payload.source)
    if not doc_id:
        doc_id = _make_uuid_from_source(payload.source)
        payload.source["uuid"] = doc_id

    if doc_id:
        try:
            before = es.get(index=payload.index, id=doc_id, _source=True)
            before_body = es_client._unwrap_es(before)
            before_source = before_body.get("_source") or {}
        except Exception:
            before_source = {}

    field_type_map = _get_index_field_type_map(payload.index, ["multidim_event_json", "route_audit_json"])
    normalized_source = _normalize_sra_source_shapes(payload.source, field_type_map)
    object_paths = _get_index_object_paths(payload.index)
    source_to_save = _coerce_source_by_mapping(normalized_source, object_paths)
    changes = diff_values(before_source, source_to_save)
    if not changes:
        return {
            "ok": True,
            "result": "noop",
            "index": payload.index,
            "id": doc_id,
            "updated_by": user.username,
            "changes": [],
            "message": "No changes detected; existing document loaded.",
            "mode": "edit_existing",
        }

    try:
        res = es.index(
            index=payload.index,
            id=doc_id,
            document=source_to_save,
            refresh="wait_for",
        )
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"寫入 Elasticsearch 失敗：{exc}") from exc
    body = es_client._unwrap_es(res)
    final_id = body.get("_id") or doc_id or ""
    append_audit_event({
        "action": "editor.create_document" if is_create else "editor.save_document",
        "username": user.username,
        "role": user.role,
        "index": payload.index,
        "id": final_id,
        "result": body.get("result"),
        "changed_paths": [item.get("path") for item in changes],
        "changes": changes,
    })
    return {"ok": True, "result": body.get("result"), "index": payload.index, "id": final_id, "updated_by": user.username, "changed_paths": [item.get("path") for item in changes]}


@router.delete("/v1/editor/doc")
def delete_document(
    index: str = Query(...),
    id: str = Query(...),
    user: AuthUser = Depends(current_user),
) -> dict[str, Any]:
    _ensure_allowed_index(index)
    try:
        res = es_client.get_es().delete(index=index, id=id, refresh="wait_for")
    except Exception as exc:
        raise HTTPException(status_code=404, detail="文件不存在或無法刪除。") from exc
    body = es_client._unwrap_es(res)
    append_audit_event({
        "action": "editor.delete_document",
        "username": user.username,
        "role": user.role,
        "index": index,
        "id": id,
        "result": body.get("result"),
        "changed_paths": [],
        "changes": [],
    })
    return {"ok": True, "index": index, "id": id, "result": body.get("result")}


@router.delete("/v1/editor/docs")
def delete_documents(payload: DeleteDocumentsPayload, user: AuthUser = Depends(current_user)) -> dict[str, Any]:
    _ensure_allowed_index(payload.index)
    ids = [str(i).strip() for i in payload.ids if str(i).strip()]
    if not ids:
        raise HTTPException(status_code=400, detail="ids 不可為空。")

    deleted: list[str] = []
    not_found: list[str] = []
    for doc_id in ids:
        try:
            res = es_client.get_es().delete(index=payload.index, id=doc_id, refresh="wait_for")
            body = es_client._unwrap_es(res)
            deleted.append(doc_id)
            append_audit_event({
                "action": "editor.delete_document",
                "username": user.username,
                "role": user.role,
                "index": payload.index,
                "id": doc_id,
                "result": body.get("result"),
                "changed_paths": [],
                "changes": [],
            })
        except Exception:
            not_found.append(doc_id)
    return {"ok": True, "index": payload.index, "deleted_ids": deleted, "not_found_ids": not_found}




@router.get("/v1/editor/usage-stats")
def usage_stats(user: AuthUser = Depends(current_user)) -> dict[str, Any]:
    records = _load_audit_records(limit=20000)
    mine = [r for r in records if str(r.get("username") or "") == user.username]

    def build(rows: list[dict[str, Any]]) -> dict[str, Any]:
        by_action: dict[str, int] = {}
        by_index: dict[str, int] = {}
        by_user: dict[str, int] = {}
        changed_paths_total = 0
        for r in rows:
            action = str(r.get("action") or "unknown")
            idx = str(r.get("index") or "-")
            uname = str(r.get("username") or "-")
            by_action[action] = by_action.get(action, 0) + 1
            by_index[idx] = by_index.get(idx, 0) + 1
            by_user[uname] = by_user.get(uname, 0) + 1
            changed_paths_total += len(r.get("changed_paths") or [])
        return {
            "total_events": len(rows),
            "changed_paths_total": changed_paths_total,
            "by_action": by_action,
            "by_index": by_index,
            "by_user": by_user,
        }

    out = {"me": build(mine)}
    if user.role == "superuser":
        out["all"] = build(records)
    return out

@router.get("/v1/editor/audit-logs")
def list_audit_logs(
    username: str | None = Query(None, description="依操作者帳號篩選"),
    action: str | None = Query(None, description="依操作類型篩選"),
    index: str | None = Query(None, description="依 index 篩選"),
    id: str | None = Query(None, description="依文件 ID 篩選"),
    changed_path: str | None = Query(None, description="依異動欄位路徑篩選（部分比對）"),
    start_ts: str | None = Query(None, description="起始時間（ISO 8601）"),
    end_ts: str | None = Query(None, description="結束時間（ISO 8601）"),
    limit: int = Query(100, ge=1, le=500),
    _: AuthUser = Depends(require_superuser),
) -> dict[str, Any]:
    records = _load_audit_records(limit=5000)
    username_q = (username or "").strip().lower()
    action_q = (action or "").strip().lower()
    index_q = (index or "").strip()
    id_q = (id or "").strip()
    path_q = (changed_path or "").strip().lower()
    start_q = (start_ts or "").strip()
    end_q = (end_ts or "").strip()

    def in_range(ts_text: str) -> bool:
        if not (start_q or end_q):
            return True
        try:
            ts = datetime.fromisoformat(ts_text.replace("Z", "+00:00"))
        except Exception:
            return False
        if start_q:
            try:
                start = datetime.fromisoformat(start_q.replace("Z", "+00:00"))
                if ts < start:
                    return False
            except Exception:
                raise HTTPException(status_code=400, detail="start_ts 格式錯誤，請使用 ISO 8601。")
        if end_q:
            try:
                end = datetime.fromisoformat(end_q.replace("Z", "+00:00"))
                if ts > end:
                    return False
            except Exception:
                raise HTTPException(status_code=400, detail="end_ts 格式錯誤，請使用 ISO 8601。")
        return True

    filtered: list[dict[str, Any]] = []
    for r in records:
        if username_q and username_q not in str(r.get("username") or "").lower():
            continue
        if action_q and action_q not in str(r.get("action") or "").lower():
            continue
        if index_q and index_q != str(r.get("index") or ""):
            continue
        if id_q and id_q != str(r.get("id") or ""):
            continue
        if path_q:
            paths = r.get("changed_paths") or []
            if not any(path_q in str(p or "").lower() for p in paths):
                continue
        if not in_range(str(r.get("ts") or "")):
            continue
        filtered.append(r)

    filtered.sort(key=lambda x: str(x.get("ts") or ""), reverse=True)
    sliced = filtered[:limit]
    return {"total": len(filtered), "limit": limit, "hits": sliced}
