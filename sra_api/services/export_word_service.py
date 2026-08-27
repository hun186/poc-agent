# api/services/export_word_service.py
from __future__ import annotations
from io import BytesIO
from typing import Optional, Tuple
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from api.core.presenters.text_sanitize import normalize_text_for_docx
from api.core.presenters.docx_helpers import (
    add_markdown_to_doc,
    add_data_agent_section,
)
from api.core.presenters.docx.config import load_docx_config_from_env
from api.core.presenters.extractors import extract_presentation
from shared.citations.citation_view_core import build_citation_core_views
from shared.http.download_filename import build_docx_filename_from_presentation

_MODEL_STAGE_ORDER = {
    stage: index
    for index, stage in enumerate(
        (
            "任務路由",
            "問題精煉",
            "資料整備",
            "證據欄位選擇",
            "報告生成",
            "風格生成",
            "一般回答生成",
            "分析管線",
        )
    )
}


def _sort_model_execution_rows(rows: list[dict]) -> list[dict]:
    """Return a stable processing-flow view of model execution records."""
    return sorted(
        rows,
        key=lambda item: _MODEL_STAGE_ORDER.get(str(item.get("stage") or "").strip(), len(_MODEL_STAGE_ORDER)),
    )


# ----------------------------
# Service
# ----------------------------

class ExportWordService:
    def build_docx_bytes(
        self,
        out: dict,
        filename: Optional[str] = "agent_output.docx",
    ) -> Tuple[bytes, str]:
        pres = extract_presentation(out)

        # 若未特別指定檔名（或仍用預設），就用「可含中文」的檔名
        if (not filename) or (filename == "agent_output.docx"):
            filename = build_docx_filename_from_presentation(
                pres,
                fallback="agent_output.docx",
                ext=".docx",
            )
        doc = Document()
        cfg = load_docx_config_from_env()
        
        # 預設字型
        style = doc.styles["Normal"]
        style.font.name = "標楷體"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "標楷體")
        style.font.size = Pt(12)

        doc.add_heading("AI 代理人輸出報告", level=1)

        # 使用者提問
        if pres.user_question:
            doc.add_heading("使用者提問", level=2)
            doc.add_paragraph(normalize_text_for_docx(pres.user_question))

        # 回答
        if pres.answer:
            doc.add_heading("AI 代理人回答", level=2)
            add_markdown_to_doc(doc, pres.answer)

        # 引用來源（改成吃同一套 shared view core）
        if pres.citations:
            doc.add_page_break()
            doc.add_heading("引用來源 / Citations", level=2)

            core_views = build_citation_core_views(pres.citations, text_limit=300)
            for v in core_views:
                doc_id = v.doc_id or ""

                # 一些顯示補充（可選）
                pid_str = f"（段落 {v.paragraph_id}）" if v.paragraph_id else ""
                date_str = f" · {v.date_iso}" if v.date_iso else ""

                # docx 文字清理（避免奇怪字元）
                short_txt = normalize_text_for_docx(v.short_text or "")

                # List Number 本身會編號，所以不用再塞 v.label
                p = doc.add_paragraph(style="List Number")
                # 第一段：標題/主旨
                p.add_run(f"{v.display}{pid_str}{date_str}")
                # 第二段：doc_id + 引文
                if doc_id:
                    p.add_run(f"\n({doc_id}) {short_txt}")
                else:
                    p.add_run(f"\n{short_txt}")

        # Data Agent
        add_data_agent_section(doc, pres.da_section, cfg=cfg)

        # 推理內容
        if pres.reasoning:
            doc.add_page_break()
            doc.add_heading("附錄：推理內容", level=2)
            add_markdown_to_doc(doc, pres.reasoning)

        # Keep model identity, completion status, and duration together at the
        # end of the document. Older saved sessions may only have model_usage,
        # so retain their original model-only display as a compatibility path.
        model_usage = pres.da_section.model_usage if pres.da_section else []
        execution_rows = _sort_model_execution_rows(pres.llm_execution or model_usage)
        if execution_rows:
            doc.add_page_break()
            doc.add_heading("各階段使用模型", level=2)
            status_labels = {
                "success": "正常完成",
                "fallback": "發生錯誤，已採用備援結果",
                "error": "執行錯誤",
            }
            for item in execution_rows:
                stage = normalize_text_for_docx(item.get("stage") or "未知階段")
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(f"{stage}：").bold = True
                model_id = item.get("model_id")
                if model_id:
                    p.add_run(normalize_text_for_docx(model_id))
                status = str(item.get("status") or "").lower()
                if status:
                    label = status_labels.get(status, normalize_text_for_docx(status))
                    p.add_run(f"；執行狀態：{label}")
                elapsed = item.get("elapsed_sec")
                if isinstance(elapsed, (int, float)):
                    p.add_run(f"；耗時：{elapsed:.3f} 秒")
                detail = item.get("detail")
                if detail:
                    p.add_run(f"\n錯誤／備註：{normalize_text_for_docx(detail)}")

        buf = BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf.getvalue(), filename
