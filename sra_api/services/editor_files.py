from __future__ import annotations

from pathlib import Path

from fastapi import HTTPException


def extract_text_from_upload(filename: str, content: bytes) -> str:
    ext = Path(filename or "").suffix.lower()
    if ext in {".txt", ".md", ".csv"}:
        return content.decode("utf-8", errors="ignore")
    if ext == ".pdf":
        try:
            import io

            from pypdf import PdfReader  # type: ignore

            reader = PdfReader(io.BytesIO(content))
            return "\n".join((p.extract_text() or "") for p in reader.pages)
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"PDF 文字擷取失敗：{exc}") from exc
    if ext in {".docx"}:
        try:
            import io

            import docx  # type: ignore

            doc = docx.Document(io.BytesIO(content))
            return "\n".join(p.text for p in doc.paragraphs)
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"Word 文字擷取失敗：{exc}") from exc
    if ext in {".xlsx"}:
        try:
            import io

            import openpyxl  # type: ignore

            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            lines: list[str] = []
            for ws in wb.worksheets:
                lines.append(f"# {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    vals = [str(c).strip() for c in row if c is not None and str(c).strip()]
                    if vals:
                        lines.append(" | ".join(vals))
            return "\n".join(lines)
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"Excel 文字擷取失敗：{exc}") from exc
    raise HTTPException(status_code=400, detail="不支援的檔案格式，請上傳 txt、md、csv、pdf、docx、xlsx。")
