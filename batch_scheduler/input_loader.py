from __future__ import annotations
from dataclasses import dataclass
from hashlib import sha256
from pathlib import Path
import csv, io, json
from openpyxl import load_workbook
from .models import BatchTaskSpec, BatchValidationError, DEFAULT_MODEL, parse_scheduled_at

FIELDS={'task_id','prompt','filename','session_id','model','scheduled_at','metadata','metadata_json'}
@dataclass(frozen=True)
class LoadedBatch:
    batch_id_hint: str | None
    source_path: Path
    source_sha256: str
    tasks: tuple[BatchTaskSpec,...]

def _task(row):
    unknown=set(row)-FIELDS
    if unknown: raise BatchValidationError(f"unknown task fields: {', '.join(sorted(unknown))}")
    metadata=row.get('metadata',{})
    if row.get('metadata_json') not in (None,''):
        try: metadata=json.loads(row['metadata_json'])
        except (TypeError,json.JSONDecodeError) as exc: raise BatchValidationError('metadata_json must be valid JSON') from exc
    if metadata in (None,''): metadata={}
    return BatchTaskSpec(str(row.get('task_id','')).strip(),str(row.get('prompt','')).strip(),row.get('filename') or None,row.get('session_id') or None,row.get('model') or DEFAULT_MODEL,parse_scheduled_at(row.get('scheduled_at')),metadata)

def load_batch(path: str|Path) -> LoadedBatch:
    path=Path(path); raw=path.read_bytes(); suffix=path.suffix.lower(); hint=None
    if suffix=='.csv': rows=list(csv.DictReader(io.StringIO(raw.decode('utf-8-sig'))))
    elif suffix=='.xlsx':
        wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True); values=wb.active.iter_rows(values_only=True)
        try: headers=[str(x) if x is not None else '' for x in next(values)]
        except StopIteration: headers=[]
        rows=[dict(zip(headers,row)) for row in values if any(x is not None for x in row)]
    elif suffix=='.json':
        try: obj=json.loads(raw.decode('utf-8-sig'))
        except (UnicodeDecodeError,json.JSONDecodeError) as exc: raise BatchValidationError('invalid JSON input') from exc
        if isinstance(obj,list): rows=obj
        elif isinstance(obj,dict) and set(obj)<= {'batch_id','tasks'} and isinstance(obj.get('tasks'),list): hint=obj.get('batch_id'); rows=obj['tasks']
        else: raise BatchValidationError('JSON must be a task list or batch envelope')
    else: raise BatchValidationError(f'unsupported input format: {suffix}')
    tasks=tuple(_task(dict(r)) for r in rows)
    ids=[t.task_id for t in tasks]
    if len(ids)!=len(set(ids)): raise BatchValidationError('duplicate task_id')
    if not tasks: raise BatchValidationError('batch contains no tasks')
    return LoadedBatch(hint,path,sha256(raw).hexdigest(),tasks)
