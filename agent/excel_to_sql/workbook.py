"""Excel workbook discovery, header mapping, and row normalization."""

from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any, Dict, Generator, List, Optional, Sequence, Tuple

from .constants import (
    DEFAULT_HASH_ALGORITHM,
    DEFAULT_HASH_HEX_LENGTH,
    EXCEL_SUFFIXES,
)
from .date_utils import is_date_mapping, normalize_cell_value, normalize_mapping_value
from .utils import (
    as_list,
    get_nested,
    is_empty_value,
    json_safe,
    normalize_header,
    normalize_header_key,
    require_nested,
    resolve_excel_path,
    resolve_path,
)


def parse_excel_column_ref(column_ref: str) -> int:
    text = str(column_ref).strip()
    if not text:
        raise ValueError("Excel 欄位參照不可為空")
    if text.isdigit():
        index = int(text)
        if index <= 0:
            raise ValueError(f"Excel 欄位序號必須大於 0：{column_ref}")
        return index
    result = 0
    for char in text.upper():
        if not "A" <= char <= "Z":
            raise ValueError(f"Excel 欄位字母格式錯誤：{column_ref}")
        result = result * 26 + (ord(char) - ord("A") + 1)
    return result


def read_header_map(ws, header_row: int) -> Dict[str, int]:
    header_cells = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
    header_map: Dict[str, int] = {}
    duplicates: Dict[str, int] = {}

    for idx, value in enumerate(header_cells, start=1):
        header = normalize_header(value)
        if not header:
            continue
        if header in header_map:
            duplicates[header] = duplicates.get(header, 1) + 1
            continue
        header_map[header] = idx

    if duplicates:
        dup_text = ", ".join(f"{k}({v})" for k, v in sorted(duplicates.items()))
        raise ValueError(f"Excel 標題列有重複欄名，請改用 source_column_index 或 source_column_letter：{dup_text}")

    return header_map


def resolve_header_candidate(
    *,
    mapping_index: int,
    header_map: Dict[str, int],
    expected_names: Sequence[str],
    match_mode: str,
) -> Tuple[int, str, str, str]:
    candidates_to_try = [name for name in expected_names if name]
    if not candidates_to_try:
        raise ValueError(f"mappings[{mapping_index}] 缺少可比對的來源欄名")

    available = ", ".join(header_map.keys())

    for expected in candidates_to_try:
        if expected in header_map:
            return header_map[expected], expected, "exact", ""

    normalized_headers: Dict[str, List[Tuple[str, int]]] = {}
    for header, idx in header_map.items():
        normalized_headers.setdefault(normalize_header_key(header), []).append((header, idx))

    if match_mode in {"smart", "normalized", "prefix", "contains"}:
        for expected in candidates_to_try:
            expected_key = normalize_header_key(expected)
            matches = normalized_headers.get(expected_key, [])
            if len(matches) == 1:
                header, idx = matches[0]
                return idx, header, "normalized", f"設定欄名 {expected} 以正規化欄名比對到 Excel 欄名 {header}"
            if len(matches) > 1:
                header, idx = sorted(matches, key=lambda item: item[1])[0]
                return (
                    idx,
                    header,
                    "ambiguous_normalized",
                    f"設定欄名 {expected} 正規化後符合多個欄名，已選用第 {idx} 欄 {header}；請檢查設定",
                )

    if match_mode in {"smart", "prefix", "contains"}:
        prefix_matches: List[Tuple[str, int, str]] = []
        for expected in candidates_to_try:
            expected_key = normalize_header_key(expected)
            if not expected_key:
                continue
            for header, idx in header_map.items():
                header_key = normalize_header_key(header)
                if header_key.startswith(expected_key):
                    prefix_matches.append((header, idx, expected))
        if len(prefix_matches) == 1:
            header, idx, expected = prefix_matches[0]
            return idx, header, "prefix", f"設定欄名 {expected} 以前綴智慧比對到 Excel 欄名 {header}"
        if len(prefix_matches) > 1:
            header, idx, expected = sorted(prefix_matches, key=lambda item: (len(item[0]), item[1]))[0]
            return (
                idx,
                header,
                "ambiguous_prefix",
                f"設定欄名 {expected} 前綴比對符合多個欄名，已選用第 {idx} 欄 {header}；請檢查設定",
            )

    if match_mode == "contains":
        contains_matches: List[Tuple[str, int, str]] = []
        for expected in candidates_to_try:
            expected_key = normalize_header_key(expected)
            if not expected_key:
                continue
            for header, idx in header_map.items():
                if expected_key in normalize_header_key(header):
                    contains_matches.append((header, idx, expected))
        if len(contains_matches) == 1:
            header, idx, expected = contains_matches[0]
            return idx, header, "contains", f"設定欄名 {expected} 以包含關係比對到 Excel 欄名 {header}"
        if len(contains_matches) > 1:
            header, idx, expected = sorted(contains_matches, key=lambda item: (len(item[0]), item[1]))[0]
            return (
                idx,
                header,
                "ambiguous_contains",
                f"設定欄名 {expected} 包含比對符合多個欄名，已選用第 {idx} 欄 {header}；請檢查設定",
            )

    raise ValueError(
        f"mappings[{mapping_index}] 找不到 Excel 欄名：{candidates_to_try[0]}；"
        f"目前標題列欄名：{available}"
    )


def read_column_mappings(cfg: Dict[str, Any], header_map: Dict[str, int]) -> List[Dict[str, Any]]:
    mappings = get_nested(cfg, "mappings", None)
    if not isinstance(mappings, list) or not mappings:
        raise ValueError("設定檔缺少必要欄位：mappings（至少一筆 source_column/dest_field 對應）")

    resolved: List[Dict[str, Any]] = []
    for index, mapping in enumerate(mappings, start=1):
        if not isinstance(mapping, dict):
            raise ValueError(f"mappings[{index}] 必須是物件")

        dest_field = mapping.get("dest_field") or mapping.get("destination_field")
        if not dest_field or not str(dest_field).strip():
            raise ValueError(f"mappings[{index}] 缺少 dest_field")

        source_name = mapping.get("source_column") or mapping.get("excel_column") or mapping.get("source_column_prefix")
        source_index = mapping.get("source_column_index")
        source_letter = mapping.get("source_column_letter")
        match_mode = str(mapping.get("match", "smart") or "smart").strip().lower()
        if mapping.get("source_column_prefix") and "match" not in mapping:
            match_mode = "prefix"
        if match_mode not in {"exact", "smart", "normalized", "prefix", "contains"}:
            raise ValueError(f"mappings[{index}] match 不支援：{match_mode}")

        source_label = ""
        if source_index is not None:
            excel_index = parse_excel_column_ref(str(source_index))
            source_label = str(source_index)
            actual_source_label = source_label
            match_type = "index"
            header_message = ""
        elif source_letter:
            excel_index = parse_excel_column_ref(str(source_letter))
            source_label = str(source_letter)
            actual_source_label = source_label
            match_type = "letter"
            header_message = ""
        elif source_name:
            source_label = normalize_header(source_name)
            expected_names = [source_label] + as_list(mapping.get("aliases"))
            excel_index, actual_source_label, match_type, header_message = resolve_header_candidate(
                mapping_index=index,
                header_map=header_map,
                expected_names=expected_names,
                match_mode=match_mode,
            )
        else:
            raise ValueError(
                f"mappings[{index}] 至少要設定 source_column、"
                "source_column_index 或 source_column_letter"
            )

        dest_field_text = str(dest_field).strip()
        value_type = str(
            mapping.get("value_type")
            or mapping.get("data_type")
            or mapping.get("type")
            or ("date" if is_date_mapping(source_label, dest_field_text) else "")
        ).strip().lower()
        if value_type and value_type not in {"date", "sql_date"}:
            raise ValueError(f"mappings[{index}] value_type 不支援：{value_type}")

        resolved_mapping = {
            "source_label": source_label,
            "actual_source_label": actual_source_label,
            "source_index": excel_index,
            "dest_field": dest_field_text,
            "required": bool(mapping.get("required", False)),
            "default": mapping.get("default", None),
            "match_type": match_type,
            "header_message": header_message,
            "value_type": value_type,
        }
        for type_key in ("dest_field_type", "destination_field_type", "sql_type"):
            if mapping.get(type_key):
                resolved_mapping[type_key] = mapping[type_key]
        resolved.append(resolved_mapping)

    return resolved


def make_hash_value(row_source_values: Dict[str, Any]) -> str:
    parts = []
    for key in sorted(row_source_values.keys()):
        value = json_safe(row_source_values[key])
        parts.append(f"{key}={'' if value is None else value}")
    basis = "|".join(parts)
    digest = hashlib.sha256(basis.encode("utf-8", errors="ignore")).hexdigest()
    if len(digest) != DEFAULT_HASH_HEX_LENGTH:
        raise ValueError(
            f"{DEFAULT_HASH_ALGORITHM} hex digest 長度應為 "
            f"{DEFAULT_HASH_HEX_LENGTH}，實際為 {len(digest)}"
        )
    return digest


def iter_excel_rows(
    *,
    excel_path: Path,
    sheet_name: str,
    header_row: int,
    start_row: int,
    end_row: int,
    mappings: Sequence[Dict[str, Any]],
    skip_empty_rows: bool,
) -> Generator[Tuple[int, Dict[str, Any], Dict[str, Any], str, str], None, None]:
    from openpyxl import load_workbook

    wb = load_workbook(filename=excel_path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Excel 找不到工作表：{sheet_name}；可用工作表：{', '.join(wb.sheetnames)}")
        ws = wb[sheet_name]

        max_col = max(int(mapping["source_index"]) for mapping in mappings)
        max_row = end_row if end_row > 0 else ws.max_row

        for row_no, row in enumerate(
            ws.iter_rows(
                min_row=start_row,
                max_row=max_row,
                min_col=1,
                max_col=max_col,
                values_only=True,
            ),
            start=start_row,
        ):
            record: Dict[str, Any] = {}
            source_values: Dict[str, Any] = {}
            non_empty = False
            invalid_message = ""

            for mapping in mappings:
                source_idx = int(mapping["source_index"])
                raw_value = row[source_idx - 1] if source_idx - 1 < len(row) else None
                value = normalize_cell_value(raw_value)

                if is_empty_value(value) and mapping.get("default") is not None:
                    value = mapping.get("default")

                if not is_empty_value(value):
                    try:
                        value = normalize_mapping_value(value, mapping)
                    except ValueError as exc:
                        invalid_message = (
                            f"Excel 第 {row_no} 列欄位日期格式錯誤："
                            f"{mapping['source_label']} -> {mapping['dest_field']}；{exc}"
                        )

                if bool(mapping.get("required", False)) and is_empty_value(value):
                    invalid_message = (
                        f"Excel 第 {row_no} 列必要欄位為空："
                        f"{mapping['source_label']} -> {mapping['dest_field']}"
                    )

                if not is_empty_value(value):
                    non_empty = True

                record[str(mapping["dest_field"])] = value
                source_values[str(mapping["source_label"])] = value

            if skip_empty_rows and not non_empty:
                continue

            yield row_no, record, source_values, make_hash_value(source_values), invalid_message
    finally:
        wb.close()


def open_workbook_header(
    config_path: str,
    cfg: Dict[str, Any],
    excel_path: Optional[Path] = None,
) -> Tuple[Path, str, int, Dict[str, int]]:
    if excel_path is None:
        excel_path_raw = require_nested(cfg, "source.excel_path")
        excel_path = resolve_excel_path(config_path, str(excel_path_raw))
    if not excel_path.exists():
        raise FileNotFoundError(f"找不到 Excel 檔案：{excel_path}")

    sheet_name = str(require_nested(cfg, "source.sheet"))
    header_row = int(get_nested(cfg, "source.header_row", 1) or 1)
    if header_row <= 0:
        raise ValueError("source.header_row 必須大於 0")

    from openpyxl import load_workbook

    wb = load_workbook(filename=excel_path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Excel 找不到工作表：{sheet_name}；可用工作表：{', '.join(wb.sheetnames)}")
        header_map = read_header_map(wb[sheet_name], header_row)
    finally:
        wb.close()

    return excel_path, sheet_name, header_row, header_map


def discover_excel_files(config_path: str, cfg: Dict[str, Any]) -> List[Path]:
    source_dir_raw = get_nested(cfg, "source.directory", "") or get_nested(cfg, "source.excel_dir", "") or ""
    if source_dir_raw:
        source_dir = resolve_path(config_path, str(source_dir_raw))
        if not source_dir.exists():
            raise FileNotFoundError(f"找不到來源目錄：{source_dir}")
        if not source_dir.is_dir():
            raise NotADirectoryError(f"來源目錄不是資料夾：{source_dir}")

        recursive = bool(get_nested(cfg, "source.recursive", False))
        pattern = str(get_nested(cfg, "source.file_pattern", "*.xlsx") or "*.xlsx")
        paths = source_dir.rglob(pattern) if recursive else source_dir.glob(pattern)
        files = [
            path.resolve()
            for path in paths
            if path.is_file() and path.suffix.lower() in EXCEL_SUFFIXES and not path.name.startswith("~$")
        ]
        return sorted(files)

    excel_path_raw = require_nested(cfg, "source.excel_path")
    return [resolve_excel_path(config_path, str(excel_path_raw))]
