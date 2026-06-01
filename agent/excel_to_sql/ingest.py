# agent/excel_to_sql/ingest.py
# cd /d D:\shared\TopicClassification\poc-agent
# python -m agent.excel_to_sql.ingest --config config/excel_to_sql/vessel_activities.example.json

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import sys
import time
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Generator, List, Optional, Sequence, Tuple

from .audit import build_header_audit_results, ensure_audit_table, insert_audit_rows
from .models import FileResult, RowResult
from .date_utils import (
    is_date_mapping,
    normalize_cell_value,
    normalize_date_text_for_parser,
    normalize_mapping_value,
    normalize_sql_date_value,
)
from .sql_utils import (
    build_conn_str_from_config,
    build_insert_sql,
    build_select_existing_hashes_sql,
    destination_column_exists,
    destination_hash_field_exists,
    detect_odbc_driver,
    ensure_destination_columns,
    ensure_destination_hash_field,
    quote_ident,
    quote_object_path,
    sql_string_literal,
    validate_hash_field_type,
)
from .utils import (
    as_list,
    chunked,
    get_nested,
    is_empty_value,
    json_safe,
    load_json_config,
    normalize_header,
    normalize_header_key,
    require_nested,
    resolve_excel_path,
    resolve_path,
)


EXCEL_SUFFIXES = {".xlsx", ".xlsm"}
DEFAULT_AUDIT_TABLE = "dbo.ExcelImportAudit"
DEFAULT_HASH_FIELD = "hash_value"
DEFAULT_HASH_ALGORITHM = "SHA-256"
DEFAULT_HASH_HEX_LENGTH = 64
DEFAULT_HASH_FIELD_TYPE = "CHAR(64)"


def enable_native_crash_tracebacks():
    """Ask Python to print a traceback if a native extension segfaults.

    A bare ``Segmentation fault (core dumped)`` is emitted by the OS after the
    interpreter dies in native code, so normal Python exceptions/retry handling
    cannot report the failing line. ``faulthandler`` gives the next crash a
    best-effort Python stack trace without changing import behavior.
    """

    try:
        import faulthandler

        if not faulthandler.is_enabled():
            faulthandler.enable(file=sys.stderr, all_threads=True)
    except Exception:
        # Diagnostics must never make ingestion fail on platforms where stderr
        # cannot be used with faulthandler.
        return


def parse_excel_column_ref(column_ref: str) -> int:
    text = str(column_ref).strip()
    if not text:
        raise ValueError("Excel 欄位參照不可為空")
    if text.isdigit():
        index = int(text)
        if index <= 0:
            raise ValueError(f"Excel 欄位序號必須大於 0：{column_ref}")
        return index
    result = 0
    for char in text.upper():
        if not "A" <= char <= "Z":
            raise ValueError(f"Excel 欄位字母格式錯誤：{column_ref}")
        result = result * 26 + (ord(char) - ord("A") + 1)
    return result


def read_header_map(ws, header_row: int) -> Dict[str, int]:
    header_cells = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
    header_map: Dict[str, int] = {}
    duplicates: Dict[str, int] = {}

    for idx, value in enumerate(header_cells, start=1):
        header = normalize_header(value)
        if not header:
            continue
        if header in header_map:
            duplicates[header] = duplicates.get(header, 1) + 1
            continue
        header_map[header] = idx

    if duplicates:
        dup_text = ", ".join(f"{k}({v})" for k, v in sorted(duplicates.items()))
        raise ValueError(f"Excel 標題列有重複欄名，請改用 source_column_index 或 source_column_letter：{dup_text}")

    return header_map


def normalize_header_key(value: Any) -> str:
    return "".join(str(value or "").strip().lower().split())


def as_list(value: Any) -> List[str]:
    if value is None:
        return []
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    text = str(value).strip()
    return [text] if text else []


def resolve_header_candidate(
    *,
    mapping_index: int,
    header_map: Dict[str, int],
    expected_names: Sequence[str],
    match_mode: str,
) -> Tuple[int, str, str, str]:
    candidates_to_try = [name for name in expected_names if name]
    if not candidates_to_try:
        raise ValueError(f"mappings[{mapping_index}] 缺少可比對的來源欄名")

    available = ", ".join(header_map.keys())

    for expected in candidates_to_try:
        if expected in header_map:
            return header_map[expected], expected, "exact", ""

    normalized_headers: Dict[str, List[Tuple[str, int]]] = {}
    for header, idx in header_map.items():
        normalized_headers.setdefault(normalize_header_key(header), []).append((header, idx))

    if match_mode in {"smart", "normalized", "prefix", "contains"}:
        for expected in candidates_to_try:
            expected_key = normalize_header_key(expected)
            matches = normalized_headers.get(expected_key, [])
            if len(matches) == 1:
                header, idx = matches[0]
                return idx, header, "normalized", f"設定欄名 {expected} 以正規化欄名比對到 Excel 欄名 {header}"
            if len(matches) > 1:
                header, idx = sorted(matches, key=lambda item: item[1])[0]
                return (
                    idx,
                    header,
                    "ambiguous_normalized",
                    f"設定欄名 {expected} 正規化後符合多個欄名，已選用第 {idx} 欄 {header}；請檢查設定",
                )

    if match_mode in {"smart", "prefix", "contains"}:
        prefix_matches: List[Tuple[str, int, str]] = []
        for expected in candidates_to_try:
            expected_key = normalize_header_key(expected)
            if not expected_key:
                continue
            for header, idx in header_map.items():
                header_key = normalize_header_key(header)
                if header_key.startswith(expected_key):
                    prefix_matches.append((header, idx, expected))
        if len(prefix_matches) == 1:
            header, idx, expected = prefix_matches[0]
            return idx, header, "prefix", f"設定欄名 {expected} 以前綴智慧比對到 Excel 欄名 {header}"
        if len(prefix_matches) > 1:
            header, idx, expected = sorted(prefix_matches, key=lambda item: (len(item[0]), item[1]))[0]
            return (
                idx,
                header,
                "ambiguous_prefix",
                f"設定欄名 {expected} 前綴比對符合多個欄名，已選用第 {idx} 欄 {header}；請檢查設定",
            )

    if match_mode == "contains":
        contains_matches: List[Tuple[str, int, str]] = []
        for expected in candidates_to_try:
            expected_key = normalize_header_key(expected)
            if not expected_key:
                continue
            for header, idx in header_map.items():
                if expected_key in normalize_header_key(header):
                    contains_matches.append((header, idx, expected))
        if len(contains_matches) == 1:
            header, idx, expected = contains_matches[0]
            return idx, header, "contains", f"設定欄名 {expected} 以包含關係比對到 Excel 欄名 {header}"
        if len(contains_matches) > 1:
            header, idx, expected = sorted(contains_matches, key=lambda item: (len(item[0]), item[1]))[0]
            return (
                idx,
                header,
                "ambiguous_contains",
                f"設定欄名 {expected} 包含比對符合多個欄名，已選用第 {idx} 欄 {header}；請檢查設定",
            )

    raise ValueError(
        f"mappings[{mapping_index}] 找不到 Excel 欄名：{candidates_to_try[0]}；"
        f"目前標題列欄名：{available}"
    )


def read_column_mappings(cfg: Dict[str, Any], header_map: Dict[str, int]) -> List[Dict[str, Any]]:
    mappings = get_nested(cfg, "mappings", None)
    if not isinstance(mappings, list) or not mappings:
        raise ValueError("設定檔缺少必要欄位：mappings（至少一筆 source_column/dest_field 對應）")

    resolved: List[Dict[str, Any]] = []
    for index, mapping in enumerate(mappings, start=1):
        if not isinstance(mapping, dict):
            raise ValueError(f"mappings[{index}] 必須是物件")

        dest_field = mapping.get("dest_field") or mapping.get("destination_field")
        if not dest_field or not str(dest_field).strip():
            raise ValueError(f"mappings[{index}] 缺少 dest_field")

        source_name = mapping.get("source_column") or mapping.get("excel_column") or mapping.get("source_column_prefix")
        source_index = mapping.get("source_column_index")
        source_letter = mapping.get("source_column_letter")
        match_mode = str(mapping.get("match", "smart") or "smart").strip().lower()
        if mapping.get("source_column_prefix") and "match" not in mapping:
            match_mode = "prefix"
        if match_mode not in {"exact", "smart", "normalized", "prefix", "contains"}:
            raise ValueError(f"mappings[{index}] match 不支援：{match_mode}")

        if source_index is not None:
            excel_index = parse_excel_column_ref(str(source_index))
            source_label = str(source_index)
            actual_source_label = source_label
            match_type = "index"
            header_message = ""
        elif source_letter:
            excel_index = parse_excel_column_ref(str(source_letter))
            source_label = str(source_letter)
            actual_source_label = source_label
            match_type = "letter"
            header_message = ""
        elif source_name:
            source_label = normalize_header(source_name)
            expected_names = [source_label] + as_list(mapping.get("aliases"))
            excel_index, actual_source_label, match_type, header_message = resolve_header_candidate(
                mapping_index=index,
                header_map=header_map,
                expected_names=expected_names,
                match_mode=match_mode,
            )
        else:
            raise ValueError(
                f"mappings[{index}] 至少要設定 source_column、"
                "source_column_index 或 source_column_letter"
            )

        dest_field_text = str(dest_field).strip()
        value_type = str(
            mapping.get("value_type")
            or mapping.get("data_type")
            or mapping.get("type")
            or ("date" if is_date_mapping(source_label, dest_field_text) else "")
        ).strip().lower()
        if value_type and value_type not in {"date", "sql_date"}:
            raise ValueError(f"mappings[{index}] value_type 不支援：{value_type}")

        resolved_mapping = {
            "source_label": source_label,
            "actual_source_label": actual_source_label,
            "source_index": excel_index,
            "dest_field": dest_field_text,
            "required": bool(mapping.get("required", False)),
            "default": mapping.get("default", None),
            "match_type": match_type,
            "header_message": header_message,
            "value_type": value_type,
        }
        for type_key in ("dest_field_type", "destination_field_type", "sql_type"):
            if mapping.get(type_key):
                resolved_mapping[type_key] = mapping[type_key]
        resolved.append(resolved_mapping)

    return resolved


def make_hash_value(row_source_values: Dict[str, Any]) -> str:
    parts = []
    for key in sorted(row_source_values.keys()):
        value = json_safe(row_source_values[key])
        parts.append(f"{key}={'' if value is None else value}")
    basis = "|".join(parts)
    digest = hashlib.sha256(basis.encode("utf-8", errors="ignore")).hexdigest()
    if len(digest) != DEFAULT_HASH_HEX_LENGTH:
        raise ValueError(
            f"{DEFAULT_HASH_ALGORITHM} hex digest 長度應為 "
            f"{DEFAULT_HASH_HEX_LENGTH}，實際為 {len(digest)}"
        )
    return digest


def iter_excel_rows(
    *,
    excel_path: Path,
    sheet_name: str,
    header_row: int,
    start_row: int,
    end_row: int,
    mappings: Sequence[Dict[str, Any]],
    skip_empty_rows: bool,
) -> Generator[Tuple[int, Dict[str, Any], Dict[str, Any], str, str], None, None]:
    from openpyxl import load_workbook

    wb = load_workbook(filename=excel_path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Excel 找不到工作表：{sheet_name}；可用工作表：{', '.join(wb.sheetnames)}")
        ws = wb[sheet_name]

        max_col = max(int(mapping["source_index"]) for mapping in mappings)
        max_row = end_row if end_row > 0 else ws.max_row

        for row_no, row in enumerate(
            ws.iter_rows(
                min_row=start_row,
                max_row=max_row,
                min_col=1,
                max_col=max_col,
                values_only=True,
            ),
            start=start_row,
        ):
            record: Dict[str, Any] = {}
            source_values: Dict[str, Any] = {}
            non_empty = False
            invalid_message = ""

            for mapping in mappings:
                source_idx = int(mapping["source_index"])
                raw_value = row[source_idx - 1] if source_idx - 1 < len(row) else None
                value = normalize_cell_value(raw_value)

                if is_empty_value(value) and mapping.get("default") is not None:
                    value = mapping.get("default")

                if not is_empty_value(value):
                    try:
                        value = normalize_mapping_value(value, mapping)
                    except ValueError as exc:
                        invalid_message = (
                            f"Excel 第 {row_no} 列欄位日期格式錯誤："
                            f"{mapping['source_label']} -> {mapping['dest_field']}；{exc}"
                        )

                if bool(mapping.get("required", False)) and is_empty_value(value):
                    invalid_message = (
                        f"Excel 第 {row_no} 列必要欄位為空："
                        f"{mapping['source_label']} -> {mapping['dest_field']}"
                    )

                if not is_empty_value(value):
                    non_empty = True

                record[str(mapping["dest_field"])] = value
                source_values[str(mapping["source_label"])] = value

            if skip_empty_rows and not non_empty:
                continue

            yield row_no, record, source_values, make_hash_value(source_values), invalid_message
    finally:
        wb.close()


def open_workbook_header(config_path: str, cfg: Dict[str, Any], excel_path: Optional[Path] = None) -> Tuple[Path, str, int, Dict[str, int]]:
    if excel_path is None:
        excel_path_raw = require_nested(cfg, "source.excel_path")
        excel_path = resolve_excel_path(config_path, str(excel_path_raw))
    if not excel_path.exists():
        raise FileNotFoundError(f"找不到 Excel 檔案：{excel_path}")

    sheet_name = str(require_nested(cfg, "source.sheet"))
    header_row = int(get_nested(cfg, "source.header_row", 1) or 1)
    if header_row <= 0:
        raise ValueError("source.header_row 必須大於 0")

    from openpyxl import load_workbook

    wb = load_workbook(filename=excel_path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Excel 找不到工作表：{sheet_name}；可用工作表：{', '.join(wb.sheetnames)}")
        header_map = read_header_map(wb[sheet_name], header_row)
    finally:
        wb.close()

    return excel_path, sheet_name, header_row, header_map


def write_log(log_dir: Path, event: Dict[str, Any]):
    log_dir.mkdir(parents=True, exist_ok=True)
    log_file = log_dir / f"excel_import_{datetime.utcnow().strftime('%Y%m%d')}.jsonl"
    payload = {"logged_at": datetime.utcnow().isoformat(timespec="seconds") + "Z", **event}
    with log_file.open("a", encoding="utf-8") as fp:
        fp.write(json.dumps(payload, ensure_ascii=False, default=json_safe) + "\n")


def unique_archive_path(target_dir: Path, source_file: Path) -> Path:
    target_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.utcnow().strftime("%Y%m%d%H%M%S")
    candidate = target_dir / f"{source_file.stem}_{timestamp}{source_file.suffix}"
    counter = 1
    while candidate.exists():
        candidate = target_dir / f"{source_file.stem}_{timestamp}_{counter}{source_file.suffix}"
        counter += 1
    return candidate


def move_to_archive(source_file: Path, target_dir: Path) -> Path:
    target = unique_archive_path(target_dir, source_file)
    shutil.move(str(source_file), str(target))
    return target


def discover_excel_files(config_path: str, cfg: Dict[str, Any]) -> List[Path]:
    source_dir_raw = get_nested(cfg, "source.directory", "") or get_nested(cfg, "source.excel_dir", "") or ""
    if source_dir_raw:
        source_dir = resolve_path(config_path, str(source_dir_raw))
        if not source_dir.exists():
            raise FileNotFoundError(f"找不到來源目錄：{source_dir}")
        if not source_dir.is_dir():
            raise NotADirectoryError(f"來源目錄不是資料夾：{source_dir}")

        recursive = bool(get_nested(cfg, "source.recursive", False))
        pattern = str(get_nested(cfg, "source.file_pattern", "*.xlsx") or "*.xlsx")
        paths = source_dir.rglob(pattern) if recursive else source_dir.glob(pattern)
        files = [
            path.resolve()
            for path in paths
            if path.is_file() and path.suffix.lower() in EXCEL_SUFFIXES and not path.name.startswith("~$")
        ]
        return sorted(files)

    excel_path_raw = require_nested(cfg, "source.excel_path")
    return [resolve_excel_path(config_path, str(excel_path_raw))]


def get_existing_hashes(cursor, *, table: str, hash_field: str, hash_values: Sequence[str]) -> set[str]:
    hashes = sorted({value for value in hash_values if value})
    if not hashes:
        return set()

    existing: set[str] = set()
    for batch in chunked(hashes, 900):
        cursor.execute(build_select_existing_hashes_sql(table, hash_field, len(batch)), batch)
        existing.update(str(row[0]) for row in cursor.fetchall())
    return existing


def insert_batch_with_fallback(
    cursor,
    *,
    insert_sql: str,
    rows: Sequence[Tuple[Any, ...]],
    row_contexts: Sequence[Tuple[int, str, Dict[str, Any]]],
) -> Tuple[List[RowResult], List[Tuple[str, Tuple[Any, ...]]]]:
    try:
        cursor.executemany(insert_sql, rows)
        return [
            RowResult(row_no=row_no, hash_value=hash_value, status="inserted", row_payload=payload)
            for row_no, hash_value, payload in row_contexts
        ], []
    except Exception as exc:
        cursor.connection.rollback()
        results: List[RowResult] = []
        inserted_rows: List[Tuple[str, Tuple[Any, ...]]] = []
        for row_values, (row_no, hash_value, payload) in zip(rows, row_contexts):
            try:
                cursor.execute(insert_sql, row_values)
                cursor.connection.commit()
                inserted_rows.append((hash_value, row_values))
                results.append(RowResult(row_no=row_no, hash_value=hash_value, status="inserted", row_payload=payload))
            except Exception as row_exc:
                cursor.connection.rollback()
                results.append(
                    RowResult(
                        row_no=row_no,
                        hash_value=hash_value,
                        status="failed",
                        message=f"批次失敗後逐筆重試仍失敗：{row_exc}; 原批次錯誤：{exc}",
                        row_payload=payload,
                    )
                )
        return results, inserted_rows


def process_excel_file(
    *,
    config_path: str,
    cfg: Dict[str, Any],
    excel_path: Path,
    conn,
    run_id: str,
    log_dir: Path,
) -> FileResult:
    result = FileResult(file_path=excel_path)
    excel_path, sheet_name, header_row, header_map = open_workbook_header(config_path, cfg, excel_path)
    mappings = read_column_mappings(cfg, header_map)

    table = str(require_nested(cfg, "destination.table"))
    hash_field = str(get_nested(cfg, "destination.hash_field", DEFAULT_HASH_FIELD) or DEFAULT_HASH_FIELD)
    auto_create_hash_field = bool(get_nested(cfg, "destination.auto_create_hash_field", True))
    auto_create_missing_fields = bool(get_nested(cfg, "destination.auto_create_missing_fields", True))
    default_field_type = str(get_nested(cfg, "destination.default_field_type", "NVARCHAR(MAX)") or "NVARCHAR(MAX)")
    hash_field_type = str(
        get_nested(cfg, "destination.hash_field_type", DEFAULT_HASH_FIELD_TYPE)
        or DEFAULT_HASH_FIELD_TYPE
    )
    create_hash_index = bool(get_nested(cfg, "destination.create_hash_index", True))
    dest_fields = [mapping["dest_field"] for mapping in mappings]
    if hash_field not in dest_fields:
        dest_fields.append(hash_field)
    insert_sql = build_insert_sql(table, dest_fields)

    batch_size = int(get_nested(cfg, "import.batch_size", 500) or 500)
    if batch_size <= 0:
        raise ValueError("import.batch_size 必須大於 0")

    start_row = int(get_nested(cfg, "source.start_row", header_row + 1) or (header_row + 1))
    end_row = int(get_nested(cfg, "source.end_row", 0) or 0)
    if start_row <= header_row:
        raise ValueError("source.start_row 必須大於 source.header_row")
    if end_row and end_row < start_row:
        raise ValueError("source.end_row 必須大於或等於 source.start_row，或設為 0 表示讀到最後")

    skip_empty_rows = bool(get_nested(cfg, "import.skip_empty_rows", True))
    fast_executemany = bool(get_nested(cfg, "import.fast_executemany", True))
    audit_enabled = bool(get_nested(cfg, "audit.enabled", True))
    audit_table = str(get_nested(cfg, "audit.table", DEFAULT_AUDIT_TABLE) or DEFAULT_AUDIT_TABLE)

    cursor = conn.cursor()
    if fast_executemany:
        cursor.fast_executemany = True

    if auto_create_missing_fields:
        created_fields = ensure_destination_columns(
            cursor,
            table=table,
            mappings=mappings,
            default_field_type=default_field_type,
        )
        if created_fields:
            conn.commit()
            created_text = ", ".join(f"{field} {field_type}" for field, field_type in created_fields.items())
            print(f"🛠️ 已自動新增目的欄位：{table}；{created_text}")
    else:
        missing_fields = [
            field
            for field in dest_fields
            if field != hash_field and not destination_column_exists(cursor, table=table, field=field)
        ]
        if missing_fields:
            raise ValueError(
                f"目的資料表 {table} 缺少欄位：{', '.join(missing_fields)}；"
                "請手動新增，或在設定檔 destination.auto_create_missing_fields 設為 true。"
            )

    if auto_create_hash_field:
        created_hash_field = ensure_destination_hash_field(
            cursor,
            table=table,
            hash_field=hash_field,
            field_type=hash_field_type,
            create_unique_index=create_hash_index,
        )
        if created_hash_field:
            conn.commit()
            index_message = "，並建立唯一索引" if create_hash_index else ""
            print(
                f"🛠️ 已自動新增去重欄位：{table}.{hash_field} "
                f"{validate_hash_field_type(hash_field_type)}{index_message}"
            )
    elif not destination_hash_field_exists(cursor, table=table, hash_field=hash_field):
        raise ValueError(
            f"目的資料表 {table} 缺少去重欄位 {hash_field}；"
            "請手動新增，或在設定檔 destination.auto_create_hash_field 設為 true。"
        )

    print(f"ℹ️ Excel: {excel_path}")
    print(f"ℹ️ Sheet: {sheet_name}")
    print(f"ℹ️ SQL: {insert_sql}")
    print(
        f"ℹ️ 去重機制：以來源欄名與內容組成字串後計算 {DEFAULT_HASH_ALGORITHM} "
        f"hex digest（固定 {DEFAULT_HASH_HEX_LENGTH} 字元 / 256-bit），寫入 {hash_field}，"
        "匯入前會檢查避免重複導入。"
    )
    if audit_enabled:
        print(
            f"ℹ️ Audit table：{audit_table}（由 audit.table 設定；"
            f"未設定時使用預設值 {DEFAULT_AUDIT_TABLE}）"
        )
    else:
        print("ℹ️ Audit table：已停用（audit.enabled=false）")

    header_audit_results = build_header_audit_results(mappings, header_row)
    if header_audit_results:
        for header_result in header_audit_results:
            print(f"⚠️ 欄名智慧比對：{header_result.message}")
            write_log(
                log_dir,
                {
                    "event": "header_smart_match",
                    "run_id": run_id,
                    "source_file": str(excel_path),
                    "message": header_result.message,
                    "payload": header_result.row_payload,
                },
            )
        if audit_enabled:
            print("ℹ️ 準備寫入欄名智慧比對 audit rows", flush=True)
            insert_audit_rows(
                cursor,
                audit_table=audit_table,
                run_id=run_id,
                source_file=excel_path,
                sheet_name=sheet_name,
                destination_table=table,
                row_results=header_audit_results,
            )
            conn.commit()
            print("ℹ️ 欄名智慧比對 audit rows 寫入完成", flush=True)

    try:
        row_iter = iter_excel_rows(
            excel_path=excel_path,
            sheet_name=sheet_name,
            header_row=header_row,
            start_row=start_row,
            end_row=end_row,
            mappings=mappings,
            skip_empty_rows=skip_empty_rows,
        )

        seen_in_run: set[str] = set()
        for raw_batch in chunked(row_iter, batch_size):
            batch_results: List[RowResult] = []
            candidates: List[Tuple[int, Dict[str, Any], str]] = []

            existing_hashes = get_existing_hashes(
                cursor,
                table=table,
                hash_field=hash_field,
                hash_values=[hash_value for _row_no, _record, _source_values, hash_value, _invalid in raw_batch],
            )

            for row_no, record, source_values, hash_value, invalid_message in raw_batch:
                payload = {"source": source_values, "destination": record}
                if invalid_message:
                    batch_results.append(
                        RowResult(
                            row_no=row_no,
                            hash_value=hash_value,
                            status="failed",
                            message=invalid_message,
                            row_payload=payload,
                        )
                    )
                    continue
                if hash_value in existing_hashes:
                    batch_results.append(
                        RowResult(
                            row_no=row_no,
                            hash_value=hash_value,
                            status="duplicate",
                            message=f"hash_value 已存在於 {table}.{hash_field}，略過匯入",
                            row_payload=payload,
                        )
                    )
                    continue
                if hash_value in seen_in_run:
                    batch_results.append(
                        RowResult(
                            row_no=row_no,
                            hash_value=hash_value,
                            status="duplicate",
                            message="同次執行中已有相同來源欄名與內容的資料列，略過匯入",
                            row_payload=payload,
                        )
                    )
                    continue
                seen_in_run.add(hash_value)
                record[hash_field] = hash_value
                candidates.append((row_no, record, hash_value))

            if candidates:
                rows = [tuple(record.get(field) for field in dest_fields) for _row_no, record, _hash_value in candidates]
                contexts = [
                    (row_no, hash_value, {"destination": record})
                    for row_no, record, hash_value in candidates
                ]
                insert_results, _inserted_rows = insert_batch_with_fallback(
                    cursor,
                    insert_sql=insert_sql,
                    rows=rows,
                    row_contexts=contexts,
                )
                batch_results.extend(insert_results)

            if audit_enabled:
                print(f"ℹ️ 準備寫入批次 audit rows：{len(batch_results)} 筆", flush=True)
                insert_audit_rows(
                    cursor,
                    audit_table=audit_table,
                    run_id=run_id,
                    source_file=excel_path,
                    sheet_name=sheet_name,
                    destination_table=table,
                    row_results=batch_results,
                )
                print("ℹ️ 批次 audit rows 寫入完成", flush=True)

            conn.commit()

            inserted = sum(1 for item in batch_results if item.status == "inserted")
            duplicate = sum(1 for item in batch_results if item.status == "duplicate")
            failed = sum(1 for item in batch_results if item.status == "failed")
            result.success += inserted
            result.duplicate += duplicate
            result.failed += failed

            print(
                f"✅ 批次完成：新增 {inserted} 筆，重複略過 {duplicate} 筆，"
                f"失敗 {failed} 筆，累計新增 {result.success} 筆"
            )
            write_log(
                log_dir,
                {
                    "event": "batch_done",
                    "run_id": run_id,
                    "source_file": str(excel_path),
                    "inserted": inserted,
                    "duplicate": duplicate,
                    "failed": failed,
                },
            )
    finally:
        cursor.close()

    print(
        f"🎯 檔案完成：{excel_path.name}，成功新增 {result.success} 筆，"
        f"重複略過 {result.duplicate} 筆，失敗 {result.failed} 筆"
    )
    return result


def get_log_dir(config_path: str, cfg: Dict[str, Any]) -> Path:
    raw = str(get_nested(cfg, "archive.log_dir", "logs/excel_import") or "logs/excel_import")
    return resolve_path(config_path, raw)


def get_import_retry_config(cfg: Dict[str, Any]) -> Tuple[int, float]:
    max_attempts = int(
        get_nested(
            cfg,
            "retry.max_attempts",
            get_nested(cfg, "import.retry_max_attempts", get_nested(cfg, "import.max_attempts", 1)),
        )
        or 1
    )
    interval_seconds = float(
        get_nested(
            cfg,
            "retry.interval_seconds",
            get_nested(cfg, "import.retry_interval_seconds", get_nested(cfg, "import.retry_delay_seconds", 0)),
        )
        or 0
    )
    if max_attempts <= 0:
        raise ValueError("retry.max_attempts 必須大於 0")
    if interval_seconds < 0:
        raise ValueError("retry.interval_seconds 不可小於 0")
    return max_attempts, interval_seconds


def process_excel_file_with_retries(
    *,
    config_path: str,
    cfg: Dict[str, Any],
    excel_path: Path,
    conn,
    run_id: str,
    log_dir: Path,
    max_attempts: int,
    interval_seconds: float,
) -> FileResult:
    last_result = FileResult(file_path=excel_path)
    for attempt in range(1, max_attempts + 1):
        try:
            if attempt > 1:
                print(f"🔁 第 {attempt}/{max_attempts} 次重試匯入：{excel_path}")
            file_result = process_excel_file(
                config_path=config_path,
                cfg=cfg,
                excel_path=excel_path,
                conn=conn,
                run_id=run_id,
                log_dir=log_dir,
            )
            if file_result.ok or attempt >= max_attempts:
                return file_result

            last_result = file_result
            print(
                f"⚠️ 檔案匯入仍有失敗列：{excel_path}；"
                f"第 {attempt}/{max_attempts} 次，{interval_seconds:g} 秒後重試"
            )
            write_log(
                log_dir,
                {
                    "event": "file_retry_scheduled",
                    "run_id": run_id,
                    "source_file": str(excel_path),
                    "attempt": attempt,
                    "max_attempts": max_attempts,
                    "interval_seconds": interval_seconds,
                    "failed": file_result.failed,
                    "error": file_result.error,
                },
            )
        except Exception as exc:
            conn.rollback()
            last_result = FileResult(file_path=excel_path, failed=1, error=str(exc))
            if attempt >= max_attempts:
                return last_result

            print(
                f"⚠️ 檔案處理失敗：{excel_path}；第 {attempt}/{max_attempts} 次；"
                f"{interval_seconds:g} 秒後重試；錯誤：{exc}"
            )
            write_log(
                log_dir,
                {
                    "event": "file_retry_scheduled",
                    "run_id": run_id,
                    "source_file": str(excel_path),
                    "attempt": attempt,
                    "max_attempts": max_attempts,
                    "interval_seconds": interval_seconds,
                    "error": str(exc),
                },
            )

        if interval_seconds > 0:
            time.sleep(interval_seconds)

    return last_result


def bulk_insert_excel_to_sql_from_config(config_path: str, cfg: Dict[str, Any]) -> tuple[int, int, int]:
    enable_native_crash_tracebacks()
    run_id = uuid.uuid4().hex
    files = discover_excel_files(config_path, cfg)
    log_dir = get_log_dir(config_path, cfg)
    if not files:
        print("ℹ️ 沒有找到可匯入的 Excel 檔案")
        write_log(log_dir, {"event": "run_done", "run_id": run_id, "message": "no files"})
        return 0, 0, 0

    import pyodbc

    conn = pyodbc.connect(build_conn_str_from_config(cfg))
    cursor = conn.cursor()
    try:
        audit_enabled = bool(get_nested(cfg, "audit.enabled", True))
        audit_table = str(get_nested(cfg, "audit.table", DEFAULT_AUDIT_TABLE) or DEFAULT_AUDIT_TABLE)
        if audit_enabled:
            ensure_audit_table(cursor, audit_table)
            conn.commit()

        truncate_before_load = bool(get_nested(cfg, "import.truncate_before_load", False))
        if truncate_before_load:
            table = str(require_nested(cfg, "destination.table"))
            truncate_sql = f"TRUNCATE TABLE {quote_object_path(table)}"
            cursor.execute(truncate_sql)
            conn.commit()
            print(f"⚠️ 已先清空目的資料表：{table}")
    finally:
        cursor.close()

    archive_enabled = bool(get_nested(cfg, "archive.enabled", bool(get_nested(cfg, "source.directory", ""))))
    success_dir = resolve_path(config_path, str(get_nested(cfg, "archive.success_dir", "archive/success") or "archive/success"))
    fail_dir = resolve_path(config_path, str(get_nested(cfg, "archive.fail_dir", "archive/failed") or "archive/failed"))

    retry_max_attempts, retry_interval_seconds = get_import_retry_config(cfg)
    if retry_max_attempts > 1:
        print(
            f"ℹ️ 匯入重試：最多嘗試 {retry_max_attempts} 次；"
            f"每次間隔 {retry_interval_seconds:g} 秒"
        )

    totals = {"success": 0, "duplicate": 0, "failed": 0}
    try:
        for excel_path in files:
            file_result = process_excel_file_with_retries(
                config_path=config_path,
                cfg=cfg,
                excel_path=excel_path,
                conn=conn,
                run_id=run_id,
                log_dir=log_dir,
                max_attempts=retry_max_attempts,
                interval_seconds=retry_interval_seconds,
            )
            if not file_result.ok:
                print(f"❌ 檔案達重試上限後仍失敗：{excel_path}；{file_result.error}")
                write_log(
                    log_dir,
                    {
                        "event": "file_failed",
                        "run_id": run_id,
                        "source_file": str(excel_path),
                        "attempts": retry_max_attempts,
                        "failed": file_result.failed,
                        "error": file_result.error,
                    },
                )

            totals["success"] += file_result.success
            totals["duplicate"] += file_result.duplicate
            totals["failed"] += file_result.failed

            if archive_enabled:
                target_dir = success_dir if file_result.ok else fail_dir
                moved_to = move_to_archive(excel_path, target_dir)
                print(f"📦 已移動檔案到：{moved_to}")
                write_log(
                    log_dir,
                    {
                        "event": "file_archived",
                        "run_id": run_id,
                        "source_file": str(excel_path),
                        "target_file": str(moved_to),
                        "status": "success" if file_result.ok else "failed",
                        "inserted": file_result.success,
                        "duplicate": file_result.duplicate,
                        "failed": file_result.failed,
                        "error": file_result.error,
                    },
                )
    finally:
        conn.close()

    print(
        f"🎯 全部完成：成功新增 {totals['success']} 筆，"
        f"重複略過 {totals['duplicate']} 筆，失敗 {totals['failed']} 筆"
    )
    write_log(log_dir, {"event": "run_done", "run_id": run_id, **totals})
    return totals["success"], totals["duplicate"], totals["failed"]


def main():
    enable_native_crash_tracebacks()

    parser = argparse.ArgumentParser(description="依設定檔從 Excel 或來源目錄匯入資料到 SQL Server")
    parser.add_argument("--config", required=True, help="JSON 設定檔路徑")
    args = parser.parse_args()

    from dotenv import load_dotenv

    load_dotenv()

    cfg = load_json_config(args.config)
    bulk_insert_excel_to_sql_from_config(args.config, cfg)


if __name__ == "__main__":
    main()
