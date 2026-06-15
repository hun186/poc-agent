"""Excel workbook discovery, header mapping, and row normalization."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

from .constants import (
    EXCEL_SUFFIXES,
)
from .utils import (
    as_list,
    get_nested,
    normalize_header,
    normalize_header_key,
    require_nested,
    resolve_excel_path,
    resolve_path,
)


@dataclass(frozen=True)
class SheetHeader:
    excel_path: Path
    sheet_name: str
    header_row: int
    header_map: Dict[str, int]
    match_type: str = "exact"
    match_message: str = ""


from .workbook_rows import (
    iter_excel_rows,
    make_hash_value,
    parse_excel_column_ref,
    read_column_mappings,
    read_header_map,
    resolve_header_candidate,
)


def _common_prefix_length(left: str, right: str) -> int:
    count = 0
    for left_char, right_char in zip(left, right):
        if left_char != right_char:
            break
        count += 1
    return count


def _common_prefix_is_meaningful(left: str, right: str) -> bool:
    prefix_len = _common_prefix_length(left, right)
    shortest = min(len(left), len(right))
    if shortest <= 0:
        return False
    return prefix_len >= min(4, shortest) and prefix_len >= max(2, shortest // 2)


def _get_expected_sheet_names(cfg: Dict[str, Any]) -> List[str]:
    raw_sheets = get_nested(cfg, "source.sheets", None)
    if raw_sheets is None:
        raw_sheets = get_nested(cfg, "source.sheet_names", None)
    if raw_sheets is None:
        raw_sheets = require_nested(cfg, "source.sheet")

    names = as_list(raw_sheets)
    names.extend(as_list(get_nested(cfg, "source.sheet_aliases", None)))
    if not names:
        raise ValueError(
            "設定檔缺少必要欄位：source.sheet（或 source.sheets 至少一筆工作表名稱）"
        )
    return names


def resolve_sheet_candidates(
    *,
    available_sheet_names: Sequence[str],
    expected_names: Sequence[str],
    match_mode: str,
) -> List[Tuple[str, str, str, str]]:
    candidates_to_try = [
        normalize_header(name) for name in expected_names if normalize_header(name)
    ]
    if not candidates_to_try:
        raise ValueError("source.sheet 缺少可比對的工作表名稱")

    mode = (match_mode or "smart").strip().lower()
    if mode not in {"exact", "smart", "normalized", "prefix", "contains"}:
        raise ValueError(f"source.sheet_match 不支援：{match_mode}")

    matches_by_sheet: Dict[str, Tuple[str, str, str, str]] = {}

    def add(sheet: str, expected: str, match_type: str, message: str = ""):
        if sheet not in matches_by_sheet:
            matches_by_sheet[sheet] = (sheet, expected, match_type, message)

    for expected in candidates_to_try:
        for sheet in available_sheet_names:
            if sheet == expected:
                add(sheet, expected, "exact", "")

    if mode == "exact":
        ordered_exact = [
            matches_by_sheet[sheet]
            for sheet in available_sheet_names
            if sheet in matches_by_sheet
        ]
        if ordered_exact:
            return ordered_exact
        raise ValueError(
            f"Excel 找不到工作表：{', '.join(candidates_to_try)}；"
            f"可用工作表：{', '.join(available_sheet_names)}"
        )

    normalized_sheets: Dict[str, List[str]] = {}
    for sheet in available_sheet_names:
        normalized_sheets.setdefault(normalize_header_key(sheet), []).append(sheet)

    if mode in {"smart", "normalized", "prefix", "contains"}:
        for expected in candidates_to_try:
            expected_key = normalize_header_key(expected)
            if not expected_key:
                continue
            for sheet in normalized_sheets.get(expected_key, []):
                add(
                    sheet,
                    expected,
                    "normalized",
                    f"設定工作表 {expected} 以正規化名稱比對到 Excel 工作表 {sheet}",
                )

    if mode in {"smart", "prefix", "contains"}:
        for expected in candidates_to_try:
            expected_key = normalize_header_key(expected)
            if not expected_key:
                continue
            for sheet in available_sheet_names:
                sheet_key = normalize_header_key(sheet)
                if (
                    sheet_key.startswith(expected_key)
                    or expected_key.startswith(sheet_key)
                    or (
                        mode == "smart"
                        and _common_prefix_is_meaningful(sheet_key, expected_key)
                    )
                ):
                    add(
                        sheet,
                        expected,
                        "prefix",
                        f"設定工作表 {expected} 以前綴智慧比對到 Excel 工作表 {sheet}",
                    )

    if mode == "contains":
        for expected in candidates_to_try:
            expected_key = normalize_header_key(expected)
            if not expected_key:
                continue
            for sheet in available_sheet_names:
                sheet_key = normalize_header_key(sheet)
                if expected_key in sheet_key or sheet_key in expected_key:
                    add(
                        sheet,
                        expected,
                        "contains",
                        f"設定工作表 {expected} 以包含關係比對到 Excel 工作表 {sheet}",
                    )

    ordered_matches = [
        matches_by_sheet[sheet]
        for sheet in available_sheet_names
        if sheet in matches_by_sheet
    ]
    if ordered_matches:
        return ordered_matches

    raise ValueError(
        f"Excel 找不到工作表：{', '.join(candidates_to_try)}；"
        f"可用工作表：{', '.join(available_sheet_names)}"
    )


def open_workbook_headers(
    config_path: str,
    cfg: Dict[str, Any],
    excel_path: Optional[Path] = None,
) -> List[SheetHeader]:
    if excel_path is None:
        excel_path_raw = require_nested(cfg, "source.excel_path")
        excel_path = resolve_excel_path(config_path, str(excel_path_raw))
    if not excel_path.exists():
        raise FileNotFoundError(f"找不到 Excel 檔案：{excel_path}")

    expected_sheet_names = _get_expected_sheet_names(cfg)
    sheet_match_mode = str(get_nested(cfg, "source.sheet_match", "smart") or "smart")
    header_row = int(get_nested(cfg, "source.header_row", 1) or 1)
    if header_row <= 0:
        raise ValueError("source.header_row 必須大於 0")

    from openpyxl import load_workbook

    wb = load_workbook(filename=excel_path, read_only=True, data_only=True)
    try:
        sheet_matches = resolve_sheet_candidates(
            available_sheet_names=wb.sheetnames,
            expected_names=expected_sheet_names,
            match_mode=sheet_match_mode,
        )
        return [
            SheetHeader(
                excel_path=excel_path,
                sheet_name=sheet_name,
                header_row=header_row,
                header_map=read_header_map(wb[sheet_name], header_row),
                match_type=match_type,
                match_message=match_message,
            )
            for sheet_name, _expected, match_type, match_message in sheet_matches
        ]
    finally:
        wb.close()


def open_workbook_header(
    config_path: str,
    cfg: Dict[str, Any],
    excel_path: Optional[Path] = None,
) -> Tuple[Path, str, int, Dict[str, int]]:
    sheet_headers = open_workbook_headers(config_path, cfg, excel_path)
    first = sheet_headers[0]
    return first.excel_path, first.sheet_name, first.header_row, first.header_map


def discover_excel_files(config_path: str, cfg: Dict[str, Any]) -> List[Path]:
    source_dir_raw = (
        get_nested(cfg, "source.directory", "")
        or get_nested(cfg, "source.excel_dir", "")
        or ""
    )
    if source_dir_raw:
        source_dir = resolve_path(config_path, str(source_dir_raw))
        if not source_dir.exists():
            raise FileNotFoundError(f"找不到來源目錄：{source_dir}")
        if not source_dir.is_dir():
            raise NotADirectoryError(f"來源目錄不是資料夾：{source_dir}")

        recursive = bool(get_nested(cfg, "source.recursive", False))
        pattern = str(get_nested(cfg, "source.file_pattern", "*.xlsx") or "*.xlsx")
        paths = source_dir.rglob(pattern) if recursive else source_dir.glob(pattern)
        files = [
            path.resolve()
            for path in paths
            if path.is_file()
            and path.suffix.lower() in EXCEL_SUFFIXES
            and not path.name.startswith("~$")
        ]
        return sorted(files)

    excel_path_raw = require_nested(cfg, "source.excel_path")
    return [resolve_excel_path(config_path, str(excel_path_raw))]
