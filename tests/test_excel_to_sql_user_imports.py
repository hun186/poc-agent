from __future__ import annotations

from io import BytesIO
import json
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from data_ingestion.excel_to_sql import user_imports
from data_ingestion.excel_to_sql import user_import_storage


def _patch_store(monkeypatch, tmp_path: Path) -> None:
    root = tmp_path / "imports"
    monkeypatch.setattr(user_import_storage, "ROOT", root)
    monkeypatch.setattr(user_import_storage, "ACTIVE_DIR", root / "active")
    monkeypatch.setattr(user_import_storage, "RECORDS_DIR", root / "records")
    monkeypatch.setattr(user_import_storage, "REGISTRY", root / "records.json")


def test_save_and_list_upload_record_uses_compact_registry(monkeypatch, tmp_path):
    _patch_store(monkeypatch, tmp_path)
    config = tmp_path / "cfg.json"
    config.write_text(
        '{"source":{"sheet":"Sheet1"},"destination":{"table":"dbo.t"},"mappings":[{"source_column":"A","dest_field":"a"}],"db":{"driver":"x","server":"s","database":"d","user":"u","password":"p"}}',
        encoding="utf-8",
    )

    record = user_imports.save_upload(
        fileobj=BytesIO(b"excel-bytes"),
        filename="sample.xlsx",
        username="alice",
        config_path=str(config),
    )

    assert record["original_filename"] == "sample.xlsx"
    assert record["stored_filename"] == f"{record['id']}.xlsx"
    assert record["status"] == "uploaded"
    assert "stored_path" not in record
    assert (
        user_imports.list_import_records(username="alice", is_superuser=False)[0]["id"]
        == record["id"]
    )
    assert (
        user_imports.list_import_records(
            username="alice",
            is_superuser=False,
            filename=record["stored_filename"],
        )[0]["id"]
        == record["id"]
    )
    assert user_imports.list_import_records(username="bob", is_superuser=False) == []

    registry = (tmp_path / "imports" / "records.json").read_text(encoding="utf-8")
    assert '"record_ids"' in registry
    assert '"original_filename"' not in registry
    record_path = tmp_path / "imports" / "records" / f"{record['id']}.json"
    assert record_path.exists()
    assert '"original_filename"' in record_path.read_text(encoding="utf-8")


def test_legacy_records_json_migrates_to_per_record_files(monkeypatch, tmp_path):
    _patch_store(monkeypatch, tmp_path)
    root = tmp_path / "imports"
    root.mkdir(parents=True)
    legacy = {
        "records": [
            {
                "id": "legacy-1",
                "original_filename": "legacy.xlsx",
                "stored_path": str(root / "active" / "legacy-1.xlsx"),
                "config_path": "cfg.json",
                "username": "alice",
                "created_at": "2026-06-15T00:00:00+00:00",
                "status": "uploaded",
            }
        ]
    }
    (root / "records.json").write_text(json.dumps(legacy), encoding="utf-8")

    data = user_import_storage.read_registry()

    assert data["records"][0]["id"] == "legacy-1"
    compact = (root / "records.json").read_text(encoding="utf-8")
    assert '"record_ids"' in compact
    assert '"records"' not in compact
    assert (root / "records" / "legacy-1.json").exists()


def test_import_center_static_contains_user_import_controls():
    dashboard_html = Path("sra_api/static/excel_to_sql_dashboard.html").read_text(
        encoding="utf-8"
    )
    html = Path("sra_api/static/excel_to_sql_import_center.html").read_text(
        encoding="utf-8"
    )
    script = Path("sra_api/static/excel_to_sql_import_records.js").read_text(
        encoding="utf-8"
    )
    assert (
        "使用者 Excel 上傳、預掃描、確認匯入 SQL 與匯入記錄管理已移至獨立頁面"
        not in dashboard_html
    )
    assert "前往 Excel 匯入作業中心" not in dashboard_html
    assert "使用者 Excel 匯入操作" not in dashboard_html
    assert "使用者 Excel 匯入操作" in html
    assert (
        "/static/excel_to_sql_import_records.js?v=20260629-task-review-permission"
        in html
    )
    assert "/v1/excel-to-sql/import-records" in script
    assert "previewImportRecord" in script
    assert "刪除記錄" not in script
    assert "deleteSelectedImportRecords" not in script
    assert "deleteSingleImportRecord" not in script
    assert "recordSelectAll" not in html
    assert "批次刪除勾選記錄" not in html
    assert "錯誤診斷" in html
    assert "匯入任務" in html
    assert "uploadTaskDetails" in html
    assert "目前匯入任務的 Excel→SQL 欄位轉換" in script
    assert "dashboardTasks" in script
    assert "renderUploadTaskDetails" in script
    assert "canReviewTaskMappings" in script
    assert "ROLE_DB_OPERATOR" in script
    config_html = Path("sra_api/static/excel_to_sql_task_configs.html").read_text(
        encoding="utf-8"
    )
    assert "建立 / 管理待審任務 JSON" not in html
    assert "建立 Excel→SQL 任務設定檔" in config_html
    assert "我建立的任務清單 / 審核進度" in config_html
    assert "資料庫連線來源任務" in config_html
    assert "department/new_task.json" in config_html
    assert "taskMappings" in config_html
    assert "sqlTypeOptions" in config_html
    assert 'id="taskDefaultFieldType" list="sqlTypeOptions"' in config_html
    assert (
        'id="taskDefaultFieldType" list="sqlTypeOptions" value="NVARCHAR(MAX)"'
        not in config_html
    )
    assert "常用 SQL 型別小提醒" in config_html
    assert "預設欄位型別與欄位對應 SQL 型別都可選擇或手動輸入" in config_html
    assert "若不確定欄位長度，建議保留預設 NVARCHAR(MAX)" in config_html
    assert "最多 255 個字元的短文字" in config_html
    assert "長文字建議 NVARCHAR(MAX)，短代碼才用 NVARCHAR(255)" in config_html
    assert "DECIMAL(18,2)" in config_html
    assert "DATE / DATETIME2" in config_html
    assert "BIT" in config_html
    assert "比對方式說明" in config_html
    assert "多欄位拼裝成 SQL 新欄位" in config_html
    assert "taskCompositeMappings" in config_html
    assert "addCompositeMappingRow" in script
    assert "collectCompositeMappings" in script
    assert "composite_mappings" in script
    assert "composite-template" in script
    assert "日期是{日期}，發生地點是{地點}，當天的敘述是{敘述}" in config_html
    assert "complete_description" in config_html + script
    assert "實體欄位模式" in config_html
    assert "會多佔用資料庫空間" in config_html
    assert "SQL View / computed expression" in config_html
    assert "尚未開放真正建 View" in config_html
    assert "拼裝來源欄位不再另外選比對方式" in config_html
    assert "相同來源欄位設定的 exact / prefix / contains 逐欄套用" in config_html
    assert "拼裝來源會沿用上方相同來源欄位的比對方式" in config_html
    assert "預掃描或確認匯入" in config_html
    assert "composite-match" not in script
    assert "item.match !== 'exact'" not in script
    assert "每個拼裝欄位都需要至少兩個來源欄位" in script
    assert "長文字建議 NVARCHAR(MAX)，短代碼才用 NVARCHAR(255)" in script
    assert "pending" in config_html
    assert "superuserTaskReview" in config_html
    assert "approveTaskConfigPath" in config_html
    assert "taskReviewEnabled" in config_html
    assert "刪除任務設定檔" in config_html
    assert "deleteReviewedTaskConfig" in script
    assert "taskReviewDetails" in config_html
    assert "review-details" in config_html
    assert "JSON 任務轉換細節" in script
    assert "Server、資料庫名稱、DB 使用者與密碼等敏感欄位已遮蔽" in script
    assert "connectionSourceNotice" in config_html
    assert "setSensitiveConnectionFieldsVisible" in script
    assert "input.disabled = !visible" in script
    assert "function renderTaskReviewDetails" in script
    assert "config_preview" in script
    assert "review_summary" in script
    assert "config-grid" in config_html
    assert "config-group" in config_html
    assert "基本資訊" in config_html
    assert "資料庫連線" in config_html
    assert "Excel 來源" in config_html
    assert "來源目錄是批次匯入時伺服器掃描 Excel 的資料夾" in config_html
    assert "SQL 目的地" in config_html
    assert "/v1/excel-to-sql/task-configs" in script
    assert "connection_source_config" in script
    assert "taskConnectionSource" in script
    assert "function createTaskConfig" in script
    assert "function applyConnectionSourceDefaults" in script
    assert "function loadTaskConfigSourceCandidates" in script
    assert "mergeTasksByConfigPath" in script
    assert "isActiveConnectionSourceTask" in script
    assert ".filter(isActiveConnectionSourceTask)" in script
    assert "connectionSourceTasks" in script
    assert "taskDbServer" in script
    assert "taskDbDatabase" in script
    assert "taskDbUser" in script
    assert "taskDbDriver" in script
    assert "taskTrustedConnection" in script
    assert "function setTaskConfigStatus" in script
    assert "function loadTaskConfigReviewOptions" in script
    assert "currentUserRole !== ROLE_SUPERUSER" in script
    assert "function approveTaskConfig" not in script
    assert "showMatchHelp" in script
    assert 'list="sqlTypeOptions"' in script
    assert "監控任務" not in html
    assert "Excel→SQL 匯入查詢儀表板" in html
    assert "任務 JSON 管理" in html
    assert "任務設定檔管理" not in html
    assert "船舶活動 Excel 匯入任務" in Path(
        "config/excel_to_sql/vessel_activities.example.json"
    ).read_text(encoding="utf-8")
    assert "原始檔名" in html
    assert "導入檔名" in html
    assert 'class="table-wrap import-records-wrap"' in html
    assert 'class="import-records-table"' in html
    assert "table-layout:fixed" in html
    assert "#importRecords .stored-file-col { width:9%; }" in html
    assert "#importRecords .actions-col { width:10%; }" in html
    assert (
        ".record-actions { display:grid; grid-template-columns: minmax(0, 1fr);" in html
    )
    assert "min-width:128px" in html
    assert ".record-actions button:nth-child(2)" not in html
    assert "repeat(3, minmax(0, 1fr))" not in html
    assert "@media (max-width: 760px)" in html
    assert "content:attr(data-label)" in html
    assert "function importRecordDiagnostics" in script
    assert "<details><summary>展開 / 收合錯誤診斷</summary>" in script
    assert "function openAuditForImportRecord" in script
    assert "source_file_filter" in script
    assert "#recentAuditSection" in script
    assert "params.set('days'" not in script
    assert "查詢已匯入條目" not in script
    assert (
        script.index("預掃描")
        < script.index("確認匯入SQL")
        < script.index("查詢已匯入列")
    )
    assert "查詢已匯入列" in script
    assert "window.__importRecordById" in script
    assert 'data-label="導入檔名"' in script
    assert 'data-label="操作"' in script
    assert 'aria-label="選取匯入記錄' not in script
    assets = Path("sra_api/routers/editor.py").read_text(encoding="utf-8")
    router = Path("sra_api/routers/excel_to_sql.py").read_text(encoding="utf-8")
    facade = Path("data_ingestion/excel_to_sql/user_imports.py").read_text(
        encoding="utf-8"
    )
    storage = Path("data_ingestion/excel_to_sql/user_import_storage.py").read_text(
        encoding="utf-8"
    )
    assert "excel_to_sql_import_records.js" in assets
    assert "excel_to_sql_dashboard_app.js" in assets
    assert '@router.delete("/v1/excel-to-sql/import-records")' not in router
    assert "def delete_records" not in storage
    assert "delete_records" not in facade


def test_preview_auto_creates_missing_destination_table(monkeypatch, tmp_path):
    from types import SimpleNamespace
    import sys
    from data_ingestion.excel_to_sql import user_import_preview

    class FakeCursor:
        def __init__(self):
            self.statements = []
            self.closed = False

        def execute(self, sql, *params):
            self.statements.append(str(sql))
            return self

        def fetchone(self):
            # destination_table_exists returns false for the missing table.
            return None

        def close(self):
            self.closed = True

    class FakeConn:
        def __init__(self):
            self.cursor_obj = FakeCursor()
            self.commits = 0
            self.closed = False

        def cursor(self):
            return self.cursor_obj

        def commit(self):
            self.commits += 1

        def close(self):
            self.closed = True

    fake_conn = FakeConn()
    monkeypatch.setitem(
        sys.modules,
        "pyodbc",
        SimpleNamespace(connect=lambda conn_str: fake_conn),
    )
    monkeypatch.setattr(
        user_import_preview,
        "open_workbook_headers",
        lambda config_path, cfg, excel_path: [
            SimpleNamespace(sheet_name="Sheet1", header_row=1, header_map={"A": "名稱"})
        ],
    )
    monkeypatch.setattr(
        user_import_preview,
        "read_column_mappings",
        lambda cfg, header_map: [{"source_column": "名稱", "dest_field": "name"}],
    )
    monkeypatch.setattr(
        user_import_preview,
        "iter_excel_rows",
        lambda **kwargs: iter([(2, {"name": "A"}, {"名稱": "A"}, "hash-a", "")]),
    )
    monkeypatch.setattr(
        user_import_preview, "get_existing_hashes", lambda *args, **kwargs: set()
    )

    result = user_import_preview._preview_file(
        "cfg.json",
        {
            "destination": {
                "table": "dbo.vessel_activities",
                "auto_create_table": True,
            },
            "db": {
                "driver": "x",
                "server": "s",
                "database": "d",
                "user": "u",
                "password": "p",
            },
            "source": {},
        },
        tmp_path / "sample.xlsx",
    )

    assert result["will_insert"] == 1
    assert any(
        "CREATE TABLE [dbo].[vessel_activities]" in sql
        for sql in fake_conn.cursor_obj.statements
    )
    assert fake_conn.commits == 1


def test_confirm_import_ensures_audit_table_before_processing(monkeypatch, tmp_path):
    from types import SimpleNamespace
    import sys
    from data_ingestion.excel_to_sql import user_import_runner

    _patch_store(monkeypatch, tmp_path)
    config = tmp_path / "cfg.json"
    config.write_text(
        '{"source":{"sheet":"Sheet1"},"destination":{"table":"dbo.t"},"audit":{"enabled":true,"table":"dbo.ExcelImportAudit"},"mappings":[{"source_column":"A","dest_field":"a"}],"db":{"driver":"x","server":"s","database":"d","user":"u","password":"p"}}',
        encoding="utf-8",
    )
    record = user_imports.save_upload(
        fileobj=BytesIO(b"excel-bytes"),
        filename="sample.xlsx",
        username="alice",
        config_path=str(config),
    )

    class FakeCursor:
        def __init__(self):
            self.statements = []
            self.closed = False

        def execute(self, sql, *params):
            self.statements.append(str(sql))
            return self

        def close(self):
            self.closed = True

    class FakeConn:
        def __init__(self):
            self.cursor_obj = FakeCursor()
            self.commits = 0
            self.closed = False

        def cursor(self):
            return self.cursor_obj

        def commit(self):
            self.commits += 1

        def close(self):
            self.closed = True

    fake_conn = FakeConn()
    captured_process_kwargs = {}

    def fake_process_excel_file(**kwargs):
        captured_process_kwargs.update(kwargs)
        return SimpleNamespace(
            success=1, duplicate=0, failed=0, ok=True, failure_summary={}
        )

    monkeypatch.setitem(
        sys.modules, "pyodbc", SimpleNamespace(connect=lambda conn_str: fake_conn)
    )
    monkeypatch.setattr(
        user_import_runner, "process_excel_file", fake_process_excel_file
    )

    result = user_import_runner.import_record(
        record["id"], username="alice", is_superuser=False
    )

    assert result["import_result"]["inserted_rows"] == 1
    assert captured_process_kwargs["cfg"]["_audit_actor"] == "alice"
    assert fake_conn.closed is True
    assert fake_conn.cursor_obj.closed is True
    assert any(
        "CREATE TABLE [dbo].[ExcelImportAudit]" in sql
        for sql in fake_conn.cursor_obj.statements
    )
    assert fake_conn.commits >= 1


def test_create_excel_to_sql_task_config_writes_enabled_dashboard_json(tmp_path):
    from data_ingestion.excel_to_sql.task_configs import create_task_config

    payload = {
        "filename": "new_task.json",
        "task_name": "新 Excel 匯入任務",
        "db_server": "sql.local",
        "db_database": "OpsDb",
        "db_user": "operator",
        "db_password": "secret-value",
        "source_sheet": "Sheet1",
        "destination_table": "dbo.ImportTarget",
        "mappings": [
            {
                "source_column": "日期",
                "dest_field": "event_date",
                "dest_field_type": "DATE",
                "required": True,
                "value_type": "date",
            },
            {
                "source_column": "敘述",
                "dest_field": "description",
                "dest_field_type": "NVARCHAR(MAX)",
                "match": "prefix",
            },
        ],
    }

    created = create_task_config(payload, actor="operator", root=tmp_path)

    config_path = Path(created["config_path"])
    data = json.loads(config_path.read_text(encoding="utf-8"))
    assert data["dashboard"]["enabled"] is False
    assert data["dashboard"]["status"] == "pending"
    assert data["dashboard"]["name"] == "新 Excel 匯入任務"
    assert data["db"]["server"] == "sql.local"
    assert data["destination"]["table"] == "dbo.ImportTarget"
    assert data["mappings"][0]["dest_field"] == "event_date"


def test_create_excel_to_sql_task_config_supports_subdirectory_and_copies_connection(
    tmp_path,
):
    from data_ingestion.excel_to_sql.task_configs import create_task_config

    source_dir = tmp_path / "shared"
    source_dir.mkdir(parents=True)
    source = source_dir / "base.json"
    source.write_text(
        json.dumps(
            {
                "db": {
                    "driver": "ODBC Driver 18 for SQL Server",
                    "server": "secret-sql.local",
                    "database": "SecretDb",
                    "trusted_connection": False,
                    "user": "secret-user",
                    "password": "secret-password",
                    "encrypt": False,
                    "trust_server_certificate": False,
                },
                "dashboard": {"enabled": True, "name": "Base"},
            }
        ),
        encoding="utf-8",
    )
    payload = {
        "filename": "business/new_task.json",
        "task_name": "業務自建任務",
        "connection_source_config": "shared/base.json",
        "source_sheet": "Sheet1",
        "destination_table": "dbo.ImportTarget",
        "mappings": [{"source_column": "A", "dest_field": "a"}],
    }

    created = create_task_config(payload, actor="operator", root=tmp_path)

    assert Path(created["config_path"]).name == "new_task.json"
    config_path = tmp_path / "business" / "new_task.json"
    data = json.loads(config_path.read_text(encoding="utf-8"))
    assert data["db"] == json.loads(source.read_text(encoding="utf-8"))["db"]
    assert data["dashboard"]["enabled"] is False
    assert data["dashboard"]["status"] == "pending"
    assert data["dashboard"]["name"] == "業務自建任務"


def test_create_excel_to_sql_task_config_direct_connection_stores_user_and_password(
    tmp_path,
):
    from data_ingestion.excel_to_sql.task_configs import create_task_config

    payload = {
        "filename": "direct.json",
        "task_name": "Direct",
        "db_server": "sql.local",
        "db_database": "OpsDb",
        "db_user": "operator",
        "db_password": "secret-value",
        "db_encrypt": False,
        "db_trust_server_certificate": False,
        "source_sheet": "Sheet1",
        "destination_table": "dbo.ImportTarget",
        "mappings": [{"source_column": "A", "dest_field": "a"}],
    }

    created = create_task_config(payload, actor="operator", root=tmp_path)

    data = json.loads(Path(created["config_path"]).read_text(encoding="utf-8"))
    assert data["db"]["user"] == "operator"
    assert data["db"]["password"] == "secret-value"
    assert data["db"]["encrypt"] is False
    assert data["db"]["trust_server_certificate"] is False


def test_create_excel_to_sql_task_config_accepts_root_prefixed_subpath(tmp_path):
    from data_ingestion.excel_to_sql.task_configs import create_task_config

    payload = {
        "filename": "config/excel_to_sql/group/root_prefixed.json",
        "task_name": "Root prefixed",
        "db_server": "sql.local",
        "db_database": "OpsDb",
        "db_user": "operator",
        "db_password": "secret-value",
        "source_sheet": "Sheet1",
        "destination_table": "dbo.ImportTarget",
        "mappings": [{"source_column": "A", "dest_field": "a"}],
    }

    create_task_config(payload, actor="operator", root=tmp_path)

    assert (tmp_path / "group" / "root_prefixed.json").exists()


def test_list_task_configs_treats_enabled_dashboard_as_active(tmp_path):
    from data_ingestion.excel_to_sql.task_configs import list_task_configs

    config = {
        "dashboard": {"enabled": True, "status": "pending", "name": "SSH active task"},
        "source": {"sheet": "Sheet1"},
        "destination": {"table": "dbo.ImportTarget"},
        "mappings": [{"source_column": "A", "dest_field": "a"}],
        "db": {
            "driver": "ODBC Driver 18 for SQL Server",
            "server": "sql.local",
            "database": "OpsDb",
            "trusted_connection": True,
        },
    }
    (tmp_path / "ssh_active.json").write_text(
        json.dumps(config, ensure_ascii=False), encoding="utf-8"
    )

    listed = list_task_configs(root=tmp_path)

    assert listed[0]["enabled"] is True
    assert listed[0]["status"] == "active"


def test_task_config_review_summary_masks_password_and_status_toggles(tmp_path):
    from data_ingestion.excel_to_sql.task_configs import (
        approve_task_config,
        create_task_config,
        list_task_configs,
        set_task_config_enabled,
    )

    payload = {
        "filename": "pending.json",
        "task_name": "Pending",
        "db_server": "sql.local",
        "db_database": "OpsDb",
        "db_user": "operator",
        "db_password": "secret-value",
        "source_sheet": "Sheet1",
        "destination_table": "dbo.ImportTarget",
        "mappings": [{"source_column": "A", "dest_field": "a"}],
    }
    create_task_config(payload, actor="operator", root=tmp_path)

    listed = list_task_configs(root=tmp_path)
    assert listed[0]["review_summary"]["source"]["sheet"] == "Sheet1"
    assert listed[0]["review_summary"]["destination"]["table"] == "dbo.ImportTarget"
    assert listed[0]["review_summary"]["mappings"][0]["source_column"] == "A"
    assert listed[0]["review_summary"]["mappings"][0]["dest_field"] == "a"
    assert listed[0]["config_preview"]["db"]["password"] == "********"
    assert listed[0]["config_preview"]["db"]["user"] == "********"
    assert listed[0]["config_preview"]["db"]["server"] == "********"
    assert listed[0]["config_preview"]["db"]["database"] == "********"
    assert listed[0]["review_summary"]["db"]["driver"] == "auto"
    assert listed[0]["review_summary"]["db"]["trusted_connection"] is False
    assert listed[0]["review_summary"]["db"]["server"] == "********"
    assert listed[0]["review_summary"]["db"]["database"] == "********"
    assert listed[0]["review_summary"]["db"]["user"] == "********"

    approved = approve_task_config("pending.json", actor="admin", root=tmp_path)

    data = json.loads((tmp_path / "pending.json").read_text(encoding="utf-8"))
    assert approved["enabled"] is True
    assert approved["status"] == "active"
    assert data["dashboard"]["enabled"] is True
    assert data["dashboard"]["status"] == "active"
    assert data["dashboard"]["approved_by"] == "admin"

    disabled = set_task_config_enabled(
        "pending.json", enabled=False, actor="admin", root=tmp_path
    )
    data = json.loads((tmp_path / "pending.json").read_text(encoding="utf-8"))
    assert disabled["enabled"] is False
    assert disabled["status"] == "pending"
    assert data["dashboard"]["enabled"] is False
    assert data["dashboard"]["status"] == "pending"
    assert data["dashboard"]["disabled_by"] == "admin"


def test_excel_to_sql_composite_mapping_combines_multiple_excel_columns(tmp_path):
    from openpyxl import Workbook
    from data_ingestion.excel_to_sql.task_configs import create_task_config
    from data_ingestion.excel_to_sql.workbook import (
        iter_excel_rows,
        open_workbook_headers,
        read_column_mappings,
    )

    workbook = Workbook()
    ws = workbook.active
    ws.title = "Sheet1"
    ws.append(["縣市", "地址", "姓名"])
    ws.append(["台北市", "中正路 1 號", "王小明"])
    excel_path = tmp_path / "sample.xlsx"
    workbook.save(excel_path)

    create_task_config(
        {
            "filename": "composite.json",
            "task_name": "Composite",
            "db_server": "sql.local",
            "db_database": "OpsDb",
            "db_user": "operator",
            "db_password": "secret-value",
            "source_sheet": "Sheet1",
            "destination_table": "dbo.ImportTarget",
            "mappings": [{"source_column": "姓名", "dest_field": "name"}],
            "composite_mappings": [
                {
                    "source_columns": ["縣市", "地址"],
                    "dest_field": "full_address",
                    "separator": " ",
                    "dest_field_type": "NVARCHAR(255)",
                }
            ],
        },
        actor="operator",
        root=tmp_path,
    )
    cfg_path = tmp_path / "composite.json"
    cfg = json.loads(cfg_path.read_text(encoding="utf-8"))
    headers = open_workbook_headers(str(cfg_path), cfg, excel_path)[0]
    mappings = read_column_mappings(cfg, headers.header_map)

    assert cfg["mappings"][1]["compose"]["source_columns"] == ["縣市", "地址"]
    assert mappings[1]["dest_field"] == "full_address"
    rows = list(
        iter_excel_rows(
            excel_path=excel_path,
            sheet_name="Sheet1",
            header_row=1,
            start_row=2,
            end_row=0,
            mappings=mappings,
            skip_empty_rows=True,
        )
    )

    assert rows[0][1]["full_address"] == "台北市 中正路 1 號"
    assert rows[0][1]["name"] == "王小明"


def test_excel_to_sql_composite_mapping_template_substitutes_excel_values(tmp_path):
    from openpyxl import Workbook
    from data_ingestion.excel_to_sql.task_configs import create_task_config
    from data_ingestion.excel_to_sql.workbook import (
        iter_excel_rows,
        open_workbook_headers,
        read_column_mappings,
    )

    workbook = Workbook()
    ws = workbook.active
    ws.title = "Sheet1"
    ws.append(["日期", "地點", "敘述"])
    ws.append(["2026-06-26", "台北", "例行巡檢完成"])
    excel_path = tmp_path / "sample-template.xlsx"
    workbook.save(excel_path)

    create_task_config(
        {
            "filename": "composite-template.json",
            "task_name": "Composite Template",
            "db_server": "sql.local",
            "db_database": "OpsDb",
            "db_user": "operator",
            "db_password": "secret-value",
            "source_sheet": "Sheet1",
            "destination_table": "dbo.ImportTarget",
            "mappings": [{"source_column": "日期", "dest_field": "event_date"}],
            "composite_mappings": [
                {
                    "source_columns": ["日期", "地點", "敘述"],
                    "dest_field": "complete_description",
                    "template": "日期是{日期}，發生地點是{地點}，當天的敘述是{敘述}",
                    "separator": " ",
                    "dest_field_type": "NVARCHAR(MAX)",
                }
            ],
        },
        actor="operator",
        root=tmp_path,
    )
    cfg_path = tmp_path / "composite-template.json"
    cfg = json.loads(cfg_path.read_text(encoding="utf-8"))
    headers = open_workbook_headers(str(cfg_path), cfg, excel_path)[0]
    mappings = read_column_mappings(cfg, headers.header_map)
    rows = list(
        iter_excel_rows(
            excel_path=excel_path,
            sheet_name="Sheet1",
            header_row=1,
            start_row=2,
            end_row=0,
            mappings=mappings,
            skip_empty_rows=True,
        )
    )

    assert (
        cfg["mappings"][1]["compose"]["template"]
        == "日期是{日期}，發生地點是{地點}，當天的敘述是{敘述}"
    )
    assert (
        rows[0][1]["complete_description"]
        == "日期是2026-06-26，發生地點是台北，當天的敘述是例行巡檢完成"
    )



def test_excel_to_sql_composite_mapping_inherits_single_column_match_modes(tmp_path):
    from openpyxl import Workbook
    from data_ingestion.excel_to_sql.task_configs import create_task_config
    from data_ingestion.excel_to_sql.workbook import (
        iter_excel_rows,
        open_workbook_headers,
        read_column_mappings,
    )

    workbook = Workbook()
    ws = workbook.active
    ws.title = "Sheet1"
    ws.append(["日期", "地點（中文）", "完整敘述內容"])
    ws.append(["2026-06-26", "台北", "例行巡檢完成"])
    excel_path = tmp_path / "sample-inherit-match.xlsx"
    workbook.save(excel_path)

    create_task_config(
        {
            "filename": "inherit-match.json",
            "task_name": "Composite Inherit Match",
            "db_server": "sql.local",
            "db_database": "OpsDb",
            "db_user": "operator",
            "db_password": "secret-value",
            "source_sheet": "Sheet1",
            "destination_table": "dbo.ImportTarget",
            "mappings": [
                {"source_column": "日期", "dest_field": "event_date"},
                {"source_column": "地點", "dest_field": "loc", "match": "prefix"},
                {"source_column": "敘述", "dest_field": "description", "match": "contains"},
            ],
            "composite_mappings": [
                {
                    "source_columns": ["日期", "地點", "敘述"],
                    "dest_field": "complete_description",
                    "template": "日期是{日期}，發生地點是{地點}，當天的敘述是{敘述}",
                }
            ],
        },
        actor="operator",
        root=tmp_path,
    )
    cfg_path = tmp_path / "inherit-match.json"
    cfg = json.loads(cfg_path.read_text(encoding="utf-8"))

    assert cfg["mappings"][3]["compose"]["source_matches"] == {
        "日期": "exact",
        "地點": "prefix",
        "敘述": "contains",
    }

    headers = open_workbook_headers(str(cfg_path), cfg, excel_path)[0]
    mappings = read_column_mappings(cfg, headers.header_map)
    rows = list(
        iter_excel_rows(
            excel_path=excel_path,
            sheet_name="Sheet1",
            header_row=1,
            start_row=2,
            end_row=0,
            mappings=mappings,
            skip_empty_rows=True,
        )
    )

    assert rows[0][1]["complete_description"] == "日期是2026-06-26，發生地點是台北，當天的敘述是例行巡檢完成"

def test_create_excel_to_sql_task_config_rejects_unsafe_filename(tmp_path):
    from data_ingestion.excel_to_sql.task_configs import create_task_config

    payload = {
        "filename": "../bad.json",
        "task_name": "Bad",
        "db_server": "sql.local",
        "db_database": "OpsDb",
        "db_user": "operator",
        "db_password": "secret-value",
        "source_sheet": "Sheet1",
        "destination_table": "dbo.ImportTarget",
        "mappings": [{"source_column": "A", "dest_field": "a"}],
    }

    try:
        create_task_config(payload, actor="operator", root=tmp_path)
    except ValueError as exc:
        assert "設定檔檔名" in str(exc)
    else:
        raise AssertionError("unsafe filename should be rejected")



def test_dashboard_tasks_hides_review_summary_below_db_operator(monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from api.routers import editor_dependencies
    from api.services.auth_store import AuthUser
    import sra_api.routers.excel_to_sql as excel_router

    app = FastAPI()
    app.include_router(excel_router.router)
    monkeypatch.setattr(
        excel_router,
        "discover_dashboard_tasks",
        lambda config_dir: [
            {
                "name": "Task",
                "config_path": "config/excel_to_sql/task.json",
                "review_summary": {
                    "mappings": [{"source_column": "日期", "dest_field": "activity_date"}]
                },
            }
        ],
    )

    def data_editor_user():
        return AuthUser(username="editor", role="data_editor", status="active")

    app.dependency_overrides[editor_dependencies.current_user] = data_editor_user
    client = TestClient(app)

    response = client.get("/v1/excel-to-sql/dashboard-tasks")

    assert response.status_code == 200
    task = response.json()["tasks"][0]
    assert task["name"] == "Task"
    assert "review_summary" not in task


def test_dashboard_tasks_keeps_review_summary_for_db_operator(monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from api.routers import editor_dependencies
    from api.services.auth_store import AuthUser
    import sra_api.routers.excel_to_sql as excel_router

    app = FastAPI()
    app.include_router(excel_router.router)
    monkeypatch.setattr(
        excel_router,
        "discover_dashboard_tasks",
        lambda config_dir: [
            {
                "name": "Task",
                "config_path": "config/excel_to_sql/task.json",
                "review_summary": {
                    "mappings": [{"source_column": "日期", "dest_field": "activity_date"}]
                },
            }
        ],
    )

    def db_operator_user():
        return AuthUser(username="operator", role="db_operator", status="active")

    app.dependency_overrides[editor_dependencies.current_user] = db_operator_user
    client = TestClient(app)

    response = client.get("/v1/excel-to-sql/dashboard-tasks")

    assert response.status_code == 200
    task = response.json()["tasks"][0]
    assert task["review_summary"]["mappings"][0]["source_column"] == "日期"

def test_excel_to_sql_task_review_api_requires_superuser(monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from api.routers import editor_dependencies
    from api.services.auth_store import AuthUser
    import sra_api.routers.excel_to_sql as excel_router

    app = FastAPI()
    app.include_router(excel_router.router)

    def db_operator_user():
        return AuthUser(username="operator", role="db_operator", status="active")

    app.dependency_overrides[editor_dependencies.current_user] = db_operator_user
    client = TestClient(app)

    list_response = client.get("/v1/excel-to-sql/task-configs")
    status_response = client.post(
        "/v1/excel-to-sql/task-configs/status",
        json={"config": "pending.json", "enabled": True},
    )
    approve_response = client.post(
        "/v1/excel-to-sql/task-configs/approve", json={"config": "pending.json"}
    )

    assert list_response.status_code == 200
    assert client.get("/v1/excel-to-sql/task-configs?scope=all").status_code == 403
    assert status_response.status_code == 403
    assert approve_response.status_code == 403


def test_task_config_owner_scoped_list_disable_and_delete(tmp_path):
    from data_ingestion.excel_to_sql.task_configs import (
        create_task_config,
        delete_task_config,
        list_task_configs,
        set_task_config_enabled,
    )

    base_payload = {
        "task_name": "Owner Task",
        "db_server": "sql.local",
        "db_database": "OpsDb",
        "db_user": "operator",
        "db_password": "secret-value",
        "source_sheet": "Sheet1",
        "destination_table": "dbo.ImportTarget",
        "mappings": [{"source_column": "A", "dest_field": "a"}],
    }
    create_task_config(
        {**base_payload, "filename": "alice.json"}, actor="alice", root=tmp_path
    )
    create_task_config(
        {**base_payload, "filename": "bob.json"}, actor="bob", root=tmp_path
    )

    alice_tasks = list_task_configs(root=tmp_path, created_by="alice")
    assert [task["created_by"] for task in alice_tasks] == ["alice"]
    assert alice_tasks[0]["config_path"].endswith("alice.json")

    disabled = set_task_config_enabled(
        alice_tasks[0]["config_path"], enabled=False, actor="alice", root=tmp_path
    )
    assert disabled["enabled"] is False
    assert disabled["status"] == "pending"

    bob_config = alice_tasks[0]["config_path"].replace("alice", "bob")
    try:
        delete_task_config(
            bob_config,
            actor="alice",
            root=tmp_path,
        )
    except PermissionError as exc:
        assert "只能刪除自己建立" in str(exc)
    else:
        raise AssertionError("operator must not delete another actor's task config")

    superuser_deleted = delete_task_config(
        bob_config, actor="root", is_superuser=True, root=tmp_path
    )
    assert superuser_deleted["deleted"] is True
    assert superuser_deleted["deleted_by"] == "root"

    deleted = delete_task_config(
        alice_tasks[0]["config_path"], actor="alice", root=tmp_path
    )
    assert deleted["deleted"] is True
    assert list_task_configs(root=tmp_path, created_by="alice") == []

def test_excel_to_sql_composite_template_does_not_keep_empty_rows(tmp_path):
    from openpyxl import Workbook
    from data_ingestion.excel_to_sql.workbook import iter_excel_rows, read_column_mappings

    workbook = Workbook()
    ws = workbook.active
    ws.title = "Sheet1"
    ws.append(["日期", "地", "敘"])
    ws.append(["2026-06-26", "台北", "例行巡檢完成"])
    ws.append([None, None, None])
    excel_path = tmp_path / "sample-template-empty-row.xlsx"
    workbook.save(excel_path)

    cfg = {
        "mappings": [
            {"source_column": "日期", "dest_field": "dat", "dest_field_type": "DATE"},
            {"source_column": "地", "dest_field": "loc", "dest_field_type": "NVARCHAR(255)"},
            {"source_column": "敘", "dest_field": "des", "dest_field_type": "NVARCHAR(MAX)"},
            {
                "dest_field": "cdes",
                "dest_field_type": "NVARCHAR(MAX)",
                "compose": {
                    "source_columns": ["日期", "地", "敘"],
                    "separator": " ",
                    "skip_empty": True,
                    "template": "日期是{日期}，發生地點是{地}，當天的敘述是{敘}",
                },
            },
        ]
    }
    mappings = read_column_mappings(cfg, {"日期": 1, "地": 2, "敘": 3})

    rows = list(
        iter_excel_rows(
            excel_path=excel_path,
            sheet_name="Sheet1",
            header_row=1,
            start_row=2,
            end_row=3,
            mappings=mappings,
            skip_empty_rows=True,
        )
    )

    assert len(rows) == 1
    assert rows[0][0] == 2
    assert rows[0][1]["cdes"] == "日期是2026-06-26，發生地點是台北，當天的敘述是例行巡檢完成"
