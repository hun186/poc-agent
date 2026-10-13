"""P2.1 Product API read/preview: RBAC, lifecycle hiding, safe projections."""
import json
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from fastapi import FastAPI
from fastapi.middleware.cors import CORSMiddleware
from fastapi.testclient import TestClient

from api.routers.editor_dependencies import current_user
from api.services.auth_store import AuthUser
from sra_api.routers.requirement_tasks import product_service, router
from sra_api.middleware.product_preview_body_limit import ProductPreviewBodyLimitMiddleware
from sra_api.services.requirement_product_service import RequirementProductService


class ProductStore:
    def __init__(self):
        self.task = {
            "task_id": "task-1", "schedule_id": "schedule-1",
            "source_id": "source-1", "execution_kind": "recurring",
            "product_state": "active", "revision": 1,
            "created_by": "operator", "created_at": None,
            "updated_by": "operator", "updated_at": None,
        }
        self.calls = []

    def list_product_tasks(self, *, limit, offset):
        self.calls.append(("tasks", limit, offset))
        return [self.task] if offset == 0 else []

    def _projection(self):
        return {
            **self.task, "schedule_name": "Monthly brief",
            "schedule_enabled": True, "schedule_active_version_id": "version-1",
            "schedule_next_run_at": None,
        }

    def list_product_task_projections(self, *, limit, offset):
        self.calls.append(("tasks", limit, offset))
        return [self._projection()] if offset == 0 and self.task["product_state"] != "provisioning" else []

    def get_product_task_projection(self, task_id):
        if task_id != self.task["task_id"] or self.task["product_state"] == "provisioning":
            raise KeyError("requirement_task_not_found")
        return self._projection()

    def get_product_task(self, task_id):
        if task_id != self.task["task_id"] or self.task["product_state"] == "provisioning":
            raise KeyError("requirement_task_not_found")
        return self.task

    def get_schedule(self, schedule_id):
        assert schedule_id == self.task["schedule_id"]
        return {
            "schedule_id": schedule_id, "name": "Monthly brief",
            "enabled": True, "active_version_id": "version-1", "next_run_at": None,
        }

    def list_product_run_bindings(self, *, limit, offset):
        self.calls.append(("runs", limit, offset))
        return [{"batch_run_id": "run-1", "task_id": "task-1"}] if offset == 0 and self.task["product_state"] != "provisioning" else []

    def list_product_run_summaries(self, *, task_id=None, limit, offset):
        self.calls.append(("run_summaries", task_id, limit, offset))
        if (offset != 0 or self.task["product_state"] == "provisioning"
                or task_id not in (None, "task-1")):
            return []
        run = self.get_batch_run(
            schedule_id="schedule-1", batch_run_id="run-1", include_items=False
        )
        return [{"task": self.task, "run": run}]

    def get_product_task_for_run(self, run_id):
        if run_id != "run-1" or self.task["product_state"] == "provisioning":
            raise KeyError("requirement_run_unavailable")
        return self.task

    def list_batch_runs(self, *, schedule_id, limit, offset):
        assert schedule_id == "schedule-1"
        return [{"batch_run_id": "run-1"}] if offset == 0 else []

    def get_batch_run(self, *, schedule_id, batch_run_id, include_items=True):
        assert (schedule_id, batch_run_id) == ("schedule-1", "run-1")
        self.calls.append(("batch_run", include_items))
        result = {
            "batch_run_id": batch_run_id, "schedule_id": schedule_id,
            "schedule_version_id": "version-1", "snapshot_id": "snapshot-1",
            "state": "handoff_pending", "trigger_type": "scheduled",
            "scheduled_for": None, "selected_count": 0, "skipped_count": 0,
            "success_count": 0, "failure_count": 0,
        }
        if include_items:
            result["items"] = []
        return result

    def list_batch_run_items(self, *, schedule_id, batch_run_id, limit, offset):
        assert (schedule_id, batch_run_id) == ("schedule-1", "run-1")
        self.calls.append(("items", limit, offset))
        return []


class PreviewService:
    def __init__(self):
        self.calls = []

    def preview_input(self, **kwargs):
        self.calls.append(kwargs)
        return {"accepted_count": 1, "has_errors": False}


def make_client(role="data_reader", *, store=None, preview=None):
    app = FastAPI()
    app.include_router(router)
    app.add_middleware(ProductPreviewBodyLimitMiddleware)
    store = store or ProductStore()
    service = RequirementProductService(schedule_store=store)
    if preview is not None:
        service.preview_input = preview.preview_input
    app.dependency_overrides[product_service] = lambda: service
    if role is not None:
        app.dependency_overrides[current_user] = lambda: AuthUser(
            username="reviewer", role=role, status="active"
        )
    return TestClient(app), store


def test_reader_list_and_get_task_only_safe_presenter_fields():
    client, store = make_client()
    listing = client.get("/v1/requirement-tasks?limit=12&offset=0")
    assert listing.status_code == 200
    assert store.calls == [("tasks", 12, 0)]
    task = listing.json()["items"][0]
    assert task["task_id"] == "task-1"
    assert task["schedule"]["active_version_id"] == "version-1"
    assert "source_config" not in json.dumps(task)
    assert client.get("/v1/requirement-tasks/task-1").json() == task


def test_reader_product_run_list_detail_and_items_use_safe_presenter():
    client, store = make_client()
    listing = client.get("/v1/requirement-runs?limit=5&offset=0")
    assert listing.status_code == 200
    run = listing.json()["items"][0]
    assert run["run_id"] == "run-1"
    assert run["status"] == "queued"
    assert run["binding"]["snapshot_id"] == "snapshot-1"
    assert "trigger_request_id" not in json.dumps(run)
    assert "artifacts" not in run["links"]  # Deferred endpoint is not advertised.
    assert "items" not in run  # History/detail must not eagerly load Run Items.
    assert client.get("/v1/requirement-runs/run-1").json() == run
    assert client.get("/v1/requirement-tasks/task-1/runs").json()["items"] == [run]
    assert client.get("/v1/requirement-runs/run-1/items?limit=7&offset=3").json() == {
        "run_id": "run-1", "items": [], "limit": 7, "offset": 3,
    }
    assert ("items", 7, 3) in store.calls
    assert all(not x[1] for x in store.calls if x[0] == "batch_run")


def test_unfinalized_product_task_and_run_are_not_visible():
    store = ProductStore()
    store.task["product_state"] = "provisioning"
    client, _ = make_client(store=store)
    assert client.get("/v1/requirement-tasks/task-1").status_code == 404
    assert client.get("/v1/requirement-runs/run-1").status_code == 404
    assert client.get("/v1/requirement-tasks/task-1/runs").status_code == 404
    assert client.get("/v1/requirement-runs").json()["items"] == []


def test_missing_auth_denied_and_preview_operator_only():
    client, _ = make_client(role=None)
    assert client.get("/v1/requirement-tasks").status_code == 401
    assert client.get("/v1/requirement-runs").status_code == 401
    reader, _ = make_client(role="data_reader", preview=PreviewService())
    response = reader.post("/v1/requirement-task-previews", data={"task": '{"profile_version_id":"p1"}', "requirements_json": "[]"})
    assert response.status_code == 403


def test_preview_upload_and_direct_requirements_are_mutually_exclusive():
    preview = PreviewService()
    client, _ = make_client(role="db_operator", preview=preview)
    data = {"task": '{"profile_version_id":"profile-version-1"}'}
    uploaded = client.post(
        "/v1/requirement-task-previews", data=data,
        files={"file": ("requirements.csv", b"A,B\n1,2", "text/csv")},
    )
    assert uploaded.status_code == 200
    assert preview.calls[0]["upload_bytes"] == b"A,B\n1,2"
    assert preview.calls[0]["upload_filename"] == "requirements.csv"
    assert preview.calls[0]["max_rows"] == 5000
    assert preview.calls[0]["max_result_rows"] == 100
    direct = client.post(
        "/v1/requirement-task-previews",
        data={**data, "requirements_json": '[{"ITEMNO":"A"}]'},
    )
    assert direct.status_code == 200
    assert preview.calls[1]["requirements"] == [{"ITEMNO": "A"}]
    assert len(preview.calls) == 2

    both = client.post(
        "/v1/requirement-task-previews",
        data={**data, "requirements_json": "[]"},
        files={"file": ("x.csv", b"x", "text/csv")},
    )
    assert both.status_code == 422
    assert both.json()["detail"] == {"code": "requirement_preview_input_invalid"}
    missing = client.post("/v1/requirement-task-previews", data=data)
    assert missing.status_code == 422
    assert len(preview.calls) == 2


def test_preview_validation_fails_before_calling_normalization():
    preview = PreviewService()
    client, _ = make_client(role="superuser", preview=preview)
    cases = [
        ({"task": "not-json", "requirements_json": "[]"}, None),
        ({"task": '{"profile_version_id":"pv"}', "requirements_json": "not-json"}, None),
        ({"task": '{"profile_version_id":"pv"}'}, ("unknown.exe", b"body", "application/octet-stream")),
        ({"task": '{"profile_version_id":"pv"}'}, ("empty.csv", b"", "text/csv")),
    ]
    for data, upload in cases:
        kwargs = {"data": data}
        if upload:
            kwargs["files"] = {"file": upload}
        response = client.post("/v1/requirement-task-previews", **kwargs)
        assert response.status_code == 422
    assert not preview.calls


def test_disabled_product_dependency_is_503_and_backend_secrets_are_never_returned():
    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[current_user] = lambda: AuthUser(
        username="operator", role="data_reader", status="active"
    )
    client = TestClient(app)
    assert client.get("/v1/requirement-tasks").json() == {
        "detail": {"code": "requirement_dependency_unavailable"}
    }

    class BrokenProductStore(ProductStore):
        def list_product_task_projections(self, *, limit, offset):
            raise RuntimeError("ODBC password=secret")

    client, _ = make_client(store=BrokenProductStore())
    response = client.get("/v1/requirement-tasks")
    assert response.status_code == 500
    assert response.json() == {"detail": {"code": "persistence_error"}}
    assert "secret" not in response.text


def test_main_registration_and_safe_product_dependency_wiring():
    from pathlib import Path
    source = Path("sra_api/main.py").read_text(encoding="utf-8")
    assert "app.include_router(requirement_tasks.router)" in source
    assert "app.add_middleware(ProductPreviewBodyLimitMiddleware)" in source
    assert "if app.state.requirement_import_service is not None:" in source
    assert "app.state.requirement_product_service = RequirementProductService(" in source


def test_store_product_query_guards_use_finalized_state_in_sql():
    from requirement_scheduler.store import ScheduleStore

    store = ScheduleStore(None, "unused")
    calls = []

    def query(sql, *params):
        calls.append((sql, params))
        return []

    store._query = query
    assert store.list_product_tasks(limit=8, offset=12) == []
    assert store.list_product_run_bindings(limit=8, offset=12) == []
    try:
        store.get_product_task_for_run("run-hidden")
        assert False, "unowned run must not resolve"
    except KeyError as exc:
        assert exc.args == ("requirement_run_unavailable",)
    assert len(calls) == 3
    for sql, _ in calls:
        assert "product_state IN ('active','disabled')" in sql
        assert "[sra_schedule].[requirement_product_tasks]" in sql
        assert "provisioning" not in sql
    assert calls[0][1] == (12, 8)
    assert calls[1][1] == (12, 8)
    assert calls[2][1] == ("run-hidden",)
    assert "JOIN" in calls[1][0] and "JOIN" in calls[2][0]


def test_preview_body_is_rejected_before_multipart_spooling_even_without_auth():
    client, _ = make_client(role=None)
    response = client.post(
        "/v1/requirement-task-previews",
        content=b"abc",
        headers={"content-type": "multipart/form-data; boundary=test",
                 "content-length": str(12 * 1024 * 1024)},
    )
    assert response.status_code == 413
    assert response.json() == {"detail": {"code": "requirement_preview_input_invalid"}}


def test_direct_json_oversize_is_rejected_by_form_parser_before_product_handler():
    preview = PreviewService()
    client, _ = make_client(role="db_operator", preview=preview)
    response = client.post(
        "/v1/requirement-task-previews",
        data={"task": '{"profile_version_id":"pv"}',
              "requirements_json": "x" * (1024 * 1024 + 1)},
    )
    # Starlette rejects over-1-MiB form fields before Product input handling.
    # Treat this pre-handler parser error as HTTP 400, not a Product 422.
    assert response.status_code == 400
    assert not preview.calls


def test_preview_rejects_compressed_xlsx_bomb_before_decoding():
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", b"A" * (21 * 1024 * 1024))
    service = RequirementProductService(requirement_import_service=object())
    with pytest.raises(ValueError, match="requirement_preview_too_large"):
        service.preview_input(
            profile_version_id="p", upload_bytes=buffer.getvalue(),
            upload_filename="compressed.xlsx", max_rows=5000, max_result_rows=100,
        )


def test_preview_csv_and_json_row_caps():
    from data_ingestion.requirements.readers import read_csv_rows, read_json_rows
    with pytest.raises(ValueError, match="requirement_preview_too_large"):
        read_csv_rows(b"ITEMNO\n1\n2\n", max_rows=1)
    with pytest.raises(ValueError, match="requirement_preview_too_large"):
        read_json_rows(b'[{"id":"1"},{"id":"2"}]', max_rows=1)


def test_backend_value_error_is_not_reported_as_user_validation():
    class InconsistentStore(ProductStore):
        def list_product_task_projections(self, *, limit, offset):
            raise ValueError("requirement_run_snapshot_binding_missing")
    client, _ = make_client(store=InconsistentStore())
    response = client.get("/v1/requirement-tasks")
    assert response.status_code == 500
    assert response.json() == {"detail": {"code": "persistence_error"}}


def test_atomic_product_projection_queries_and_paged_run_items():
    from requirement_scheduler.store import ScheduleStore
    store = ScheduleStore(None, "unused")
    calls = []
    store._query = lambda sql, *args: calls.append((sql, args)) or []
    assert store.list_product_task_projections(limit=6, offset=2) == []
    with pytest.raises(KeyError, match="requirement_task_not_found"):
        store.get_product_task_projection("hidden")
    assert store.list_batch_run_items(
        schedule_id="schedule-1", batch_run_id="run-1", limit=7, offset=3,
    ) == []
    assert all("JOIN" in sql for sql, _ in calls)
    assert "product_state IN ('active','disabled')" in calls[0][0]
    for sql, _ in calls[:2]:
        assert "[requirement_product_tasks] AS t WITH (HOLDLOCK)" in sql
        assert "[requirement_schedules] AS s WITH (HOLDLOCK)" in sql
    assert calls[0][1] == (2, 6)
    assert calls[1][1] == ("hidden",)
    assert "OFFSET ? ROWS FETCH NEXT ? ROWS ONLY" in calls[2][0]
    assert calls[2][1] == ("schedule-1", "run-1", 3, 7)


def test_missing_joined_schedule_field_is_masked_as_server_error():
    class MalformedProjectionStore(ProductStore):
        def get_product_task_projection(self, task_id):
            projection = super().get_product_task_projection(task_id)
            del projection["schedule_name"]
            return projection

    client, _ = make_client(store=MalformedProjectionStore())
    response = client.get("/v1/requirement-tasks/task-1")
    assert response.status_code == 500
    assert response.json() == {"detail": {"code": "persistence_error"}}


def test_explicit_not_found_remains_404():
    client, _ = make_client()
    response = client.get("/v1/requirement-tasks/does-not-exist")
    assert response.status_code == 404
    assert response.json() == {"detail": {"code": "requirement_task_not_found"}}


def test_run_list_uses_bounded_bulk_projection_without_n_plus_one():
    client, store = make_client()
    assert client.get("/v1/requirement-runs?limit=40&offset=0").status_code == 200
    assert ("run_summaries", None, 40, 0) in store.calls
    assert not any(call[0] == "runs" for call in store.calls)
    assert client.get("/v1/requirement-tasks/task-1/runs?limit=10").status_code == 200
    assert ("run_summaries", "task-1", 10, 0) in store.calls


def test_xlsx_preview_limits_total_decoded_cells():
    from types import SimpleNamespace
    from openpyxl import Workbook
    from data_ingestion.requirements.readers import read_xlsx_rows

    profile = SimpleNamespace(
        worksheet=SimpleNamespace(name=None, header_row=1, first_data_row=2)
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([f"field{i}" for i in range(128)])
    for _ in range(400):
        sheet.append(["v"] * 128)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    with pytest.raises(ValueError, match="requirement_preview_too_large"):
        read_xlsx_rows(stream.getvalue(), profile, max_rows=5000)


def test_store_run_history_projection_one_sql_statement_and_task_filter():
    from requirement_scheduler.store import ScheduleStore
    store = ScheduleStore(None, "unused")
    statements = []
    store._query = lambda sql, *params: statements.append((sql, params)) or []
    assert store.list_product_run_summaries(limit=40, offset=7) == []
    assert store.list_product_run_summaries(task_id="task-1", limit=4, offset=2) == []
    assert len(statements) == 2
    for sql, _ in statements:
        assert "JOIN" in sql
        assert "requirement_product_tasks" in sql
        assert "requirement_batch_runs" in sql
        assert "requirement_batch_run_items" in sql  # Inlined failure evidence
        assert "product_state IN ('active','disabled')" in sql
        assert "OFFSET ? ROWS FETCH NEXT ? ROWS ONLY" in sql
    assert statements[0][1] == (7, 40)
    assert statements[1][1] == ("task-1", 2, 4)


def test_malformed_but_valid_zip_xlsx_is_client_validation_not_server_error():
    from types import SimpleNamespace
    from data_ingestion.requirements.readers import read_bytes

    archive_data = BytesIO()
    with ZipFile(archive_data, "w", ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", "<worksheet/>")
    profile = SimpleNamespace(
        format="xlsx",
        worksheet=SimpleNamespace(name=None, header_row=1, first_data_row=2),
    )
    with pytest.raises(ValueError, match="^source_format_invalid$"):
        read_bytes(archive_data.getvalue(), "invalid.xlsx", profile, max_rows=5000)


def test_malformed_xlsx_workbook_sheet_id_type_error_maps_to_input_error():
    from types import SimpleNamespace
    from openpyxl import Workbook
    from data_ingestion.requirements.readers import read_bytes

    workbook_bytes = BytesIO()
    workbook = Workbook()
    workbook.active.append(["ITEMNO"])
    workbook.active.append(["1"])
    workbook.save(workbook_bytes)
    workbook.close()

    malformed_bytes = BytesIO()
    with ZipFile(BytesIO(workbook_bytes.getvalue())) as source:
        with ZipFile(malformed_bytes, "w", ZIP_DEFLATED) as invalid:
            for member in source.infolist():
                data = source.read(member.filename)
                if member.filename == "xl/workbook.xml":
                    assert b'sheetId="1"' in data
                    data = data.replace(b'sheetId="1"', b'sheetId="invalid"', 1)
                invalid.writestr(member, data)

    profile = SimpleNamespace(
        format="xlsx",
        worksheet=SimpleNamespace(name=None, header_row=1, first_data_row=2),
    )
    with pytest.raises(ValueError, match="^source_format_invalid$"):
        read_bytes(malformed_bytes.getvalue(), "malformed-sheet-id.xlsx", profile, max_rows=5000)


def test_cors_allowed_origin_survives_preparser_413_body_rejection():
    app = FastAPI()
    app.add_middleware(ProductPreviewBodyLimitMiddleware)
    app.add_middleware(
        CORSMiddleware,
        allow_origins=["https://ui.example.invalid"],
        allow_credentials=True,
        allow_methods=["POST"],
        allow_headers=["Authorization", "Content-Type"],
    )
    client = TestClient(app)
    response = client.post(
        "/v1/requirement-task-previews",
        content=b"x" * (11 * 1024 * 1024 + 1),
        headers={
            "Origin": "https://ui.example.invalid",
            "Content-Type": "multipart/form-data; boundary=test",
        },
    )
    assert response.status_code == 413
    assert response.json() == {"detail": {"code": "requirement_preview_input_invalid"}}
    assert response.headers["access-control-allow-origin"] == "https://ui.example.invalid"
    assert response.headers["access-control-allow-credentials"] == "true"


def test_application_registers_cors_outside_product_preview_body_limiter():
    from pathlib import Path
    source = Path("sra_api/main.py").read_text(encoding="utf-8")
    # Starlette executes the last registered middleware first.
    assert source.index("app.add_middleware(ProductPreviewBodyLimitMiddleware)") < (
        source.index("app.add_middleware(\n        CORSMiddleware,")
    )


def test_preview_rejects_sparse_wide_csv_before_materializing_rows():
    from data_ingestion.requirements.readers import read_csv_rows

    # 400 sparse rows x 128 headers exceed the shared 50,000-cell cap,
    # although only one field on each row contains a value.
    headers = [f"column_{index}" for index in range(128)]
    records = [",".join(headers)]
    records.extend(["value" + "," * 127] * 400)
    content = ("\n".join(records) + "\n").encode("utf-8")
    assert len(content) < 10 * 1024 * 1024
    with pytest.raises(ValueError, match="^requirement_preview_too_large$"):
        read_csv_rows(content, max_rows=5000)


def test_preview_rejects_wide_csv_header_and_preserves_internal_reader():
    from data_ingestion.requirements.readers import read_csv_rows

    content = (",".join(f"c{i}" for i in range(129)) + "\n" +
               ",".join("x" for _ in range(129)) + "\n").encode()
    with pytest.raises(ValueError, match="^requirement_preview_too_large$"):
        read_csv_rows(content, max_rows=5000)
    internal_rows, _ = read_csv_rows(content)
    assert len(internal_rows) == 1
    assert len(internal_rows[0]) == 129


def test_preview_rejects_wide_json_before_normalization():
    from data_ingestion.requirements.readers import read_json_rows

    payload = [{f"c{i}": "x" for i in range(129)}]
    content = json.dumps(payload).encode("utf-8")
    with pytest.raises(ValueError, match="^requirement_preview_too_large$"):
        read_json_rows(content, max_rows=5000)
    internal_rows, _ = read_json_rows(content)
    assert len(internal_rows[0]) == 129


def test_preview_csv_and_json_cell_caps_accept_small_inputs():
    from data_ingestion.requirements.readers import read_csv_rows, read_json_rows

    csv_rows, csv_numbers = read_csv_rows(b"ITEMNO,TITLE\n1,Example\n", max_rows=5000)
    assert csv_rows == [{"ITEMNO": "1", "TITLE": "Example"}]
    assert csv_numbers == [2]
    json_rows, json_numbers = read_json_rows(b'[{"ITEMNO": "1"}]', max_rows=5000)
    assert json_rows == [{"ITEMNO": "1"}]
    assert json_numbers == [1]


def test_json_preview_bounded_decode_rejects_millions_of_small_rows():
    from data_ingestion.requirements.readers import read_json_rows

    # A compact array can have far more than 5,000 records while fitting
    # comfortably inside the uploaded-file size cap. Stop before decoding it.
    many = b"[" + b"{}," * 500_000 + b"{}]"
    assert len(many) < 10 * 1024 * 1024
    with pytest.raises(ValueError, match="^requirement_preview_too_large$"):
        read_json_rows(many, max_rows=5000)


def test_json_preview_streamed_rows_preserve_nested_and_escaped_values():
    from data_ingestion.requirements.readers import read_json_rows

    examples = [
        {"ITEMNO": "one", "text": 'embedded "quotes", ] and }', "details": {"tags": [1, 2]}},
        {"ITEMNO": "two", "text": "escaped \\ backslash and \n newline"},
    ]
    raw = (" \n\t" + json.dumps(examples, ensure_ascii=False) + " \n").encode("utf-8")
    decoded, numbers = read_json_rows(raw, max_rows=5000)
    assert decoded == examples
    assert numbers == [1, 2]


def test_json_preview_rejects_trailing_commas_and_oversized_single_object():
    from data_ingestion.requirements.readers import read_json_rows

    for malformed in (b'[{},]', b'[{},,{}]', b'[{}]{"other":true}', b'{"obj":{}}'):
        with pytest.raises(ValueError, match="^source_format_invalid$"):
            read_json_rows(malformed, max_rows=5000)
    one_huge_row = b'[{"text":"' + b"x" * (1024 * 1024 + 20) + b'"}]'
    with pytest.raises(ValueError, match="^requirement_preview_too_large$"):
        read_json_rows(one_huge_row, max_rows=5000)


def test_json_internal_reader_keeps_full_array_without_preview_cap():
    from data_ingestion.requirements.readers import read_json_rows

    rows, numbers = read_json_rows(b"[{}, {}, {}]")
    assert rows == [{}, {}, {}]
    assert numbers == [1, 2, 3]


def test_json_preview_enforces_single_object_budget_in_utf8_bytes():
    from data_ingestion.requirements.readers import read_json_rows

    # 400k CJK code points encode to ~1.2 MiB despite fewer than
    # 1 MiB of decoded Python string code points.
    multibyte_json = json.dumps(
        [{"text": "界" * 400_000}], ensure_ascii=False
    ).encode("utf-8")
    assert 1024 * 1024 < len(multibyte_json) < 10 * 1024 * 1024
    with pytest.raises(ValueError, match="^requirement_preview_too_large$"):
        read_json_rows(multibyte_json, max_rows=5000)

    # A genuinely small multibyte JSON object remains valid.
    small_json = json.dumps(
        [{"text": "界" * 1000}], ensure_ascii=False
    ).encode("utf-8")
    rows, numbers = read_json_rows(small_json, max_rows=5000)
    assert rows == [{"text": "界" * 1000}]
    assert numbers == [1]


def test_json_preview_rejects_noncomma_nonbracket_element_delimiters():
    from data_ingestion.requirements.readers import read_json_rows

    malformed = [
        b'[{} } {"x":1}]',
        b'[{} } {}]',
        b'[{} {"x":1}]',
        b'[{} : {"x":1}]',
    ]
    for raw in malformed:
        with pytest.raises(ValueError, match="^source_format_invalid$"):
            read_json_rows(raw, max_rows=5000)

    accepted, positions = read_json_rows(
        b'[{}, {"x":1}]', max_rows=5000
    )
    assert accepted == [{}, {"x": 1}]
    assert positions == [1, 2]


@pytest.mark.parametrize(
    "payload",
    [
        b'[{"text":"\\ud800"}]',
        b'[{"\\ud800":"text"}]',
        b'[{"text":{"nested":"\\ud800"}}]',
        b'[{"text":{"\\ud800":"value"}}]',
        b'[{"text":["normal", "\\udc00"]}]',
    ],
)
def test_preview_json_unpaired_surrogate_escape_is_invalid_input(payload):
    from data_ingestion.requirements.readers import read_json_rows

    # Even though Python json.loads accepts isolated surrogate escapes,
    # the preview cannot re-encode them to UTF-8 for byte accounting.
    with pytest.raises(ValueError, match="^source_format_invalid$"):
        read_json_rows(payload, max_rows=5000)


def test_preview_json_valid_surrogate_pair_and_cjk_are_accepted():
    from data_ingestion.requirements.readers import read_json_rows

    payload = b'[{"emoji":"\\ud83d\\ude00","text":"\\u754c"}]'
    rows, numbers = read_json_rows(payload, max_rows=5000)
    assert rows == [{"emoji": "\U0001F600", "text": "\u754c"}]
    assert numbers == [1]



def test_malformed_xlsx_content_types_xml_is_input_error_with_lxml():
    from types import SimpleNamespace
    from openpyxl import Workbook
    import openpyxl.xml
    from data_ingestion.requirements.readers import read_bytes

    # lxml is optional in openpyxl; exercise its XMLSyntaxError path when
    # available rather than inadvertently testing only stdlib ParseError.
    if not openpyxl.xml.LXML:
        pytest.skip("optional openpyxl lxml parser not enabled")

    valid = BytesIO()
    book = Workbook()
    book.active.append(["ITEMNO"])
    book.active.append(["1"])
    book.save(valid)
    book.close()

    malformed = BytesIO()
    with ZipFile(BytesIO(valid.getvalue())) as source:
        with ZipFile(malformed, "w", ZIP_DEFLATED) as target:
            for member in source.infolist():
                payload = source.read(member.filename)
                if member.filename == "[Content_Types].xml":
                    assert b"<Types" in payload
                    payload = payload.replace(b"<Types", b'<Types bad="&"', 1)
                target.writestr(member, payload)

    profile = SimpleNamespace(
        format="xlsx",
        worksheet=SimpleNamespace(name=None, header_row=1, first_data_row=2),
    )
    with pytest.raises(ValueError, match="^source_format_invalid$"):
        read_bytes(
            malformed.getvalue(), "broken-content-types.xlsx", profile,
            max_rows=5000,
        )


def test_json_preview_bounds_nested_object_expansion_before_row_decode():
    from data_ingestion.requirements.readers import read_json_rows

    # Around 9 KiB on wire but 3,000 individual Python dict objects. 1 MiB
    # byte and 128 top-level column budgets alone cannot bound this.
    nested = b'[{"ITEMNO":"1","data":[' + b"{}," * 2999 + b"{}]}]"
    assert len(nested) < 1024 * 1024
    with pytest.raises(ValueError, match="^requirement_preview_too_large$"):
        read_json_rows(nested, max_rows=5000)


def test_json_preview_bounds_aggregate_nested_nodes_across_rows():
    from data_ingestion.requirements.readers import read_json_rows

    # Each row stays below the per-item node cap, but the collection
    # exceeds the total structural-token budget.
    row = b'{"ITEMNO":"1","data":[' + b"0," * 2999 + b"0]}"
    raw = b"[" + b",".join([row] * 40) + b"]"
    assert len(raw) < 1024 * 1024
    with pytest.raises(ValueError, match="^requirement_preview_too_large$"):
        read_json_rows(raw, max_rows=5000)


def test_json_preview_accepts_small_nested_objects_and_strings_with_punctuation():
    from data_ingestion.requirements.readers import read_json_rows

    record = [{"ITEMNO": "1", "data": [{}, {"codes": [1, 2, 3]}],
               "note": "literal brackets [ ] and commas , : { }"}]
    payload = json.dumps(record, ensure_ascii=False).encode("utf-8")
    rows, numbers = read_json_rows(payload, max_rows=5000)
    assert rows == record
    assert numbers == [1]


def test_internal_json_import_retains_previous_nested_input_behavior():
    from data_ingestion.requirements.readers import read_json_rows

    raw = b'[{"data":[' + b"{}," * 2999 + b"{}]}]"
    rows, numbers = read_json_rows(raw)
    assert len(rows) == 1 and len(rows[0]["data"]) == 3000
    assert numbers == [1]


def test_json_preview_two_valid_top_level_rows_decode_separately():
    from data_ingestion.requirements.readers import read_json_rows

    rows, numbers = read_json_rows(
        b'[{"ITEMNO":"A"},{"ITEMNO":"B","TITLE":"ok"}]',
        max_rows=5000,
    )
    assert rows == [
        {"ITEMNO": "A"},
        {"ITEMNO": "B", "TITLE": "ok"},
    ]
    assert numbers == [1, 2]


def test_direct_json_preview_rejects_nested_amplification_before_service():
    preview = PreviewService()
    client, _ = make_client(role="db_operator", preview=preview)
    nested = '[{"ITEMNO":"A","data":[' + "{}," * 2999 + "{}]}]"
    assert len(nested.encode("utf-8")) < 1024 * 1024
    response = client.post(
        "/v1/requirement-task-previews",
        data={
            "task": '{"profile_version_id":"pv"}',
            "requirements_json": nested,
        },
    )
    assert response.status_code == 422
    assert response.json() == {"detail": {"code": "requirement_preview_too_large"}}
    assert preview.calls == []


def test_direct_json_preview_preserves_multirow_input_to_service():
    preview = PreviewService()
    client, _ = make_client(role="db_operator", preview=preview)
    response = client.post(
        "/v1/requirement-task-previews",
        data={
            "task": '{"profile_version_id":"pv"}',
            "requirements_json": '[{"ITEMNO":"A"},{"ITEMNO":"B"}]',
        },
    )
    assert response.status_code == 200
    assert len(preview.calls) == 1
    assert preview.calls[0]["requirements"] == [
        {"ITEMNO": "A"}, {"ITEMNO": "B"}
    ]


def test_direct_json_preview_malformed_input_still_uses_input_error_code():
    preview = PreviewService()
    client, _ = make_client(role="db_operator", preview=preview)
    response = client.post(
        "/v1/requirement-task-previews",
        data={
            "task": '{"profile_version_id":"pv"}',
            "requirements_json": '[{"ITEMNO":"A"},]',
        },
    )
    assert response.status_code == 422
    assert response.json() == {"detail": {"code": "requirement_preview_input_invalid"}}
    assert preview.calls == []


def test_preview_task_metadata_size_checked_before_eager_json_decode():
    preview = PreviewService()
    client, _ = make_client(role="db_operator", preview=preview)
    # Legal multipart size, but this task metadata contains thousands of
    # hidden objects beyond the two allowed schema fields.
    task_metadata = '{"profile_version_id":"pv","extra":[' + "{}," * 4000 + "{}]}"
    assert 4096 < len(task_metadata.encode("utf-8")) < 1024 * 1024
    response = client.post(
        "/v1/requirement-task-previews",
        data={"task": task_metadata, "requirements_json": "[]"},
    )
    assert response.status_code == 422
    assert response.json() == {"detail": {"code": "requirement_preview_input_invalid"}}
    assert preview.calls == []


def test_preview_task_metadata_accepts_declared_maximum_lengths():
    preview = PreviewService()
    client, _ = make_client(role="db_operator", preview=preview)
    metadata = json.dumps({"profile_version_id": "p" * 128, "name": "名" * 240})
    response = client.post(
        "/v1/requirement-task-previews",
        data={"task": metadata, "requirements_json": "[]"},
    )
    assert response.status_code == 200
    assert len(preview.calls) == 1
    assert preview.calls[0]["profile_version_id"] == "p" * 128


def test_xlsx_preview_rejects_large_central_directory_before_zipfile(monkeypatch):
    from sra_api.services import requirement_product_service as module

    data = BytesIO()
    with ZipFile(data, "w", ZIP_DEFLATED) as archive:
        for i in range(2001):
            archive.writestr(f"empty-{i:04d}.txt", b"")

    def forbid_open(*args, **kwargs):
        pytest.fail("ZipFile allocated ZipInfo before raw metadata preflight")

    monkeypatch.setattr(module, "ZipFile", forbid_open)
    service = RequirementProductService(requirement_import_service=object())
    with pytest.raises(ValueError, match="^requirement_preview_too_large$"):
        service.preview_input(
            profile_version_id="pv",
            upload_bytes=data.getvalue(),
            upload_filename="many-members.xlsx",
            max_rows=5000,
            max_result_rows=100,
        )


def test_xlsx_preview_rejects_understated_central_directory_count():
    from sra_api.services.requirement_product_service import (
        _preflight_preview_zip_entries,
    )
    from struct import pack_into

    data = BytesIO()
    with ZipFile(data, "w", ZIP_DEFLATED) as archive:
        for i in range(2001):
            archive.writestr(f"file-{i:04d}", b"")
    mutated = bytearray(data.getvalue())
    eocd = mutated.rfind(b"PK\x05\x06")
    assert eocd > 0
    # Rewrite the untrusted EOCD declared counts to try to hide 2,001 rows.
    pack_into("<H", mutated, eocd + 8, 1)
    pack_into("<H", mutated, eocd + 10, 1)
    with pytest.raises(ValueError, match="^requirement_preview_too_large$"):
        _preflight_preview_zip_entries(bytes(mutated))


def test_xlsx_preview_raw_central_directory_accepts_small_archive():
    from sra_api.services.requirement_product_service import (
        _preflight_preview_zip_entries,
    )

    data = BytesIO()
    with ZipFile(data, "w", ZIP_DEFLATED) as archive:
        archive.writestr("xl/workbook.xml", b"<workbook/>")
        archive.writestr("xl/worksheets/sheet1.xml", b"<worksheet/>")
    assert _preflight_preview_zip_entries(data.getvalue()) is None


def test_empty_xlsx_sheet_is_client_input_error():
    from types import SimpleNamespace
    from openpyxl import Workbook
    from data_ingestion.requirements.readers import read_bytes

    book = Workbook()
    data = BytesIO()
    book.save(data)
    book.close()
    profile = SimpleNamespace(
        format="xlsx",
        worksheet=SimpleNamespace(name=None, header_row=1, first_data_row=2),
    )
    with pytest.raises(ValueError, match="^source_format_invalid$"):
        read_bytes(data.getvalue(), "empty.xlsx", profile, max_rows=5000)
