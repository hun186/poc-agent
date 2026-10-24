from __future__ import annotations

import ast
import hashlib
import importlib.util
import io
import json
import os
import re
import sys
import zipfile
from datetime import datetime, timedelta, timezone
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from operator_e2e.requirement_e2e_support import (
    E2EConfig,
    ExitCode,
    HarnessFailure,
    StageStatus,
    atomic_write_json,
    bounded_log_tail,
    choose_fixture,
    classify_new_batch,
    load_e2e_environment,
    recover_ambiguous_write,
    redact,
    result_document,
    validate_artifact,
    validate_live_platform,
    validate_next_run,
    validate_preview,
    validate_remote_items,
    validate_terminal,
    verify_worker_identity,
)


@pytest.fixture(autouse=True)
def _reviewed_runtime_authority(monkeypatch):
    monkeypatch.setattr(
        "operator_e2e.test_requirement_schedule_e2e._provenance_source",
        lambda _root, _path, expected=None: {
            "repository": "hun186/poc-agent", "reviewed_runtime_sha": expected or "a" * 40,
            "provenance_fingerprint": "f" * 64, "identity_verified": True,
        },
    )


def test_support_imports_when_fcntl_is_unavailable(monkeypatch):
    source = Path("operator_e2e/requirement_e2e_support.py")
    real_find_spec = importlib.util.find_spec
    monkeypatch.setattr(
        importlib.util,
        "find_spec",
        lambda name, *args, **kwargs: None if name == "fcntl" else real_find_spec(name, *args, **kwargs),
    )
    spec = importlib.util.spec_from_file_location("requirement_e2e_support_without_fcntl", source)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    monkeypatch.setitem(sys.modules, spec.name, module)
    spec.loader.exec_module(module)
    assert module.fcntl is None


def test_unsupported_platform_fails_before_harness_run(tmp_path, monkeypatch, capsys):
    from operator_e2e.test_requirement_schedule_e2e import Harness, main

    output = tmp_path / "evidence"
    run_calls = []
    monkeypatch.setattr(Harness, "run", lambda self: run_calls.append(self))
    monkeypatch.setattr(
        "operator_e2e.test_requirement_schedule_e2e.validate_live_platform",
        lambda *_: (_ for _ in ()).throw(HarnessFailure(
            "unsupported_local_platform", stage="CONFIG_PREFLIGHT", exit_code=ExitCode.PREFLIGHT,
        )),
    )

    code = main(["--excel", "unused.xlsx", "--env-file", str(tmp_path / "missing"), "--output-dir", str(output)])

    assert code == 2
    assert run_calls == []
    result = json.loads(next(output.glob("*/result.json")).read_text())
    assert result["failed_stage"] == "CONFIG_PREFLIGHT"
    assert result["failure_category"] == "unsupported_local_platform"
    assert "Linux deployment host" in capsys.readouterr().out


def test_linux_platform_preflight_accepts_required_primitives(monkeypatch):
    import operator_e2e.requirement_e2e_support as support

    monkeypatch.setattr(support.sys, "platform", "linux")
    monkeypatch.setattr(support.os, "name", "posix")
    monkeypatch.setattr(support.os, "geteuid", lambda: 1000)
    monkeypatch.setattr(support, "fcntl", SimpleNamespace(flock=lambda *_: None))
    monkeypatch.setattr(support.Path, "is_dir", lambda self: str(self) == "/proc")
    validate_live_platform()


@pytest.mark.parametrize("phase0_keys", [False, True])
def test_missing_requirement_env_file_has_actionable_reminder(tmp_path, monkeypatch, capsys, phase0_keys):
    from operator_e2e.test_requirement_schedule_e2e import Harness, main

    for key in tuple(os.environ):
        if key.startswith("SRA_E2E_"):
            monkeypatch.delenv(key, raising=False)
    if phase0_keys:
        for name in ("OPENWEBUI_KEY", "READ_ONLY_KEY", "REPORT_KEY"):
            monkeypatch.setenv(f"SRA_E2E_{name}", "not-a-requirement-credential")
    run_calls = []
    monkeypatch.setattr(Harness, "run", lambda self: run_calls.append(self))
    output = tmp_path / "evidence"

    assert main(["--excel", "unused.xlsx", "--env-file", str(tmp_path / "missing.env"), "--output-dir", str(output)]) == 2
    assert run_calls == []
    result = json.loads(next(output.glob("*/result.json")).read_text())
    assert result["failure_category"] == "configuration_file_missing"
    assert result["failed_stage"] == "CONFIG_PREFLIGHT"
    stdout = capsys.readouterr().out
    assert "operator_e2e/e2e.env.example" in stdout
    assert "operator_e2e/e2e.env.local" in stdout


def test_complete_requirement_process_environment_does_not_require_env_file(tmp_path, monkeypatch):
    from operator_e2e.test_requirement_schedule_e2e import Harness, main

    for key in tuple(os.environ):
        if key.startswith("SRA_E2E_"):
            monkeypatch.delenv(key, raising=False)
    monkeypatch.setenv("SRA_E2E_MODEL", "model")
    monkeypatch.setenv("SRA_E2E_EDITOR_TOKEN", "authorized-token")
    monkeypatch.setenv("SRA_E2E_EXPECTED_RUNTIME_SHA", "919422bce992e757a89767f66beb8e042707a976")
    monkeypatch.setattr(
        "config.deployment_bootstrap.requirement_handoff_output_root",
        lambda _root: tmp_path / "artifacts",
    )
    monkeypatch.setattr(Harness, "run", lambda self: None)

    assert main([
        "--excel", "unused.xlsx", "--env-file", str(tmp_path / "missing.env"),
        "--output-dir", str(tmp_path / "evidence"),
    ]) == ExitCode.PASS


def test_cli_help_explains_e2e_external_modes_and_production_boundary(capsys):
    from operator_e2e.test_requirement_schedule_e2e import parse_args

    assert parse_args([]).external_mode is None
    assert parse_args(["--external-mode", "live"]).external_mode == "live"

    with pytest.raises(SystemExit) as exc:
        parse_args(["--help"])

    assert exc.value.code == 0
    help_text = capsys.readouterr().out
    assert "External capability modes apply ONLY to verified Operator E2E test jobs" in help_text
    assert "Production/user Report Jobs are different" in help_text
    assert "If --external-mode is omitted" in help_text
    assert "uses the authenticated mode already" in help_text
    assert "Changing this CLI option does not reconfigure" in help_text


def test_harness_omitted_external_mode_adopts_authenticated_deployment_mode(tmp_path, monkeypatch):
    import operator_e2e.test_requirement_schedule_e2e as module

    class FakeApi:
        token = "token"
        def get(self, path, **_kwargs):
            if path == "/healthz":
                return {"status": "ok"}
            if path in ("/v1/requirement-import-profiles", "/v1/requirement-schedules"):
                return {"items": []}
            raise AssertionError(path)

    class FakeEvidenceProvider:
        def __init__(self, *_args, **_kwargs):
            pass
        def preflight(self, _verification):
            return {
                "acceptance": {"requested_mode": "live", "classification": "FULL_LIVE"},
                "workers": {
                    "handoff": {"alive": True, "identity_verified": True},
                    "schedule": {"alive": True, "identity_verified": True},
                },
                "deployment_source": {
                    "identity_verified": True,
                    "provenance_fingerprint": "f" * 64,
                },
            }

    result = result_document("run")
    result["operator_source"] = {
        "identity_verified": True,
        "provenance_fingerprint": "f" * 64,
    }
    config = SimpleNamespace(
        mode="local", artifact_root=tmp_path, output_dir=tmp_path,
        token="token", password=None, expected_runtime_sha=None,
    )
    harness = module.Harness(config, tmp_path / "unused.xlsx", "run", result)
    harness.api = FakeApi()
    monkeypatch.setattr(module, "LocalDeploymentEvidenceProvider", FakeEvidenceProvider)
    monkeypatch.setattr(harness, "validate_input", lambda: (_ for _ in ()).throw(RuntimeError("stop-after-mode")))
    monkeypatch.setattr(module, "_validate_deployment_source", lambda *_args, **_kwargs: None)

    with pytest.raises(RuntimeError, match="stop-after-mode"):
        harness.run()

    assert harness.external_mode == "live"
    assert result["acceptance"]["requested_mode"] == "live"
    runtime_stage = next(item for item in result["stages"] if item["name"] == "RUNTIME_WORKERS")
    assert runtime_stage["status"] == StageStatus.PASS


def test_harness_omitted_external_mode_fails_closed_when_deployment_mode_missing(tmp_path, monkeypatch):
    import operator_e2e.test_requirement_schedule_e2e as module

    class FakeApi:
        token = "token"
        def get(self, path, **_kwargs):
            if path == "/healthz":
                return {"status": "ok"}
            if path in ("/v1/requirement-import-profiles", "/v1/requirement-schedules"):
                return {"items": []}
            raise AssertionError(path)

    class FakeEvidenceProvider:
        def __init__(self, *_args, **_kwargs):
            pass
        def preflight(self, _verification):
            return {
                "acceptance": {},
                "workers": {
                    "handoff": {"alive": True, "identity_verified": True},
                    "schedule": {"alive": True, "identity_verified": True},
                },
            }

    result = result_document("run")
    config = SimpleNamespace(
        mode="local", artifact_root=tmp_path, output_dir=tmp_path,
        token="token", password=None, expected_runtime_sha=None,
    )
    harness = module.Harness(config, tmp_path / "unused.xlsx", "run", result)
    harness.api = FakeApi()
    monkeypatch.setattr(module, "LocalDeploymentEvidenceProvider", FakeEvidenceProvider)

    with pytest.raises(HarnessFailure, match="external_capability_mode_mismatch") as failure:
        harness.run()

    assert failure.value.stage == "RUNTIME_WORKERS"
    assert result["acceptance"]["requested_mode"] is None
    assert result["acceptance"]["deployed_mode"] is None


def test_harness_explicit_external_mode_still_fails_closed_on_mismatch(tmp_path, monkeypatch):
    import operator_e2e.test_requirement_schedule_e2e as module

    class FakeApi:
        token = "token"
        def get(self, path, **_kwargs):
            if path == "/healthz":
                return {"status": "ok"}
            if path in ("/v1/requirement-import-profiles", "/v1/requirement-schedules"):
                return {"items": []}
            raise AssertionError(path)

    class FakeEvidenceProvider:
        def __init__(self, *_args, **_kwargs):
            pass
        def preflight(self, _verification):
            return {
                "acceptance": {"requested_mode": "live", "classification": "FULL_LIVE"},
                "workers": {
                    "handoff": {"alive": True, "identity_verified": True},
                    "schedule": {"alive": True, "identity_verified": True},
                },
            }

    result = result_document("run")
    config = SimpleNamespace(
        mode="local", artifact_root=tmp_path, output_dir=tmp_path,
        token="token", password=None, expected_runtime_sha=None,
    )
    harness = module.Harness(config, tmp_path / "unused.xlsx", "run", result, external_mode="auto")
    harness.api = FakeApi()
    monkeypatch.setattr(module, "LocalDeploymentEvidenceProvider", FakeEvidenceProvider)

    with pytest.raises(HarnessFailure, match="external_capability_mode_mismatch") as failure:
        harness.run()

    assert failure.value.stage == "RUNTIME_WORKERS"
    assert result["acceptance"] == {
        "requested_mode": "auto",
        "deployed_mode": "live",
        "remediation_hint": "set REQUIREMENT_E2E_EXTERNAL_MODE and restart API/workers",
    }


def test_summary_reports_effective_deployed_external_mode_when_cli_omits_it():
    from operator_e2e.test_requirement_schedule_e2e import _summary_text, parse_args

    args = parse_args([])
    result = {
        "overall_status": "PASS",
        "acceptance": {"requested_mode": "live", "classification": "FULL_LIVE"},
        "verification": {
            "source_attestation": "VERIFIED",
            "deployment_acceptance": "ESTABLISHED",
        },
    }

    summary = _summary_text(result, args, "FULL_LIVE", ExitCode.PASS)

    assert "EXTERNAL_MODE = live" in summary
    assert "FRESH_LIVE_REQUIREMENT_E2E_PASS = YES" in summary


def test_summary_reports_unresolved_when_cli_and_deployment_modes_are_missing():
    from operator_e2e.test_requirement_schedule_e2e import _summary_text, parse_args

    args = parse_args([])
    result = {
        "overall_status": "FAIL",
        "failure_category": "external_capability_mode_mismatch",
        "acceptance": {
            "requested_mode": None,
            "deployed_mode": None,
            "remediation_hint": "set REQUIREMENT_E2E_EXTERNAL_MODE and restart API/workers",
        },
        "verification": {
            "source_attestation": "VERIFIED",
            "deployment_acceptance": "NOT_ESTABLISHED",
        },
    }

    summary = _summary_text(result, args, "UNRESOLVED", ExitCode.PREFLIGHT)

    assert "EXTERNAL_MODE = UNRESOLVED" in summary
    assert "REQUESTED_EXTERNAL_MODE = \n" in summary
    assert "DEPLOYED_EXTERNAL_MODE = \n" in summary


def test_runbook_separates_release_acceptance_from_external_mode_assertion():
    runbook = Path("docs/operations/requirement-automatic-operator-e2e.md").read_text(encoding="utf-8")
    assert "Formal release/deployment acceptance requires `--verification release`" in runbook
    assert "independent of whether `--external-mode` is omitted or supplied" in runbook

def test_common_environment_example_is_safe_and_complete():
    example = Path("operator_e2e/e2e.env.example").read_text(encoding="utf-8")
    for name in (
        "BASE_URL", "OPENWEBUI_KEY", "READ_ONLY_KEY", "REPORT_KEY", "EDITOR_USERNAME",
        "EDITOR_PASSWORD", "EDITOR_TOKEN", "MODEL", "REQUIREMENT_ARTIFACT_ROOT",
        "EXPECTED_SOURCE_SHA",
    ):
        assert f"SRA_E2E_{name}" in example
    assert "SRA_E2E_EXPECTED_RUNTIME_SHA is optional" in example
    assert "must NOT automatically be reused" in example
    assert "SRA_E2E_EDITOR_TOKEN = SRA_E2E_REPORT_KEY" in example
    assert "Operator E2E TEST policy only" in example
    assert "never inherit these E2E mocks" in example
    ignore_rules = Path(".gitignore").read_text(encoding="utf-8").splitlines()
    assert "operator_e2e/e2e.env.local" in ignore_rules
    assert "operator_e2e/e2e.env.example" not in ignore_rules


def test_phase0_base_url_prefers_common_name_and_retains_legacy_fallback(monkeypatch):
    source = Path("operator_e2e/test_phase0_api_routing_evidence.py").read_text(encoding="utf-8")
    tree = ast.parse(source)
    assignment = next(node for node in tree.body if isinstance(node, ast.Assign) and any(
        isinstance(target, ast.Name) and target.id == "BASE" for target in node.targets
    ))
    expression = compile(ast.Expression(assignment.value), "<phase0-base>", "eval")
    monkeypatch.setenv("SRA_E2E_BASE_URL", "http://common")
    monkeypatch.setenv("SRA_E2E_API_BASE", "http://legacy")
    assert eval(expression, {"os": os}) == "http://common"
    monkeypatch.delenv("SRA_E2E_BASE_URL")
    assert eval(expression, {"os": os}) == "http://legacy"


def test_env_contract_rejects_foreign_key_and_process_env_wins(tmp_path, monkeypatch):
    env_file = tmp_path / "e2e.env.local"
    env_file.write_text("OTHER=x\nSRA_E2E_MODEL=file\n", encoding="utf-8")
    with pytest.raises(HarnessFailure, match="configuration_error"):
        load_e2e_environment(env_file, {})
    env_file.write_text("SRA_E2E_MODEL=file\n", encoding="utf-8")
    loaded = load_e2e_environment(env_file, {
        "SRA_E2E_MODEL": "process",
        "REQUIREMENT_HANDOFF_OUTPUT_ROOT": str(tmp_path / "production-artifacts"),
    })
    assert loaded["SRA_E2E_MODEL"] == "process"
    assert "REQUIREMENT_HANDOFF_OUTPUT_ROOT" not in loaded


def test_config_auth_and_required_values(tmp_path):
    base = {"SRA_E2E_MODEL": "m", "SRA_E2E_REQUIREMENT_ARTIFACT_ROOT": str(tmp_path)}
    with pytest.raises(HarnessFailure):
        E2EConfig.from_environment(base)
    with pytest.raises(HarnessFailure):
        E2EConfig.from_environment({**base, "SRA_E2E_EDITOR_TOKEN": "t", "SRA_E2E_EDITOR_USERNAME": "u", "SRA_E2E_EDITOR_PASSWORD": "p"})
    config = E2EConfig.from_environment({**base, "SRA_E2E_EDITOR_TOKEN": "t"})
    assert config.model == "m" and config.exit_code_for_failure == ExitCode.PREFLIGHT


def test_local_config_uses_production_artifact_root_without_duplicate_e2e_setting(tmp_path):
    production_root = tmp_path / "production-artifacts"
    config = E2EConfig.from_environment(
        {"SRA_E2E_MODEL": "m", "SRA_E2E_EDITOR_TOKEN": "t"},
        production_artifact_root=production_root,
    )
    assert config.artifact_root == production_root.resolve()


def test_local_config_rejects_artifact_root_that_disagrees_with_production(tmp_path):
    with pytest.raises(HarnessFailure, match="configuration_error") as failure:
        E2EConfig.from_environment(
            {
                "SRA_E2E_MODEL": "m",
                "SRA_E2E_EDITOR_TOKEN": "t",
                "SRA_E2E_REQUIREMENT_ARTIFACT_ROOT": str(tmp_path / "operator-artifacts"),
            },
            production_artifact_root=tmp_path / "production-artifacts",
        )
    assert failure.value.stage == "CONFIG_PREFLIGHT"


def test_local_config_accepts_matching_artifact_root_assertion(tmp_path):
    production_root = tmp_path / "production-artifacts"
    config = E2EConfig.from_environment(
        {
            "SRA_E2E_MODEL": "m",
            "SRA_E2E_EDITOR_TOKEN": "t",
            "SRA_E2E_REQUIREMENT_ARTIFACT_ROOT": str(
                production_root / ".." / "production-artifacts"
            ),
        },
        production_artifact_root=production_root,
    )
    assert config.artifact_root == production_root.resolve()


@pytest.mark.parametrize(
    ("name", "value"),
    [
        ("SRA_E2E_EXPECTED_SOURCE_SHA", "a" * 40),
        ("SRA_E2E_EXPECTED_RUNTIME_SHA", "a" * 64),
        ("SRA_E2E_EXPECTED_SOURCE_SHA", "z" * 64),
        ("SRA_E2E_EXPECTED_RUNTIME_SHA", "z" * 40),
    ],
)
def test_config_rejects_malformed_source_and_runtime_shas(tmp_path, name, value):
    env = {
        "SRA_E2E_MODEL": "m",
        "SRA_E2E_REQUIREMENT_ARTIFACT_ROOT": str(tmp_path),
        "SRA_E2E_EDITOR_TOKEN": "t",
        name: value,
    }
    with pytest.raises(HarnessFailure, match="configuration_error") as failure:
        E2EConfig.from_environment(env)
    assert failure.value.stage == "CONFIG_PREFLIGHT"


def test_input_sha_validation_and_evidence_are_input_only(tmp_path):
    from operator_e2e.test_requirement_schedule_e2e import Harness

    excel = Path("operator_e2e/input/requirement_e2e_3items.xlsx")
    digest = hashlib.sha256(excel.read_bytes()).hexdigest()
    result = result_document("run")
    config = SimpleNamespace(
        output_dir=tmp_path, token="token", password=None, expected_source_sha=digest,
    )
    harness = Harness(config, excel, "run", result)
    harness.validate_input()
    assert result["input"]["sha256"] == digest
    assert result["input"]["expected_sha256"] == digest

    mismatch = Harness(
        SimpleNamespace(**{**config.__dict__, "expected_source_sha": "0" * 64}),
        excel,
        "run-2",
        result_document("run-2"),
    )
    with pytest.raises(HarnessFailure, match="input_sha256_mismatch") as failure:
        mismatch.validate_input()
    assert failure.value.stage == "INPUT_VALIDATION"


def test_runtime_release_authority_is_not_xlsx_source_sha():
    config = E2EConfig.from_environment({
        "SRA_E2E_MODEL": "model", "SRA_E2E_EDITOR_TOKEN": "token",
        "SRA_E2E_EXPECTED_SOURCE_SHA": "1" * 64,
    }, "remote")
    assert config.expected_source_sha == "1" * 64
    assert config.expected_runtime_sha is None


def test_runtime_sha_mismatch_fails_before_harness_mutations(tmp_path, monkeypatch):
    from operator_e2e.test_requirement_schedule_e2e import Harness, main

    artifact_root = tmp_path / "artifact-root"
    artifact_root.mkdir()
    env_file = tmp_path / "e2e.env.local"
    env_file.write_text(
        f"SRA_E2E_MODEL=model\nSRA_E2E_REQUIREMENT_ARTIFACT_ROOT={artifact_root}\n"
        f"SRA_E2E_EDITOR_TOKEN=token\nSRA_E2E_EXPECTED_RUNTIME_SHA={'b' * 40}\n",
        encoding="utf-8",
    )
    mutations = []
    monkeypatch.setattr(Harness, "run", lambda self: mutations.append("run"))
    monkeypatch.setattr(
        "operator_e2e.test_requirement_schedule_e2e._provenance_source",
        lambda *_args: (_ for _ in ()).throw(HarnessFailure(
            "legacy_runtime_sha_mismatch", stage="CONFIG_PREFLIGHT", exit_code=ExitCode.PREFLIGHT,
        )),
    )

    output = tmp_path / "evidence"
    code = main([
        "--excel", "unused.xlsx", "--env-file", str(env_file), "--output-dir", str(output),
    ])
    assert code == ExitCode.PREFLIGHT
    assert mutations == []
    result = json.loads(next(output.glob("*/result.json")).read_text())
    assert result["failed_stage"] == "CONFIG_PREFLIGHT"
    assert result["failure_category"] == "legacy_runtime_sha_mismatch"


def test_dirty_reviewed_runtime_fails_before_harness_mutations(tmp_path, monkeypatch):
    from operator_e2e.test_requirement_schedule_e2e import Harness, main

    reviewed_sha = "a" * 40
    monkeypatch.setenv("SRA_E2E_EXPECTED_RUNTIME_SHA", reviewed_sha)
    artifact_root = tmp_path / "artifact-root"
    artifact_root.mkdir()
    env_file = tmp_path / "e2e.env.local"
    env_file.write_text(
        f"SRA_E2E_MODEL=model\nSRA_E2E_REQUIREMENT_ARTIFACT_ROOT={artifact_root}\n"
        f"SRA_E2E_EDITOR_TOKEN=token\nSRA_E2E_EXPECTED_RUNTIME_SHA={reviewed_sha}\n",
        encoding="utf-8",
    )
    run_calls = []
    monkeypatch.setattr(Harness, "run", lambda self: run_calls.append(self))
    monkeypatch.setattr(
        "operator_e2e.test_requirement_schedule_e2e._provenance_source",
        lambda *_args: (_ for _ in ()).throw(HarnessFailure(
            "operator_runtime_source_mismatch", stage="CONFIG_PREFLIGHT", exit_code=ExitCode.PREFLIGHT,
        )),
    )

    output = tmp_path / "evidence"
    code = main([
        "--excel", "unused.xlsx", "--env-file", str(env_file), "--output-dir", str(output),
    ])

    assert code == ExitCode.PREFLIGHT
    assert len(run_calls) == 0
    result = json.loads(next(output.glob("*/result.json")).read_text())
    assert result["failed_stage"] == "CONFIG_PREFLIGHT"
    assert result["failure_category"] == "operator_runtime_source_mismatch"


def test_matching_release_fingerprint_passes_source_validation():
    from operator_e2e.test_requirement_schedule_e2e import _validate_deployment_source

    _validate_deployment_source({"identity_verified": True, "provenance_fingerprint": "a" * 64}, "a" * 64)


def test_runbook_uses_offline_provenance_authority_and_stable_xlsx_digest():
    runbook = Path("docs/operations/requirement-automatic-operator-e2e.md").read_text(encoding="utf-8")
    assert "REQUIREMENT_E2E_PROVENANCE.json" in runbook
    assert "neither `.git` nor GitHub access is required" in runbook
    assert (
        "SRA_E2E_EXPECTED_SOURCE_SHA="
        "816ffac9039530a86955c590b0fa68830757a620229d85630f5e42c12f542403"
    ) in runbook


def test_redaction_removes_headers_fields_and_known_literals():
    secret = "SENTINEL-SECRET"
    value = {"Authorization": f"Bearer {secret}", "nested": f"password={secret} token={secret} PWD={secret} User ID=bob"}
    encoded = json.dumps(redact(value, [secret]))
    assert secret not in encoded and "bob" not in encoded and encoded.count("REDACTED") >= 2


def test_result_schema_and_atomic_evidence_does_not_overwrite(tmp_path):
    doc = result_document("run-1")
    assert doc["schema_version"] == "requirement_e2e_result_v2"
    assert all(stage["status"] == StageStatus.NOT_REACHED for stage in doc["stages"])
    target = tmp_path / "result.json"
    atomic_write_json(target, doc)
    with pytest.raises(FileExistsError):
        atomic_write_json(target, doc)


@pytest.mark.parametrize("items,expected", [([], "create"), ([{"fingerprint": "x"}], "reuse")])
def test_fixture_create_and_reuse(items, expected):
    assert choose_fixture(items, lambda item: item["fingerprint"] == "x")[0] == expected


def test_fixture_drift_and_ambiguity_fail_closed():
    with pytest.raises(HarnessFailure, match="fixture_drift"):
        choose_fixture([{"fingerprint": "bad"}], lambda item: item["fingerprint"] == "x")
    with pytest.raises(HarnessFailure, match="fixture_ambiguous"):
        choose_fixture([{"fingerprint": "x"}, {"fingerprint": "x"}], lambda item: True)


def test_ambiguous_write_rereads_before_retry():
    assert recover_ambiguous_write(lambda: [{"id": "winner"}]) == {"id": "winner"}
    assert recover_ambiguous_write(lambda: []) is None
    with pytest.raises(HarnessFailure, match="ambiguous_write"):
        recover_ambiguous_write(lambda: [{"id": "a"}, {"id": "b"}])


def test_preview_requires_exactly_three():
    validate_preview({"has_errors": False, "accepted_count": 3, "rejected_count": 0})
    with pytest.raises(HarnessFailure, match="preview_rejected"):
        validate_preview({"has_errors": False, "accepted_count": 2, "rejected_count": 0})


def _batch(**changes):
    value = {"batch_run_id": "b", "trigger_type": "scheduled", "schedule_version_id": "v", "scheduled_for": "2026-01-01T00:00:00+00:00", "snapshot_id": "s", "selected_count": 3, "skipped_count": 0}
    value.update(changes)
    return value


def test_automatic_batch_identity_and_rejections():
    assert classify_new_batch([_batch()], set(), "v", "s", "2026-01-01T00:00:00+00:00")["batch_run_id"] == "b"
    with pytest.raises(HarnessFailure, match="snapshot_identity_mismatch"):
        classify_new_batch([_batch(snapshot_id="wrong")], set(), "v", "s", "2026-01-01T00:00:00+00:00")
    with pytest.raises(HarnessFailure, match="multiple_scheduled_runs"):
        classify_new_batch([_batch(batch_run_id="a"), _batch(batch_run_id="b")], set(), "v", "s", "2026-01-01T00:00:00+00:00")


def test_remote_job_and_terminal_validation():
    items = [{"selection_status": "selected", "remote_job_id": f"j{i}", "remote_job_state": "completed", "execution_state": "completed"} for i in range(3)]
    validate_remote_items(items)
    validate_terminal({"state": "completed", "success_count": 3, "failure_count": 0, "items": items})
    items[2]["remote_job_id"] = "j1"
    with pytest.raises(HarnessFailure, match="remote_job_count_mismatch"):
        validate_remote_items(items)
    items[2].update(remote_job_id="j2", remote_job_state="failed")
    with pytest.raises(HarnessFailure, match="remote_execution_failed"):
        validate_terminal({"state": "running", "success_count": 2, "failure_count": 1, "items": items})


def _docx(path: Path, member="word/document.xml"):
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr(member, "<document/>")


def test_artifact_fingerprint_docx_and_sha(tmp_path):
    root = tmp_path / "root"; root.mkdir(); path = root / "a.docx"; _docx(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    copied = tmp_path / "copy.docx"
    evidence = validate_artifact(root, hashlib.sha256(str(root.resolve()).encode()).hexdigest(), "a.docx", digest, copied)
    assert evidence["actual_sha256"] == digest and copied.read_bytes() == path.read_bytes()
    with pytest.raises(HarnessFailure, match="artifact_sha256_mismatch"):
        validate_artifact(root, hashlib.sha256(str(root.resolve()).encode()).hexdigest(), "a.docx", "0" * 64, tmp_path / "bad.docx")
    evil = root / "evil.docx"; _docx(evil, "../evil")
    with pytest.raises(HarnessFailure, match="artifact_invalid"):
        validate_artifact(root, hashlib.sha256(str(root.resolve()).encode()).hexdigest(), "evil.docx", hashlib.sha256(evil.read_bytes()).hexdigest(), tmp_path / "evil-copy.docx")


def test_worker_identity_verification(tmp_path, monkeypatch):
    identity = tmp_path / "identity"
    identity.write_text(f"role=schedule\npid={os.getpid()}\nprocess_start_time=123\n", encoding="utf-8")
    monkeypatch.setattr("operator_e2e.requirement_e2e_support._proc_start_time", lambda pid: "123")
    monkeypatch.setattr("operator_e2e.requirement_e2e_support._proc_uid", lambda pid: os.geteuid())
    monkeypatch.setattr("operator_e2e.requirement_e2e_support._proc_cwd", lambda pid: tmp_path.resolve())
    # The canonical verifier now reads bounded NUL-delimited /proc argv.
    # Stub the reader actually used, not the obsolete flattened string helper.
    import operator_e2e.requirement_e2e_support as support
    argv = b"python\0-m\0requirement_scheduler.cli\0schedule-worker\0"
    monkeypatch.setattr(
        support, "open",
        lambda path, mode: (
            io.BytesIO(argv) if path == f"/proc/{os.getpid()}/cmdline"
            and mode == "rb" else pytest.fail("unexpected proc read")
        ),
        raising=False,
    )
    assert verify_worker_identity(identity, "schedule", tmp_path)["alive"] is True
    identity.write_text(f"role=schedule\npid={os.getpid()}\nprocess_start_time=999\n", encoding="utf-8")
    with pytest.raises(HarnessFailure, match="worker_unavailable"):
        verify_worker_identity(identity, "schedule", tmp_path)


def test_worker_identity_rejects_wrong_command_and_cwd(tmp_path, monkeypatch):
    identity = tmp_path / "identity"
    identity.write_text(f"role=handoff\npid={os.getpid()}\nprocess_start_time=123\n", encoding="utf-8")
    monkeypatch.setattr("operator_e2e.requirement_e2e_support._proc_start_time", lambda pid: "123")
    monkeypatch.setattr("operator_e2e.requirement_e2e_support._proc_uid", lambda pid: os.geteuid())
    monkeypatch.setattr("operator_e2e.requirement_e2e_support._proc_cwd", lambda pid: tmp_path.resolve())
    import operator_e2e.requirement_e2e_support as support
    # A real NUL-delimited unrelated argv must be rejected.
    argv = b"python\0unrelated.py\0handoff-worker\0"
    monkeypatch.setattr(
        support, "open",
        lambda path, mode: (
            io.BytesIO(argv) if path == f"/proc/{os.getpid()}/cmdline"
            and mode == "rb" else pytest.fail("unexpected proc read")
        ),
        raising=False,
    )
    with pytest.raises(HarnessFailure, match="worker_unavailable"):
        verify_worker_identity(identity, "handoff", tmp_path)
    # Also preserve the original cwd-negative contract with an otherwise
    # recognized worker launch command.
    monkeypatch.setattr(support, "_proc_cwd", lambda pid: tmp_path / "wrong-cwd")
    worker_argv = b"python\0-m\0requirement_scheduler.cli\0handoff-worker\0"
    monkeypatch.setattr(
        support, "open",
        lambda path, mode: (
            io.BytesIO(worker_argv) if path == f"/proc/{os.getpid()}/cmdline"
            and mode == "rb" else pytest.fail("unexpected proc read")
        ),
    )
    with pytest.raises(HarnessFailure, match="worker_unavailable"):
        verify_worker_identity(identity, "handoff", tmp_path)


def test_bounded_log_tail_is_capped_and_redacted(tmp_path):
    secret = "SENTINEL-LOG-SECRET"
    path = tmp_path / "worker.log"
    path.write_text("\n".join(["old"] * 120 + [f"Authorization: Bearer {secret}", f"password={secret}"]), encoding="utf-8")
    tail = bounded_log_tail(path, [secret], max_lines=3, max_bytes=65536)
    assert secret not in tail and len(tail.splitlines()) == 3


def test_next_run_advancement_contract():
    scheduled_for = datetime(2026, 1, 1, tzinfo=timezone.utc)
    validate_next_run(scheduled_for, datetime(2026, 1, 8, tzinfo=timezone.utc))
    with pytest.raises(HarnessFailure, match="automatic_claim_identity_mismatch"):
        validate_next_run(scheduled_for, scheduled_for + timedelta(minutes=1))


def test_exit_codes_are_stable():
    assert tuple(int(code) for code in ExitCode) == (0, 2, 3, 4, 5)


def test_entrypoint_contains_no_manual_run_or_worker_execution():
    source = Path("operator_e2e/test_requirement_schedule_e2e.py").read_text(encoding="utf-8")
    assert 'post(f"/v1/requirement-schedules/{schedule_id}/runs"' not in source.lower()
    assert "schedule-worker --once" not in source
    assert "coordinator.run_once" not in source


def test_tracked_xlsx_fixture_has_canonical_three_item_shape():
    path = Path("operator_e2e/input/requirement_e2e_3items.xlsx")
    book = load_workbook(path, read_only=True)
    rows = list(book.active.iter_rows(values_only=True)); book.close()
    assert rows[0] == (
        "ITEMNO", "TITLE", "CONTENT ", "UPDT", "DeadlineDT",
        "CreateApeno", "UpdateApeno", "CrtDT",
    )
    assert len(rows[1:]) == 3


def test_config_failure_still_writes_safe_evidence(tmp_path, monkeypatch):
    from operator_e2e.test_requirement_schedule_e2e import main
    for key in tuple(os.environ):
        if key.startswith("SRA_E2E_"): monkeypatch.delenv(key, raising=False)
    output = tmp_path / "evidence"
    code = main(["--excel", "missing.xlsx", "--env-file", str(tmp_path / "missing.env"), "--output-dir", str(output)])
    assert code == ExitCode.PREFLIGHT
    runs = list(output.iterdir())
    assert len(runs) == 1
    assert {item.name for item in runs[0].iterdir()} == {"result.json", "summary.txt", "diagnostics.json", "artifacts"}


def test_cleanup_failure_forces_exit_five_and_evidence(tmp_path, monkeypatch):
    from operator_e2e.test_requirement_schedule_e2e import Harness, main
    artifact_root = tmp_path / "artifacts-root"; artifact_root.mkdir()
    monkeypatch.setenv("SRA_E2E_MODEL", "model")
    monkeypatch.setenv("SRA_E2E_REQUIREMENT_ARTIFACT_ROOT", str(artifact_root))
    monkeypatch.setenv("SRA_E2E_EDITOR_TOKEN", "SENTINEL-TOKEN")
    monkeypatch.setenv("SRA_E2E_EXPECTED_RUNTIME_SHA", "919422bce992e757a89767f66beb8e042707a976")
    monkeypatch.setattr(Harness, "run", lambda self: setattr(self, "enabled_schedule_id", "schedule-1"))
    monkeypatch.setattr(Harness, "cleanup", lambda self: (_ for _ in ()).throw(RuntimeError("disable failed")))
    output = tmp_path / "evidence"
    code = main(["--excel", "unused.xlsx", "--env-file", str(tmp_path / "missing.env"), "--output-dir", str(output)])
    assert code == ExitCode.CLEANUP
    result = json.loads(next(output.glob("*/result.json")).read_text())
    assert result["failure_category"] == "cleanup_failed"
    assert "SENTINEL-TOKEN" not in json.dumps(result)


def test_disabled_schedule_batch_keeps_polling_identity(monkeypatch, tmp_path):
    from operator_e2e.test_requirement_schedule_e2e import Harness

    schedule_id, batch_id = "schedule-durable", "batch-durable"
    initial = (datetime.now(timezone.utc) + timedelta(seconds=30)).replace(microsecond=0)
    selected = [
        {
            "run_item_id": f"item-{index}", "requirement_id": f"req-{index}",
            "selection_status": "selected", "remote_job_id": f"job-{index}",
            "dispatched_at": initial.isoformat(), "remote_job_state": "completed",
            "execution_state": "completed",
        }
        for index in range(3)
    ]
    batch = {
        "schedule_id": schedule_id, "batch_run_id": batch_id, "trigger_type": "scheduled",
        "schedule_version_id": "version-1", "scheduled_for": initial.isoformat(),
        "snapshot_id": "snapshot-1", "selected_count": 3, "skipped_count": 0,
        "success_count": 3, "failure_count": 0, "state": "completed", "items": selected,
    }

    class FakeApi:
        def __init__(self): self.poll_paths, self.enabled, self.version_payload = [], False, None
        def get(self, path, **_kwargs):
            if path == f"/v1/requirement-schedules/{schedule_id}/runs?limit=100":
                return {"items": [{"batch_run_id": batch_id}]} if self.enabled else {"items": []}
            if path == f"/v1/requirement-schedules/{schedule_id}":
                return {"versions": [], "next_run_at": (initial + timedelta(days=7)).isoformat()}
            if path == f"/v1/requirement-schedules/{schedule_id}/runs/{batch_id}":
                self.poll_paths.append(path); return batch
            raise AssertionError(path)
        def post(self, path, **_kwargs):
            if path.endswith("/versions"):
                self.version_payload = _kwargs["json"]
                return {"schedule_version_id": "version-1", "version_number": 1}
            if path.endswith("/approve"): return {}
            raise AssertionError(path)
        def patch(self, path, *, json, **_kwargs):
            assert path == f"/v1/requirement-schedules/{schedule_id}"
            if json["enabled"]:
                self.enabled = True
                return {"enabled": True, "active_version_id": "version-1", "next_run_at": initial.isoformat()}
            return {"enabled": False}

    config = SimpleNamespace(anchor_lead_seconds=30, allowed_anchor_window_seconds=60,
                             auto_claim_timeout_seconds=1, handoff_timeout_seconds=1,
                             remote_timeout_seconds=1, model="model", output_dir=tmp_path,
                             token="token", password=None)
    result = result_document("run")
    harness = Harness(config, tmp_path / "unused.xlsx", "run", result)
    harness.api = FakeApi()
    detail = harness.schedule_and_wait({"schedule_id": schedule_id}, "snapshot-1", {f"req-{i}" for i in range(3)})

    assert harness.enabled_schedule_id is None
    assert harness.api.version_payload["model"] == "requirement-e2e:model"
    assert harness.wait_terminal(detail)["batch_run_id"] == batch_id
    assert harness.api.poll_paths and set(harness.api.poll_paths) == {f"/v1/requirement-schedules/{schedule_id}/runs/{batch_id}"}


def test_local_main_uses_deployment_state_artifact_root_without_duplicate_env_setting(tmp_path, monkeypatch):
    from operator_e2e.test_requirement_schedule_e2e import Harness, main

    for key in tuple(os.environ):
        if key.startswith("SRA_E2E_"):
            monkeypatch.delenv(key, raising=False)
    production_root = tmp_path / "production-artifacts"
    output = tmp_path / "evidence"
    env_file = tmp_path / "e2e.env.local"
    env_file.write_text(
        "SRA_E2E_MODEL=model\n"
        "SRA_E2E_EDITOR_TOKEN=token\n",
        encoding="utf-8",
    )
    observed = {}
    monkeypatch.setattr(
        "config.deployment_bootstrap.requirement_handoff_output_root",
        lambda _root: production_root,
    )
    monkeypatch.setattr(
        Harness, "run",
        lambda self: observed.update(artifact_root=self.config.artifact_root),
    )

    assert main([
        "--excel", "unused.xlsx", "--env-file", str(env_file),
        "--output-dir", str(output),
    ]) == ExitCode.PASS
    assert observed["artifact_root"] == production_root.resolve()


def test_env_file_output_directory_is_authoritative_for_all_evidence(tmp_path, monkeypatch):
    from operator_e2e.test_requirement_schedule_e2e import Harness, main

    for key in tuple(os.environ):
        if key.startswith("SRA_E2E_"): monkeypatch.delenv(key, raising=False)
    artifact_root = tmp_path / "artifact-root"; artifact_root.mkdir()
    output = tmp_path / "env-output"
    env_file = tmp_path / "e2e.env.local"
    env_file.write_text(
        f"SRA_E2E_MODEL=model\nSRA_E2E_REQUIREMENT_ARTIFACT_ROOT={artifact_root}\n"
        f"SRA_E2E_EDITOR_TOKEN=token\nSRA_E2E_OUTPUT_DIR={output}\n"
        f"SRA_E2E_EXPECTED_RUNTIME_SHA=919422bce992e757a89767f66beb8e042707a976\n", encoding="utf-8",
    )
    observed = {}
    monkeypatch.setattr(Harness, "run", lambda self: observed.update(output=self.output, config=self.config.output_dir))

    assert main(["--excel", "unused.xlsx", "--env-file", str(env_file)]) == ExitCode.PASS
    runs = list(output.iterdir())
    assert observed["config"] == output and observed["output"].parent == output
    assert len([path for path in runs if path.is_dir()]) == 1
    assert {path.name for path in observed["output"].iterdir()} == {"result.json", "summary.txt", "diagnostics.json", "artifacts"}
    assert (output / ".lock").exists()


def test_import_ambiguous_write_ignores_historical_matching_snapshots(tmp_path):
    import requests
    from operator_e2e.test_requirement_schedule_e2e import Harness

    old_ids = {"historical-1", "historical-2", "active-before"}
    new_id = "new-snapshot"

    class FakeApi:
        imported = False
        import_post_count = 0
        def post(self, path, **_kwargs):
            if path.endswith("/preview") or "/preview?" in path:
                return {"has_errors": False, "accepted_count": 3, "rejected_count": 0, "source_sha256": "same-sha"}
            if path.endswith("/import"):
                self.import_post_count += 1
                self.imported = True; raise requests.Timeout()
            raise AssertionError(path)
        def get(self, path, **_kwargs):
            if path == "/v1/requirement-sets":
                active = new_id if self.imported else "active-before"
                return {"items": [{"source_id": "source-1", "requirement_set_id": "set-1", "active_snapshot_id": active}]}
            if path == "/v1/requirement-sets/set-1/snapshots":
                ids = old_ids | ({new_id} if self.imported else set())
                return {"items": [
                    {
                        "snapshot_id": sid,
                        "requirement_set_id": "set-1",
                        "import_run_id": "new-import-run" if sid == new_id else f"import-{sid}",
                        "source_snapshot_sha256": "same-sha",
                        "import_profile_version_id": "profile-v1",
                        "created_at": "2026-09-15T00:00:00+00:00",
                        "activated_at": "2026-09-15T00:00:00+00:00",
                    }
                    for sid in sorted(ids)
                ]}
            if path == f"/v1/requirement-snapshots/{new_id}/items":
                return {"items": [{"requirement_id": f"req-{i}", "expires_at": None} for i in range(3)]}
            raise AssertionError(path)

    excel = tmp_path / "input.xlsx"; excel.write_bytes(b"xlsx")
    config = SimpleNamespace(output_dir=tmp_path, token="token", password=None)
    harness = Harness(config, excel, "run", result_document("run")); harness.api = FakeApi()
    reqset, items = harness.import_requirements({"source_id": "source-1"}, {"profile_version_id": "profile-v1"})
    assert reqset["active_snapshot_id"] == new_id
    assert harness.result["requirement_import"]["snapshot_id"] == new_id
    assert harness.result["requirement_import"]["import_run_id"] == "new-import-run"
    assert harness.api.import_post_count == 1
    assert len(items) == 3
