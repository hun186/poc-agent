from __future__ import annotations

import hashlib
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from data_ingestion.requirements.models import ImportProfileDocument
from data_ingestion.requirements.profile import normalize_rows
from data_ingestion.requirements.readers import read_bytes
from operator_e2e.requirement_e2e_support import CANONICAL_SOURCE_SHA256, HarnessFailure, result_document
from operator_e2e.test_requirement_schedule_e2e import DEFAULT_EXCEL, Harness, PROFILE_PATH, _summary_text


REAL_SOURCE = next(Path("data").rglob("requirement_e2e_three.xlsx"))
PLACEHOLDERS = (
    "first stable Requirement E2E topic",
    "second stable Requirement E2E topic",
    "third stable Requirement E2E topic",
)


def _workbook_rows(path: Path) -> tuple[list[str], list[tuple]]:
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.active
        values = list(sheet.iter_rows(values_only=True))
        return [str(value) for value in values[0]], values[1:]
    finally:
        book.close()


def _harness(tmp_path: Path) -> Harness:
    config = SimpleNamespace(output_dir=tmp_path, token=None, password=None)
    return Harness(config, DEFAULT_EXCEL.resolve(), "real-content", result_document("real-content"))


def test_canonical_workbook_is_exact_real_source_and_has_three_business_rows():
    canonical = DEFAULT_EXCEL.read_bytes()
    assert canonical == REAL_SOURCE.read_bytes()
    assert hashlib.sha256(canonical).hexdigest() == CANONICAL_SOURCE_SHA256
    headers, rows = _workbook_rows(DEFAULT_EXCEL)
    assert headers == [
        "ITEMNO", "TITLE", "CONTENT ", "UPDT", "DeadlineDT",
        "CreateApeno", "UpdateApeno", "CrtDT",
    ]
    assert len(rows) == 3
    content = [str(row[2]) for row in rows]
    assert all(marker in content[0] for marker in ("加薩", "人質", "人道援助"))
    assert all(marker in content[1] for marker in ("臺灣海峽", "灰色地帶"))
    assert all(marker in content[2] for marker in ("美中戰略競爭", "印太"))
    assert not any(placeholder in text for placeholder in PLACEHOLDERS for text in content)


def test_production_profile_maps_real_content_to_prompt_content():
    profile = ImportProfileDocument.model_validate(json.loads(PROFILE_PATH.read_text(encoding="utf-8")))
    rows, numbers = read_bytes(DEFAULT_EXCEL.read_bytes(), DEFAULT_EXCEL.name, profile)
    candidate = normalize_rows(rows, profile, numbers)
    assert not candidate.diagnostics
    assert len(candidate.rows) == 3
    expected = {str(row[0]): str(row[2]) for row in _workbook_rows(DEFAULT_EXCEL)[1]}
    assert {row.external_item_id: row.prompt_content for row in candidate.rows} == expected
    assert profile.mapping["prompt_content"].source == "CONTENT"
    assert "CONTENT " in profile.mapping["prompt_content"].aliases


def test_canonical_snapshot_prompt_guard_accepts_exact_content_and_rejects_drift(tmp_path):
    harness = _harness(tmp_path)
    expected = harness._canonical_prompt_content()
    assert expected is not None and len(expected) == 3
    items = [
        {"external_item_id": external_id, "prompt_content": prompt}
        for external_id, prompt in expected.items()
    ]
    harness._validate_canonical_prompt_content(items)
    items[0] = {**items[0], "prompt_content": "synthetic replacement"}
    with pytest.raises(HarnessFailure, match="prompt_content_mismatch") as failure:
        harness._validate_canonical_prompt_content(items)
    assert failure.value.stage == "REQUIREMENT_ITEMS"
    assert "synthetic replacement" not in json.dumps(harness.result, ensure_ascii=False)


@pytest.mark.parametrize(("row_index", "markers", "query_markers"), [
    (0, ("加薩", "人道援助"), ("加薩", "人道援助")),
    (1, ("臺灣海峽", "灰色地帶"), ("臺灣海峽", "灰色地帶")),
    (2, ("美中戰略競爭", "印太"), ("美中戰略競爭", "印太")),
])
def test_deterministic_boundaries_preserve_real_requirement_semantics(row_index, markers, query_markers):
    from agent.data_agent.builder import DeterministicE2ECodeModel
    from agent.llm_utils.ollama_client import DeterministicE2ELLMClient
    from agent.refiners.query_refiner import REFINER_SYSTEM_PROMPT
    from agent.routing import TASK_ROUTER_SYSTEM

    prompt = str(_workbook_rows(DEFAULT_EXCEL)[1][row_index][2])
    assert all(marker in prompt for marker in markers)
    client = DeterministicE2ELLMClient()
    routed = json.loads(client.chat([
        {"role": "system", "content": TASK_ROUTER_SYSTEM},
        {"role": "user", "content": prompt},
    ]))
    assert routed["mode"] == "SRA_REPORT"
    refined = json.loads(client.chat([
        {"role": "system", "content": REFINER_SYSTEM_PROMPT},
        {"role": "user", "content": f"【使用者問句】\n{prompt}"},
    ]))["refined_question"]
    assert all(marker in refined for marker in markers)
    action = DeterministicE2ECodeModel().generate([
        {"role": "user", "content": refined},
    ]).content
    assert "es_search_news" in action
    assert all(marker in action for marker in query_markers)
    assert not any(placeholder in action for placeholder in PLACEHOLDERS)


def test_runtime_semantic_adapters_do_not_depend_on_stable_topic_placeholders():
    sources = "\n".join(
        Path(path).read_text(encoding="utf-8")
        for path in ("agent/data_agent/builder.py", "agent/llm_utils/ollama_client.py")
    )
    assert not any(placeholder in sources for placeholder in PLACEHOLDERS)


def test_input_failures_have_distinct_safe_categories(tmp_path):
    config = SimpleNamespace(
        output_dir=tmp_path, token=None, password=None, expected_source_sha="0" * 64,
    )
    missing = Harness(config, tmp_path / "missing.xlsx", "missing", result_document("missing"))
    with pytest.raises(HarnessFailure, match="input_file_missing"):
        missing.validate_input()
    invalid_path = tmp_path / "invalid.xlsx"
    invalid_path.write_bytes(b"not an xlsx")
    invalid = Harness(config, invalid_path, "invalid", result_document("invalid"))
    with pytest.raises(HarnessFailure, match="input_xlsx_invalid"):
        invalid.validate_input()
    custom = tmp_path / "custom.xlsx"
    custom.write_bytes(DEFAULT_EXCEL.read_bytes())
    mismatch = Harness(config, custom, "mismatch", result_document("mismatch"))
    with pytest.raises(HarnessFailure, match="input_sha256_mismatch"):
        mismatch.validate_input()
    assert set(mismatch.result["input"]) >= {
        "path", "expected_sha256", "sha256", "remediation_hint",
    }


def test_external_mode_mismatch_summary_is_self_diagnosing_without_secrets():
    result = {
        "overall_status": "FAIL",
        "failure_category": "external_capability_mode_mismatch",
        "verification": {"source_attestation": "VERIFIED", "deployment_acceptance": "NOT_ESTABLISHED"},
        "acceptance": {
            "requested_mode": "offline", "deployed_mode": "auto",
            "remediation_hint": "set REQUIREMENT_E2E_EXTERNAL_MODE and restart API/workers",
        },
    }
    summary = _summary_text(
        result, SimpleNamespace(verification="functional", mode="local", external_mode="offline"),
        "UNRESOLVED", 2,
    )
    assert "REQUESTED_EXTERNAL_MODE = offline" in summary
    assert "DEPLOYED_EXTERNAL_MODE = auto" in summary
    assert "restart API/workers" in summary
