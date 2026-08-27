from io import BytesIO
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document

from api.services.export_word_service import ExportWordService
from agent.agent import _collect_llm_execution, _collect_model_usage


def _paragraph_texts(docx_bytes: bytes) -> list[str]:
    doc = Document(BytesIO(docx_bytes))
    return [paragraph.text for paragraph in doc.paragraphs]


def test_collect_model_usage_preserves_pipeline_order_and_deduplicates() -> None:
    trace = {
        "llm": {
            "router": [{"model_id": "router-model"}],
            "field_selector": [{"model_id": "selector-model"}],
            "report": [{"model_id": "report-model"}],
        },
        "query_refiner": {"report": {"model_id": "refiner-model"}},
        "data_agent": {
            "runs": [
                {"model_id": "ollama_chat/data-model", "model_id_effective": "data-model"},
                {"model_id_effective": "data-model"},
            ]
        },
    }

    assert _collect_model_usage(trace) == [
        {"stage": "任務路由", "model_id": "router-model"},
        {"stage": "問題精煉", "model_id": "refiner-model"},
        {"stage": "資料整備", "model_id": "data-model"},
        {"stage": "證據欄位選擇", "model_id": "selector-model"},
        {"stage": "報告生成", "model_id": "report-model"},
    ]


def test_collect_llm_execution_preserves_processing_flow_order() -> None:
    trace = {
        "llm": {
            "router": [{"model_id": "router"}],
            "field_selector": [{"model_id": "selector"}],
            "report": [{"model_id": "report"}],
        },
        "query_refiner": {"report": {"enabled": True, "model_id": "refiner"}},
        "data_agent": {"runs": [{"model_id_effective": "data"}]},
    }

    assert [row["stage"] for row in _collect_llm_execution(trace)] == [
        "任務路由",
        "問題精煉",
        "資料整備",
        "證據欄位選擇",
        "報告生成",
    ]


def test_export_docx_adds_answer_heading_after_user_question() -> None:
    docx_bytes, filename = ExportWordService().build_docx_bytes(
        {
            "question": "這份資料的重點是什麼？",
            "answer": "重點摘要如下。",
        }
    )

    assert filename.endswith(".docx")
    paragraphs = _paragraph_texts(docx_bytes)
    assert "使用者提問" in paragraphs
    assert "AI 代理人回答" in paragraphs
    assert paragraphs.index("使用者提問") < paragraphs.index("AI 代理人回答")
    assert paragraphs.index("AI 代理人回答") < paragraphs.index("重點摘要如下。")


def test_export_docx_answer_heading_does_not_remove_numbered_report_sections() -> None:
    docx_bytes, _ = ExportWordService().build_docx_bytes(
        {
            "question": "請綜整漁船祥發23號的活動樣態。",
            "answer": "\n".join(
                [
                    "**一、資料概況**",
                    "- 文件數 / 片段數：6 / 6",
                    "",
                    "**二、活動時間線（依日期排列）**",
                    "- 2026-05-25：東石漁港外海活動。",
                    "",
                    "**三、趨勢分析**",
                    "1. 作業區域變化",
                ]
            ),
        }
    )

    doc = Document(BytesIO(docx_bytes))
    paragraph_texts = [paragraph.text for paragraph in doc.paragraphs]
    paragraph_styles = {paragraph.text: paragraph.style.name for paragraph in doc.paragraphs}

    assert paragraph_texts.index("AI 代理人回答") < paragraph_texts.index("一、資料概況")
    assert paragraph_texts.index("一、資料概況") < paragraph_texts.index("二、活動時間線（依日期排列）")
    assert paragraph_texts.index("二、活動時間線（依日期排列）") < paragraph_texts.index("三、趨勢分析")
    assert paragraph_styles["一、資料概況"] == "Heading 2"
    assert paragraph_styles["二、活動時間線（依日期排列）"] == "Heading 2"
    assert paragraph_styles["三、趨勢分析"] == "Heading 2"


def test_export_docx_answer_heading_preserves_markdown_table_smart_merge() -> None:
    docx_bytes, _ = ExportWordService().build_docx_bytes(
        {
            "question": "請綜整漁船祥發23號的活動樣態。",
            "answer": "\n".join(
                [
                    "**二、活動時間線（依日期排列）**",
                    "| 船名 | 海域 | 日期 | 活動 |",
                    "| --- | --- | --- | --- |",
                    "| 祥發23號 | 東石外海 | 2026-05-25 | 作業 |",
                    "| 祥發23號 | 東石外海 | 2026-05-28 | 移動 |",
                    "| 祥發23號 | 八斗子 | 2026-05-29 | 返港 |",
                ]
            ),
        }
    )

    doc = Document(BytesIO(docx_bytes))
    paragraph_texts = [paragraph.text for paragraph in doc.paragraphs]

    assert paragraph_texts.index("AI 代理人回答") < paragraph_texts.index("二、活動時間線（依日期排列）")
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert [table.cell(0, col).text for col in range(4)] == ["船名", "海域", "日期", "活動"]
    assert table.cell(1, 0)._tc is table.cell(2, 0)._tc
    assert table.cell(1, 1)._tc is table.cell(2, 1)._tc
    assert table.cell(1, 2)._tc is not table.cell(2, 2)._tc


def test_export_docx_lists_stage_models_at_end_of_data_agent_section() -> None:
    docx_bytes, _ = ExportWordService().build_docx_bytes(
        {
            "question": "請整理本週資料。",
            "answer": "整理完成。",
            "data_agent": {
                "model_usage": [
                    {"stage": "任務路由", "model_id": "router-model"},
                    {"stage": "資料整備", "model_id": "data-model"},
                    {"stage": "報告生成", "model_id": "report-model"},
                ],
                "runs": [{"mode": "SRA_REPORT", "question_raw": "請整理本週資料。"}],
            },
        }
    )

    paragraphs = _paragraph_texts(docx_bytes)
    heading_index = paragraphs.index("各階段使用模型")
    assert paragraphs[heading_index + 1:heading_index + 4] == [
        "任務路由：router-model",
        "資料整備：data-model",
        "報告生成：report-model",
    ]
    assert heading_index > paragraphs.index("原始問題：請整理本週資料。")
    assert heading_index == len(paragraphs) - 4


def test_export_docx_displays_relative_time_reference() -> None:
    docx_bytes, _ = ExportWordService().build_docx_bytes(
        {
            "question": "請整理去年資料。",
            "answer": "整理完成。",
            "data_agent": {
                "time_context": {
                    "enabled": True,
                    "reference_time": "2026-08-25T12:00:00+08:00",
                    "timezone": "Asia/Taipei",
                    "previous_year": "2025",
                },
                "runs": [{"mode": "SRA_REPORT", "question_raw": "請整理去年資料。"}],
            },
        }
    )

    assert "相對時間解析基準：2026-08-25T12:00:00+08:00 (Asia/Taipei；去年=2025)" in _paragraph_texts(docx_bytes)


def test_collect_llm_execution_reports_success_and_fallback() -> None:
    trace = {
        "llm": {"router": [{"model_id": "router", "elapsed_sec": 1.25}]},
        "query_refiner": {
            "report": {
                "enabled": True,
                "model_id": "refiner",
                "note": "error=TimeoutError: timed out",
                "elapsed_sec": 3.0,
            }
        },
        "data_agent": {"runs": [{"model_id_effective": "data", "elapsed_sec": 2.0}]},
    }

    assert _collect_llm_execution(trace) == [
        {"stage": "任務路由", "status": "success", "model_id": "router", "elapsed_sec": 1.25},
        {
            "stage": "問題精煉",
            "status": "fallback",
            "model_id": "refiner",
            "elapsed_sec": 3.0,
            "detail": "TimeoutError: timed out",
        },
        {"stage": "資料整備", "status": "success", "model_id": "data", "elapsed_sec": 2.0},
    ]


def test_export_docx_integrates_llm_execution_into_model_list_at_the_end() -> None:
    docx_bytes, _ = ExportWordService().build_docx_bytes(
        {
            "question": "請整理資料。",
            "answer": "完成。",
            "reasoning": "先分析資料。",
            "llm_execution": [
                {"stage": "任務路由", "status": "success", "model_id": "router", "elapsed_sec": 1.25},
                {
                    "stage": "問題精煉",
                    "status": "fallback",
                    "model_id": "refiner",
                    "detail": "TimeoutError: timed out",
                },
            ],
        }
    )

    paragraphs = _paragraph_texts(docx_bytes)
    heading = paragraphs.index("各階段使用模型")
    assert heading > paragraphs.index("附錄：推理內容")
    assert "附錄：LLM 各階段執行情況" not in paragraphs
    assert paragraphs[heading + 1:] == [
        "任務路由：router；執行狀態：正常完成；耗時：1.250 秒",
        "問題精煉：refiner；執行狀態：發生錯誤，已採用備援結果\n錯誤／備註：TimeoutError: timed out",
    ]


def test_export_docx_sorts_model_list_by_processing_flow() -> None:
    docx_bytes, _ = ExportWordService().build_docx_bytes(
        {
            "question": "請整理資料。",
            "answer": "完成。",
            "llm_execution": [
                {"stage": "證據欄位選擇", "status": "success", "model_id": "selector"},
                {"stage": "報告生成", "status": "success", "model_id": "report"},
                {"stage": "問題精煉", "status": "success", "model_id": "refiner"},
                {"stage": "資料整備", "status": "success", "model_id": "data"},
                {"stage": "任務路由", "status": "success", "model_id": "router"},
            ],
        }
    )

    paragraphs = _paragraph_texts(docx_bytes)
    heading = paragraphs.index("各階段使用模型")
    assert [text.split("：", 1)[0] for text in paragraphs[heading + 1:]] == [
        "任務路由",
        "問題精煉",
        "資料整備",
        "證據欄位選擇",
        "報告生成",
    ]
